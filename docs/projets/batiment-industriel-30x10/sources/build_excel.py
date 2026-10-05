# -*- coding: utf-8 -*-
"""Classeur Excel modifiable (formules) - Bâtiment industriel 30 x 10 m. Auteur : Moulo Jean Claude."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL
import modele as Mo
import entreprise as ENT
from openpyxl.drawing.image import Image as XLImage

OUTX = sys.argv[1]
Mo.compute()
AUTEUR = 'Moulo Jean Claude'
F = 'Arial'
BLEU_IN = Font(name=F, size=10, color='0000FF')
NOIR = Font(name=F, size=10)
GRAS = Font(name=F, size=10, bold=True)
BLANC = Font(name=F, size=10, bold=True, color='FFFFFF')
JAUNE = PatternFill('solid', fgColor='FFF2CC')
ENTETE = PatternFill('solid', fgColor='1F3A5F')
LOT = PatternFill('solid', fgColor='D6E4F0')
SOUS = PatternFill('solid', fgColor='F2F2F2')
TOT = PatternFill('solid', fgColor='BDD7EE')
fin = Side(style='thin', color='999999')
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)
WRAP = Alignment(wrap_text=True, vertical='top')
FCFA = '#,##0;(#,##0);"-"'
QTE = '#,##0.00;(#,##0.00);"-"'
QTE3 = '#,##0.000;(#,##0.000);"-"'

wb = Workbook()
wb.properties.creator = AUTEUR
wb.properties.lastModifiedBy = AUTEUR
wb.properties.subject = ENT.NOM
wb.properties.title = 'Bâtiment industriel 30 x 10 m - Métré, déboursé sec et DQE'

def nom(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

def feuille(titre, sous_titre, larg, ws=None):
    ws = ws or wb.create_sheet(titre)
    ws.title = titre
    ws['A1'] = f'{ENT.NOM} - {ENT.ACTIVITE} - {ENT.LIGNE_CONTACT}'
    ws['A1'].font = Font(name=F, size=12, bold=True, color='1F3A5F')
    ws['A2'] = (f'{sous_titre} - Bâtiment industriel 30,00 x 10,00 m - 04/10/2026 - '
                f'établi par {ENT.AUTEUR}, {ENT.FONCTION} - BAEL 91 mod. 99 / EC2')
    ws['A2'].font = Font(name=F, size=10, italic=True, color='1F3A5F')
    for i, w in enumerate(larg, 1):
        ws.column_dimensions[CL(i)].width = w
    ws.oddHeader.left.text = f'{ENT.NOM} - établi par {ENT.AUTEUR}'
    ws.oddHeader.right.text = sous_titre
    ws.oddFooter.left.text = f'{ENT.NOM} - N° CC {ENT.CC} - RCCM {ENT.RCCM}'
    ws.oddFooter.right.text = 'Page &P / &N'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    return ws

def entete(ws, r, cols):
    for i, h in enumerate(cols, 1):
        x = ws.cell(r, i, h); x.font = BLANC; x.fill = ENTETE; x.border = BORD
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = ws.cell(r + 1, 1)

def put(ws, r, col, v, font=NOIR, fmt=None, fill=None, al=None):
    x = ws.cell(r, col, v); x.font = font; x.border = BORD
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    if al: x.alignment = al
    return x

def qname(code):
    return 'q_' + code.replace('.', '_')

# =============================================================== LISEZ-MOI
ws = feuille('Lisez-moi', 'Notice du classeur', [3, 28, 95], wb.active)
lignes = [
    ('ENTREPRISE', ''),
    ('Raison sociale', f'{ENT.NOM} ({ENT.SIGLE}) - {ENT.ACTIVITE}'),
    ('Siège', ENT.VILLE),
    ('Téléphone', ENT.TEL),
    ('Email', ENT.EMAIL),
    ('N° CC / RCCM', f'{ENT.CC} / {ENT.RCCM}'),
    ('Banque', ENT.BANQUE),
    ('Régime fiscal', ENT.IMPOT),
    ('', ''),
    ('Objet', 'Métré, déboursé sec et devis quantitatif et estimatif (DQE) d\'un bâtiment industriel 30 x 10 m, '
              'murs en agglos 15 pleins en fondation (1 m) et 15 creux en élévation (6 m), semelle filante 60 x 15, piliers IPE 220 tous les 5 m encadrés de 2 poteaux BA, contreventements en cornières 50x50x5, '
              'entrée 6 m, toiture 2 versants à 20 %, '
              'dallage BA 20 cm nappe simple HA10, semelles isolées 80 x 80 x 30 en HA12 avec tiges d\'ancrage.'),
    ('Auteur', 'Moulo Jean Claude - Technicien Génie Civil BTP - Abidjan'),
    ('Date', '04/10/2026 - prix du marché d\'Abidjan à confirmer par cotations fournisseurs'),
    ('', ''),
    ('COMMENT MODIFIER', ''),
    ('Cellules bleues sur fond jaune', 'Données modifiables : dimensions, épaisseurs, prix unitaires, rendements, coefficients. '
                                        'Tout le reste se recalcule automatiquement.'),
    ('Texte noir', 'Formules - ne pas écraser.'),
    ('Sol de bonne portance', 'Feuille « Paramètres » : mettre p_EPLAT = 0,20 (une seule couche de latérite) au lieu de 0,40.'),
    ('Changer un prix', 'Feuille « Prix » colonne D : la modification se répercute sur tous les sous-détails, le DS et le DQE.'),
    ('Changer un rendement', 'Feuille « Sous-détails » colonne F (quantité par unité d\'ouvrage).'),
    ('', ''),
    ('FEUILLES', ''),
    ('Paramètres', 'Dimensions, hypothèses, coefficients K et TVA, grandeurs calculées (noms p_...).'),
    ('Prix', 'Prix de base : matériaux, main-d\'œuvre (h.j), matériel (noms px_...).'),
    ('Aciers', 'Nomenclature des aciers HA6 / HA10 / HA12 par lot, nombre de barres de 12 m.'),
    ('Métré', 'Avant-métré : formule de chaque quantité (noms q_...).'),
    ('Sous-détails', 'Décomposition matériaux / main-d\'œuvre / matériel de chaque prix unitaire.'),
    ('Déboursé sec', 'Déboursé sec par article et par lot, répartition par nature.'),
    ('DQE', 'Devis quantitatif et estimatif : PU = DS x K arrondi à 5 FCFA, récapitulatif HT / TVA 18 % / TTC.'),
    ('Matériaux et MO', 'Liste de commande des matériaux et besoins en main-d\'œuvre et matériel.'),
    ('', ''),
    ('HYPOTHÈSES CLÉS', ''),
    ('Piliers', '17 piliers en IPE 220 (S235) de ±0,00 à +6,00 sur platine 15 mm et 4 tiges d\'ancrage M20 ; '
                '2 poteaux BA 15x15 (4 HA10, cadres HA6 e = 15) de part et d\'autre de chaque IPE.'),
    ('Contreventements', 'Cornières 50x50x5 : 4 palées verticales en croix (travées d\'extrémité des long pans) '
                         'et poutre au vent dans 2 travées de toiture.'),
    ('Semelles isolées', '80 x 80 x 30 cm, fond de fouille à -1,20 m, nappe HA12 e = 15 cm dans les 2 sens, '
                         'fût BA 40 x 40 (4 HA12) jusqu\'à ±0,00 recevant les 4 tiges d\'ancrage M20 de la platine de l\'IPE.'),
    ('Dallage', '20 cm, béton dosé à 350 kg/m³, nappe simple HA10 e = 20 cm (p_NNAP = 2 pour double nappe), chaises HA10, polyane, joints sciés maille 5 m.'),
    ('Forme', 'Purge et rechargement en latérite par couches de 20 cm compactées à 95 % OPM : 40 cm si mauvaise portance '
              '(valeur par défaut), 20 cm sinon.'),
    ('Options', 'Lots 6 à 8 (enduits, fermes + pannes Z 120x2 + tôles bacs 5 ondes, portail) chiffrés séparément.'),
]
img = XLImage(ENT.LOGO); img.height, img.width = 95, 140
ws.add_image(img, 'D1')
ws.column_dimensions['D'].width = 22
r = 4
for a, b in lignes:
    ws.cell(r, 2, a).font = GRAS if (a.isupper() or a in ('Objet', 'Auteur', 'Date')) else NOIR
    ws.cell(r, 3, b).font = NOIR; ws.cell(r, 3).alignment = WRAP
    if a == 'Cellules bleues sur fond jaune':
        ws.cell(r, 2).font = BLEU_IN; ws.cell(r, 2).fill = JAUNE
    r += 1

# =============================================================== PARAMÈTRES
ws = feuille('Paramètres', 'Paramètres et hypothèses', [12, 50, 13, 8, 55])
entete(ws, 4, ['Nom', 'Désignation', 'Valeur', 'Unité', 'Commentaire / source'])
r = 5; grp = None
for n, lib, v, u, com, g in Mo.PARAMS:
    if g != grp:
        grp = g
        put(ws, r, 1, g.upper(), GRAS, fill=LOT)
        for col in range(2, 6): put(ws, r, col, None, fill=LOT)
        r += 1
    put(ws, r, 1, n); put(ws, r, 2, lib)
    fmt = '0.0%' if u == '%' else ('0.0000' if n == 'p_K' else '#,##0.00#')
    if isinstance(v, str):
        put(ws, r, 3, v, NOIR, fmt)
    else:
        put(ws, r, 3, v, BLEU_IN, fmt, JAUNE)
    put(ws, r, 4, u, al=Alignment(horizontal='center')); put(ws, r, 5, com)
    nom(n, f"'Paramètres'!$C${r}")
    r += 1

# =============================================================== PRIX
ws = feuille('Prix', 'Prix de base', [14, 58, 9, 16, 14])
entete(ws, 4, ['Code', 'Désignation', 'Unité', 'Prix unitaire (FCFA)', 'Catégorie'])
r = 5; cat = None
for code, lib, u, pu, ct in Mo.PRIX:
    if ct != cat:
        cat = ct
        put(ws, r, 1, ct.upper(), GRAS, fill=LOT)
        for col in range(2, 6): put(ws, r, col, None, fill=LOT)
        r += 1
    put(ws, r, 1, code); put(ws, r, 2, lib); put(ws, r, 3, u, al=Alignment(horizontal='center'))
    put(ws, r, 4, pu, BLEU_IN, FCFA, JAUNE); put(ws, r, 5, ct)
    nom('px_' + code, f"'Prix'!$D${r}")
    r += 1
ws.cell(r + 1, 2, 'Source : estimation du marché d\'Abidjan, octobre 2026 - à actualiser avec les cotations des fournisseurs.').font = \
    Font(name=F, size=9, italic=True)

# =============================================================== ACIERS
ws = feuille('Aciers', 'Nomenclature des aciers', [7, 75, 9, 14, 10, 14])
entete(ws, 4, ['Lot', 'Élément', 'Nuance', 'Longueur (m)', 'kg/m', 'Poids (kg)'])
r0 = 5
for i, (lot, elt, nu, f) in enumerate(Mo.ACIERS):
    r = r0 + i
    put(ws, r, 1, lot, al=Alignment(horizontal='center')); put(ws, r, 2, elt); put(ws, r, 3, nu, al=Alignment(horizontal='center'))
    put(ws, r, 4, f, fmt=QTE); put(ws, r, 5, '=' + Mo.KGNAME[nu], fmt='0.000'); put(ws, r, 6, f'=D{r}*E{r}', fmt=QTE)
r1 = r0 + len(Mo.ACIERS) - 1
RA = lambda col: f"'Aciers'!${col}${r0}:${col}${r1}"
r = r1 + 2
entete_r = r
for i, h in enumerate(['Lot', 'Récapitulatif', 'Nuance', 'Poids net (kg)', 'Chutes', 'Poids à commander (kg)', 'Barres de 12 m'], 1):
    x = ws.cell(r, i, h); x.font = BLANC; x.fill = ENTETE; x.border = BORD
    x.alignment = Alignment(horizontal='center', wrap_text=True)
ws.column_dimensions['G'].width = 14
r += 1
ACIER_CELL = {}
for lot in sorted({a[0] for a in Mo.ACIERS}):
    for nu in ('HA12', 'HA10', 'HA6'):
        if not any(a[0] == lot and a[2] == nu for a in Mo.ACIERS):
            continue
        put(ws, r, 1, lot, al=Alignment(horizontal='center')); put(ws, r, 2, f'Lot {lot}'); put(ws, r, 3, nu)
        put(ws, r, 4, f'=SUMIFS({RA("F")},{RA("A")},A{r},{RA("C")},C{r})', fmt=QTE)
        put(ws, r, 5, '=p_CHUTES', fmt='0.00'); put(ws, r, 6, f'=D{r}*E{r}', fmt=QTE)
        put(ws, r, 7, f'=ROUNDUP(F{r}/(12*{Mo.KGNAME[nu]}),0)', fmt='#,##0')
        ACIER_CELL[(lot, nu)] = f"'Aciers'!$F${r}"
        r += 1
for nu in ('HA12', 'HA10', 'HA6'):
    put(ws, r, 2, f'TOTAL {nu}', GRAS, fill=TOT); put(ws, r, 3, nu, GRAS, fill=TOT)
    put(ws, r, 4, f'=SUMIFS({RA("F")},{RA("C")},C{r})', GRAS, QTE, TOT); put(ws, r, 5, '=p_CHUTES', GRAS, '0.00', TOT)
    put(ws, r, 6, f'=D{r}*E{r}', GRAS, QTE, TOT); put(ws, r, 7, f'=ROUNDUP(F{r}/(12*{Mo.KGNAME[nu]}),0)', GRAS, '#,##0', TOT)
    put(ws, r, 1, None, fill=TOT)
    r += 1

# =============================================================== MÉTRÉ
ws = feuille('Métré', 'Avant-métré', [7, 6, 70, 6, 70, 14])
entete(ws, 4, ['N°', 'Lot', 'Désignation', 'U', 'Formule (noms de la feuille Paramètres)', 'Quantité'])
r = 5
for lot, lname, opt in Mo.LOTS:
    put(ws, r, 1, f'LOT {lot}', GRAS, fill=LOT); put(ws, r, 3, lname, GRAS, fill=LOT)
    for col in (2, 4, 5, 6): put(ws, r, col, None, fill=LOT)
    r += 1
    for a in [a for a in Mo.ARTS if a['lot'] == lot]:
        put(ws, r, 1, a['code'], al=Alignment(horizontal='center')); put(ws, r, 2, lot, al=Alignment(horizontal='center'))
        put(ws, r, 3, a['des'], al=WRAP); put(ws, r, 4, a['u'], al=Alignment(horizontal='center'))
        if isinstance(a['f'], tuple):
            _, l, nu = a['f']
            put(ws, r, 5, f'Nomenclature des aciers - lot {l} - {nu} (y compris chutes)', al=WRAP)
            put(ws, r, 6, '=' + ACIER_CELL[(l, nu)], fmt=QTE)
        else:
            put(ws, r, 5, a['f'].lstrip('='), al=WRAP)
            put(ws, r, 6, a['f'], fmt=QTE)
        nom(qname(a['code']), f"'Métré'!$F${r}")
        r += 1

# =============================================================== SOUS-DÉTAILS
ws = feuille('Sous-détails', 'Sous-détails de prix (déboursé sec unitaire)', [10, 14, 52, 12, 7, 11, 12, 13, 12, 13, 8])
entete(ws, 4, ['N°', 'Nature', 'Composante', 'Code prix', 'U', 'Qté / unité', 'P.U. (FCFA)', 'Montant / unité',
               'Qté ouvrage', 'Qté totale', 'Option'])
r = 5
SD = 'Sous-détails'
for a in Mo.ARTS:
    put(ws, r, 1, None, GRAS, fill=LOT)
    put(ws, r, 2, f"Article {a['code']} - {a['des']} (unité : {a['u']})", GRAS, fill=LOT)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
    r += 1; s = r
    for nat, lib, code, u, q in a['comp']:
        put(ws, r, 1, 'Art. ' + a['code'], al=Alignment(horizontal='center')); put(ws, r, 2, nat); put(ws, r, 3, lib)
        put(ws, r, 4, code); put(ws, r, 5, u, al=Alignment(horizontal='center'))
        put(ws, r, 6, q, BLEU_IN, QTE3, JAUNE); put(ws, r, 7, '=px_' + code, fmt=FCFA); put(ws, r, 8, f'=F{r}*G{r}', fmt=FCFA)
        put(ws, r, 9, '=' + qname(a['code']), fmt=QTE); put(ws, r, 10, f'=F{r}*I{r}', fmt=QTE)
        put(ws, r, 11, 'OUI' if Mo.OPTION[a['lot']] else 'NON', al=Alignment(horizontal='center'))
        r += 1
    e = r - 1
    put(ws, r, 1, 'Art. ' + a['code'], al=Alignment(horizontal='center')); put(ws, r, 2, 'Matériel'); put(ws, r, 3, 'Petit outillage (% de la main-d\'œuvre)')
    put(ws, r, 4, None); put(ws, r, 5, '%', al=Alignment(horizontal='center')); put(ws, r, 6, '=p_PO', fmt='0.0%')
    put(ws, r, 7, f'=SUMIFS(H{s}:H{e},B{s}:B{e},"Main-d\'œuvre")', fmt=FCFA); put(ws, r, 8, f'=F{r}*G{r}', fmt=FCFA)
    for col in (9, 10, 11): put(ws, r, col, None)
    r += 1
    put(ws, r, 3, f"DÉBOURSÉ SEC UNITAIRE / {a['u']}", GRAS, fill=SOUS)
    put(ws, r, 8, f'=SUM(H{s}:H{r-1})', GRAS, FCFA, SOUS)
    for col in (1, 2, 4, 5, 6, 7, 9, 10, 11): put(ws, r, col, None, fill=SOUS)
    r += 2

SDL = r
SDR = lambda col: f"'{SD}'!${col}$5:${col}${SDL}"
# =============================================================== DÉBOURSÉ SEC
ws = feuille('Déboursé sec', 'Déboursé sec', [7, 60, 6, 11, 13, 13, 12, 13, 15, 15, 14, 16, 8])
entete(ws, 4, ['N°', 'Désignation', 'U', 'Quantité', 'Matériaux / u', 'Main-d\'œuvre / u', 'Matériel / u', 'DS unitaire',
               'Matériaux total', 'Main-d\'œuvre total', 'Matériel total', 'DS total (FCFA)', 'Option'])
def sds(col, crit_nat, r):
    return f'=SUMIFS({SDR("H")},{SDR("A")},"Art. "&$A{r},{SDR("B")},"{crit_nat}")'
r = 5
DSROW = {}; DS_LOT = {}
for lot, lname, opt in Mo.LOTS:
    put(ws, r, 1, f'LOT {lot}', GRAS, fill=LOT); put(ws, r, 2, lname, GRAS, fill=LOT)
    for col in range(3, 14): put(ws, r, col, None, fill=LOT)
    r += 1; s = r
    for a in [a for a in Mo.ARTS if a['lot'] == lot]:
        put(ws, r, 1, a['code'], al=Alignment(horizontal='center')); put(ws, r, 2, a['des'], al=WRAP)
        put(ws, r, 3, a['u'], al=Alignment(horizontal='center')); put(ws, r, 4, '=' + qname(a['code']), fmt=QTE)
        put(ws, r, 5, sds('E', 'Matériaux', r), fmt=FCFA); put(ws, r, 6, sds('F', 'Main-d\'œuvre', r), fmt=FCFA)
        put(ws, r, 7, sds('G', 'Matériel', r), fmt=FCFA); put(ws, r, 8, f'=E{r}+F{r}+G{r}', fmt=FCFA)
        for i, col in enumerate('IJKL'):
            put(ws, r, 9 + i, f'=$D{r}*{"EFGH"[i]}{r}', fmt=FCFA)
        put(ws, r, 13, 'OUI' if opt else 'NON', al=Alignment(horizontal='center'))
        DSROW[a['code']] = r
        r += 1
    e = r - 1
    put(ws, r, 2, f'Sous-total déboursé sec lot {lot}', GRAS, fill=SOUS)
    for col in range(1, 14):
        if col in (9, 10, 11, 12):
            put(ws, r, col, f'=SUM({CL(col)}{s}:{CL(col)}{e})', GRAS, FCFA, SOUS)
        elif col != 2:
            put(ws, r, col, None, fill=SOUS)
    DS_LOT[lot] = r
    r += 2
put(ws, r, 2, 'DÉBOURSÉ SEC GROS ŒUVRE ET STRUCTURE (lots 0 à 5)', GRAS, fill=TOT)
for col in (9, 10, 11, 12):
    put(ws, r, col, f'=SUMIFS({CL(col)}5:{CL(col)}{r-1},$M$5:$M${r-1},"NON")', GRAS, FCFA, TOT)
rgo = r; r += 1
put(ws, r, 2, 'DÉBOURSÉ SEC OPTIONS (lots 6 à 8)', GRAS, fill=TOT)
for col in (9, 10, 11, 12):
    put(ws, r, col, f'=SUMIFS({CL(col)}5:{CL(col)}{rgo-1},$M$5:$M${rgo-1},"OUI")', GRAS, FCFA, TOT)
r += 1
put(ws, r, 2, 'DÉBOURSÉ SEC TOTAL', GRAS, fill=TOT)
for col in (9, 10, 11, 12):
    put(ws, r, col, f'={CL(col)}{rgo}+{CL(col)}{rgo+1}', GRAS, FCFA, TOT)
r += 1
put(ws, r, 2, 'Répartition gros œuvre (%)', GRAS)
for col in (9, 10, 11, 12):
    put(ws, r, col, f'=IF($L${rgo}=0,0,{CL(col)}{rgo}/$L${rgo})', GRAS, '0.0%')
DS_GO = f"'Déboursé sec'!$L${rgo}"; DS_OP = f"'Déboursé sec'!$L${rgo+1}"

# =============================================================== DQE
ws = feuille('DQE', 'Devis quantitatif et estimatif', [7, 66, 6, 12, 15, 17, 8])
entete(ws, 4, ['N°', 'Désignation', 'U', 'Quantité', 'P.U. HT (FCFA)', 'Montant HT (FCFA)', 'Option'])
r = 5; LOTROW = {}
for lot, lname, opt in Mo.LOTS:
    put(ws, r, 1, f'LOT {lot}', GRAS, fill=LOT); put(ws, r, 2, lname, GRAS, fill=LOT)
    for col in range(3, 8): put(ws, r, col, None, fill=LOT)
    r += 1; s = r
    for a in [a for a in Mo.ARTS if a['lot'] == lot]:
        put(ws, r, 1, a['code'], al=Alignment(horizontal='center')); put(ws, r, 2, a['des'], al=WRAP)
        put(ws, r, 3, a['u'], al=Alignment(horizontal='center')); put(ws, r, 4, '=' + qname(a['code']), fmt=QTE)
        put(ws, r, 5, f"=ROUND('Déboursé sec'!H{DSROW[a['code']]}*p_K/5,0)*5", fmt=FCFA)
        put(ws, r, 6, f'=D{r}*E{r}', fmt=FCFA); put(ws, r, 7, 'OUI' if opt else 'NON', al=Alignment(horizontal='center'))
        r += 1
    put(ws, r, 2, f'Sous-total lot {lot}', GRAS, fill=SOUS)
    put(ws, r, 6, f'=SUM(F{s}:F{r-1})', GRAS, FCFA, SOUS)
    for col in (1, 3, 4, 5, 7): put(ws, r, col, None, fill=SOUS)
    LOTROW[lot] = r
    r += 2
put(ws, r, 2, 'RÉCAPITULATIF', GRAS); r += 1
rec0 = r
for lot, lname, opt in Mo.LOTS:
    if opt: continue
    put(ws, r, 1, f'Lot {lot}'); put(ws, r, 2, lname); put(ws, r, 6, f'=F{LOTROW[lot]}', fmt=FCFA); r += 1
put(ws, r, 2, 'TOTAL GROS ŒUVRE HT (A)', GRAS, fill=TOT); put(ws, r, 6, f'=SUM(F{rec0}:F{r-1})', GRAS, FCFA, TOT); rA = r; r += 1
put(ws, r, 2, 'TVA'); put(ws, r, 5, '=p_TVA', fmt='0%'); put(ws, r, 6, f'=F{rA}*p_TVA', fmt=FCFA); r += 1
put(ws, r, 2, 'TOTAL GROS ŒUVRE TTC', GRAS, fill=TOT); put(ws, r, 6, f'=F{rA}+F{r-1}', GRAS, FCFA, TOT); rAT = r; r += 2
rec1 = r
for lot, lname, opt in Mo.LOTS:
    if not opt: continue
    put(ws, r, 1, f'Lot {lot}'); put(ws, r, 2, lname); put(ws, r, 6, f'=F{LOTROW[lot]}', fmt=FCFA); r += 1
put(ws, r, 2, 'TOTAL OPTIONS HT (B)', GRAS, fill=TOT); put(ws, r, 6, f'=SUM(F{rec1}:F{r-1})', GRAS, FCFA, TOT); rB = r; r += 2
put(ws, r, 2, 'TOTAL GÉNÉRAL HT (A + B)', GRAS, fill=TOT); put(ws, r, 6, f'=F{rA}+F{rB}', GRAS, FCFA, TOT); rG = r; r += 1
put(ws, r, 2, 'TVA'); put(ws, r, 5, '=p_TVA', fmt='0%'); put(ws, r, 6, f'=F{rG}*p_TVA', fmt=FCFA); r += 1
put(ws, r, 2, 'TOTAL GÉNÉRAL TTC (A + B)', GRAS, fill=TOT); put(ws, r, 6, f'=F{rG}+F{r-1}', GRAS, FCFA, TOT); rGT = r; r += 2
put(ws, r, 2, 'Ratio gros œuvre HT par m² d\'emprise'); put(ws, r, 6, f'=F{rA}/(p_L*p_B)', fmt=FCFA); r += 1
put(ws, r, 2, 'Marge brute gros œuvre (PV HT - déboursé sec)'); put(ws, r, 6, f'=F{rA}-{DS_GO}', fmt=FCFA); r += 2
ws.cell(r, 2, 'Fait à Abidjan, le 04/10/2026 - Établi par Moulo Jean Claude, Technicien Génie Civil BTP').font = Font(name=F, size=10, bold=True)
ws.cell(r + 1, 2, f'Pour {ENT.NOM} : cachet et signature').font = Font(name=F, size=10, bold=True)
ws.cell(r + 2, 2, f'{ENT.NOM} - {ENT.LIGNE_LEGALE}').font = Font(name=F, size=8, italic=True)
img2 = XLImage(ENT.LOGO); img2.height, img2.width = 75, 110
ws.add_image(img2, 'H1')
CELLS = {'GO_HT': ('DQE', f'F{rA}'), 'GO_TTC': ('DQE', f'F{rAT}'), 'OP_HT': ('DQE', f'F{rB}'), 'TG_TTC': ('DQE', f'F{rGT}'),
         'DS_GO': ('Déboursé sec', f'L{rgo}'), 'DS_OP': ('Déboursé sec', f'L{rgo+1}')}

# =============================================================== MATÉRIAUX ET MO
ws = feuille('Matériaux et MO', 'Besoins en matériaux, main-d\'œuvre et matériel', [14, 55, 8, 15, 15, 13, 16, 16, 22])
entete(ws, 4, ['Code', 'Désignation', 'U', 'Qté gros œuvre', 'Qté avec options', 'P.U. (FCFA)', 'Montant gros œuvre',
               'Montant avec options', 'Observation'])
r = 5; cat = None
SDJ = SDR('J'); SDD = SDR('D'); SDK = SDR('K')
for code, lib, u, pu, ct in Mo.PRIX:
    if ct != cat:
        cat = ct
        put(ws, r, 1, ct.upper(), GRAS, fill=LOT)
        for col in range(2, 10): put(ws, r, col, None, fill=LOT)
        r += 1
    put(ws, r, 1, code); put(ws, r, 2, lib); put(ws, r, 3, u, al=Alignment(horizontal='center'))
    put(ws, r, 4, f'=SUMIFS({SDJ},{SDD},A{r},{SDK},"NON")', fmt=QTE)
    put(ws, r, 5, f'=SUMIFS({SDJ},{SDD},A{r})', fmt=QTE)
    put(ws, r, 6, '=px_' + code, fmt=FCFA); put(ws, r, 7, f'=D{r}*F{r}', fmt=FCFA); put(ws, r, 8, f'=E{r}*F{r}', fmt=FCFA)
    obs = None
    if code in ('ha6', 'ha10', 'ha12'):
        kg = {'ha6': 'p_KG6', 'ha10': 'p_KG10', 'ha12': 'p_KG12'}[code]
        obs = f'=ROUNDUP(D{r}/(12*{kg}),0)&" barres de 12 m (GO)"'
    elif code == 'ciment':
        obs = f'=ROUND(D{r}*50/1000,1)&" t (gros œuvre)"'
    elif code == 'agglo':
        obs = '=ROUNDUP(D' + str(r) + ',0)&" agglos (GO)"'
    put(ws, r, 9, obs)
    r += 1
put(ws, r, 2, 'TOTAL (contrôle = déboursé sec, hors petit outillage)', GRAS, fill=TOT)
put(ws, r, 7, f'=SUM(G5:G{r-1})', GRAS, FCFA, TOT); put(ws, r, 8, f'=SUM(H5:H{r-1})', GRAS, FCFA, TOT)
for col in (1, 3, 4, 5, 6, 9): put(ws, r, col, None, fill=TOT)

wb.calculation.fullCalcOnLoad = True
wb.save(OUTX)
import json
json.dump(CELLS, open(OUTX + '.cells.json', 'w'))
print('ok', OUTX)
