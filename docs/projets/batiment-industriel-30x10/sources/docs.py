# -*- coding: utf-8 -*-
import os, math, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import calc as C

OUT = os.environ['OUT']; IMG = os.environ['IMG']
AUTEUR = 'Moulo Jean Claude'
DATE = '04/10/2026'
PROJET = 'Construction d\'un bâtiment industriel 30,00 x 10,00 m'
BLEU = RGBColor(0x1F, 0x3A, 0x5F)

# ------------------------------------------------------------------ utilitaires
def fnum(x, d=0):
    s = f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')
    return s.replace(' ', ' ')

def lettres(n):
    U = ['zéro', 'un', 'deux', 'trois', 'quatre', 'cinq', 'six', 'sept', 'huit', 'neuf', 'dix', 'onze', 'douze',
         'treize', 'quatorze', 'quinze', 'seize']
    D = {2: 'vingt', 3: 'trente', 4: 'quarante', 5: 'cinquante', 6: 'soixante'}
    def lt100(n):
        if n <= 16: return U[n]
        if n < 20: return 'dix-' + U[n - 10]
        d, u = divmod(n, 10)
        if d in (7, 9):
            base = 'soixante' if d == 7 else 'quatre-vingt'
            r = n - (60 if d == 7 else 80)
            if r == 11 and d == 7: return 'soixante et onze'
            return base + '-' + lt100(r)
        if d == 8: return 'quatre-vingts' if u == 0 else 'quatre-vingt-' + U[u]
        if u == 0: return D[d]
        if u == 1: return D[d] + ' et un'
        return D[d] + '-' + U[u]
    def lt1000(n, final=True):
        c, r = divmod(n, 100)
        out = []
        if c:
            out.append('cent' if c == 1 else U[c] + (' cents' if r == 0 and final else ' cent'))
        if r: out.append(lt100(r))
        return ' '.join(out)
    n = int(round(n))
    if n == 0: return 'zéro'
    parts = []
    mi, r = divmod(n, 1_000_000)
    mil, r = divmod(r, 1000)
    if mi: parts.append(lt1000(mi) + (' million' if mi == 1 else ' millions'))
    if mil: parts.append('mille' if mil == 1 else lt1000(mil, final=False) + ' mille')
    if r: parts.append(lt1000(r))
    return ' '.join(parts)

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol)
    tcPr.append(sh)

def field(run, code):
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)

def new_doc(titre_doc, landscape=False):
    d = Document()
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    for lvl, sz in ((1, 14), (2, 12), (3, 11)):
        h = d.styles[f'Heading {lvl}']; h.font.name = 'Calibri'; h.font.size = Pt(sz); h.font.color.rgb = BLEU
    s = d.sections[0]
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    else:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = Cm(2.6); s.bottom_margin = Cm(1.8)
    s.header_distance = Cm(0.8)
    # en-tête
    hd = s.header; p = hd.paragraphs[0]
    t = hd.add_table(rows=1, cols=2, width=s.page_width - s.left_margin - s.right_margin)
    c0, c1 = t.rows[0].cells
    r = c0.paragraphs[0].add_run('MOULO JEAN CLAUDE'); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = BLEU
    c0.add_paragraph('Technicien Génie Civil BTP - Abidjan, Côte d\'Ivoire').runs[0].font.size = Pt(8)
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p1.add_run(titre_doc); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = BLEU
    p2 = c1.add_paragraph(PROJET); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p2.runs[0].font.size = Pt(8)
    p3 = c1.add_paragraph(f'Date : {DATE} - Normes : BAEL 91 mod. 99 / EC2'); p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p3.runs[0].font.size = Pt(8)
    for c in (c0, c1): shade(c, 'EAF0F7')
    p.text = ''
    # pied de page
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f'Document établi par {AUTEUR} - Technicien Génie Civil BTP  |  Page '); r.font.size = Pt(8)
    r = fp.add_run(); r.font.size = Pt(8); field(r, 'PAGE')
    r = fp.add_run(' / '); r.font.size = Pt(8)
    r = fp.add_run(); r.font.size = Pt(8); field(r, 'NUMPAGES')
    cp = d.core_properties; cp.author = AUTEUR; cp.last_modified_by = AUTEUR; cp.title = f'{titre_doc} - {PROJET}'
    cp.created = datetime.datetime(2026, 10, 4)
    return d

def titre(d, t, st=None):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = BLEU
    if st:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(st); r.font.size = Pt(11); r.italic = True

def table(d, head, rows, widths=None, align=None, bold_rows=(), fill_rows=None, fs=8.5):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(fs); r.font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; shade(c, '1F3A5F')
    fill_rows = fill_rows or {}
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(fs)
            if k in bold_rows: r.bold = True
            if align and align[i] == 'r': cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if align and align[i] == 'c': cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if k in fill_rows: shade(cells[i], fill_rows[k])
    # répéter l'en-tête
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t

def puces(d, items):
    for it in items:
        d.add_paragraph(it, style='List Bullet')

MO_K = 'Main-d\'œuvre'
LOTNAME = {l: n for l, n, _ in C.LOTS}
OPTION = {l: o for l, _, o in C.LOTS}
q = C.V

# =================================================================== 1. ESQUISSE
def esquisse():
    d = new_doc('ESQUISSE', landscape=True)
    titre(d, 'ESQUISSE - BÂTIMENT INDUSTRIEL 30,00 x 10,00 m',
          'Piliers IPE 220 tous les 5,00 m + poteaux BA - Murs en agglos 15 pleins - Dallage BA 20 cm - Toiture 2 versants pente 20 %')
    d.add_heading('1. Données du projet', 1)
    table(d, ['Paramètre', 'Valeur retenue'], [
        ['Dimensions en plan (entre axes)', '30,00 m x 10,00 m - emprise 300 m², surface utile ≈ 294 m²'],
        ['Ossature', f'{C.N_PIL} piliers en IPE 220 (S235) tous les 5,00 m, de ±0,00 à +6,00, sur platine 320x220x15 '
                     'et 4 tiges d\'ancrage M20 ; de part et d\'autre de chaque IPE, 2 poteaux BA 15 x 15 (4 HA10, cadres HA6)'],
        ['Contreventements', 'Cornières 50 x 50 x 5 : 4 palées verticales en croix de Saint-André (travées d\'extrémité des long pans) '
                             '+ poutre au vent en toiture dans les 2 travées d\'extrémité'],
        ['Maçonnerie', 'Agglos 15 pleins (40 x 20 x 15) - fondation h = 1,00 m, élévation 6,00 m'],
        ['Chaînages', 'Bas (±0,00/+0,20), intermédiaire (+3,00/+3,20), haut (+5,80/+6,00) - 15 x 20, 4 HA10, cadres HA6 e = 20'],
        ['Fondations', 'Semelles isolées 80 x 80 x 30 sous piliers, fond de fouille à -1,20 m, nappe HA12 e = 15 dans les 2 sens, '
                       'fût BA 40 x 40 (4 HA12) jusqu\'à ±0,00 ; semelle filante 40 x 15 (4 HA10) sous les murs'],
        ['Dallage', 'Béton armé 20 cm dosé à 350 kg/m³, double nappe HA10 FeE500 e = 20 cm, polyane, joints sciés 5 x 5 m'],
        ['Forme sous dallage', 'Purge et rechargement en latérite par couches de 20 cm compactées à 95 % OPM : '
                               '40 cm si le sol est de mauvaise portance, 20 cm sinon'],
        ['Entrée', 'Ouverture 6,00 m (pignon Est), portail 6,00 x 4,50 m, linteau BA 15 x 40'],
        ['Toiture (option)', 'Deux versants, pente 20 % → faîtage à +7,00 m ; fermes, pannes Z 120 x 2 (entraxe ≤ 1,20 m), '
                             'tôles bacs 5 ondulations'],
    ], widths=[5, 20.5])
    d.add_heading('2. Interprétations et hypothèses à valider', 1)
    puces(d, [
        '« Autour de chaque pilier deux poteaux » : chaque pilier IPE 220 est encadré par 2 poteaux BA 15 x 15 '
        '(un de chaque côté, dans l\'épaisseur du mur) qui lient la maçonnerie et les chaînages à l\'ossature métallique.',
        'Les tiges d\'ancrage (4 M20 par pilier) sont scellées dans un fût BA 40 x 40 coulé sur la semelle 80 x 80 ; '
        'la platine est à ±0,00 et noyée dans le dallage.',
        'Disposition : 7 piliers par long pan, 1 au centre du pignon Ouest, 2 piliers d\'encadrement de l\'entrée '
        'sur le pignon Est → 17 piliers.',
        'Rechargement en latérite chiffré pour 40 cm (cas défavorable) ; mettre p_EPLAT = 0,20 dans le fichier Excel '
        'si les essais montrent une bonne portance.',
        'Le gros œuvre comprend les lots 0 à 5 (y compris la structure métallique) ; la charpente de toiture et la couverture, '
        'les enduits et le portail sont chiffrés en options (lots 6 à 8).',
        'Note vocale WhatsApp : elle n\'a pas pu être transcrite ; seuls les messages écrits ont été utilisés.',
    ])
    d.add_heading('3. Plans d\'esquisse', 1)
    for f, cap in [('01_plan.png', 'Plan d\'implantation - piliers IPE 220, poteaux BA, contreventements'),
                   ('02_facade_long_pan.png', 'Façade Nord (long pan) - palées de stabilité'),
                   ('03_pignon_entree.png', 'Pignon Est - façade d\'entrée'),
                   ('04_coupe_AA.png', 'Coupe transversale A-A'),
                   ('05_details.png', 'Détails type')]:
        d.add_page_break()
        if f in ('03_pignon_entree.png', '04_coupe_AA.png', '05_details.png'):
            d.add_picture(f'{IMG}/{f}', height=Cm(15.5))
        else:
            d.add_picture(f'{IMG}/{f}', width=Cm(26.0))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = d.add_paragraph(cap); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic = True
    d.add_page_break()
    d.add_heading('4. Tableau des éléments structuraux', 1)
    table(d, ['Élément', 'Section / dimensions', 'Armatures / profilé', 'Armatures transversales / assemblage', 'Matériau'], [
        ['Semelle isolée (17 u)', '0,80 x 0,80 x 0,30, fond -1,20', '2 x 6 HA12 e = 15 cm (retours 20 cm)', '-', 'Béton 350 kg/m³'],
        ['Fût sous platine (17 u)', '0,40 x 0,40, de -0,85 à ±0,00', '4 HA12 en L', 'Cadres HA6 e = 15 cm', 'Béton 350 kg/m³'],
        ['Tiges d\'ancrage', '4 par pilier (68 u)', 'M20 L = 600 coudées', '2 écrous + rondelle, gabarit de pose', 'Acier 4.6 / S235'],
        ['Pilier (17 u)', 'IPE 220, ±0,00 → +6,00', 'Platine de pied 320 x 220 x 15 + raidisseurs', 'Calage mortier sans retrait', 'S235, peint'],
        ['Poteaux BA (34 u)', '15 x 15, de -0,85 à +6,00', '4 HA10', 'Cadres HA6 e = 15 cm', 'Béton 350 kg/m³'],
        ['Contreventements', 'Croix en L 50 x 50 x 5', '4 palées + poutre au vent (2 travées)', 'Goussets 8 mm, boulons M12', 'S235, peint'],
        ['Semelle filante', '0,40 x 0,15', '4 HA10 filants', 'Répartiteurs HA6 e = 20 cm', 'Béton 350 kg/m³'],
        ['Chaînages bas / interm. / haut', '15 x 20', '4 HA10', 'Cadres HA6 e = 20 cm', 'Béton 350 kg/m³'],
        ['Linteau d\'entrée', '15 x 40, portée 6,00 m', '2 lits 3 HA10 (inf.) + 2 HA10 (sup.)', 'Cadres HA6 e = 15 cm', 'Béton 350 kg/m³'],
        ['Dallage', 'ép. 20 cm, 294 m²', '2 nappes HA10 e = 20 (2 sens)', 'Chaises HA10 1/m²', 'Béton 350 kg/m³'],
        ['Pannes (option)', 'Z 120 x 2, 12 lignes', 'Entraxe ≤ 1,20 m', 'Échantignoles boulonnées', 'Galvanisé'],
        ['Couverture (option)', '≈ 347 m²', 'Tôles bacs 5 ondulations', 'Faîtière, fixations', '-'],
    ], widths=[4.2, 4.6, 5.2, 6.2, 3.2])
    d.add_heading('5. Vérifications sommaires (BAEL 91 mod. 99 / EC2 / EC3)', 1)
    table(d, ['Vérification', 'Calcul', 'Conclusion'], [
        ['Semelle 80 x 80 - sol', 'Nser ≈ 42 kN (G) + 25 kN (Q) = 67 kN → σ = 0,067 / 0,64 = 0,105 MPa ≈ 1,05 bar < 1,5 bar', 'Vérifié (sol à confirmer)'],
        ['Semelle 80 x 80 - aciers', 'Nu ≈ 1,35 x 42 + 1,5 x 25 = 94 kN ; bielles : As = Nu (A - a) / (8 d fsu) = 0,094 x 0,40 / (8 x 0,25 x 435) '
         '= 0,43 cm² < 6 HA12 = 6,79 cm²', 'Vérifié (largement)'],
        ['IPE 220 - flexion (vent)', 'w = 1,5 x 0,6 kN/m² x 5 m = 4,5 kN/m ; M = 4,5 x 6² / 8 = 20,3 kN·m < Mpl = 285 cm³ x 235 = 67 kN·m',
         'Vérifié'],
        ['IPE 220 - flambement axe faible', 'l0 = 3,00 m (tenu par les chaînages) ; λ̄z = 1,29 ; χ ≈ 0,42 → Nb,Rd ≈ 330 kN > 94 kN', 'Vérifié'],
        ['Cornière 50 x 50 x 5 (traction)', 'A = 4,80 cm² → Npl = 113 kN ; effort diagonal estimé ≈ 19 kN', 'Vérifié'],
        ['Panne Z 120 x 2', 'qu ≈ 1,14 kN/m sur 5,00 m → M ≈ 3,6 kN·m ; W requis ≈ 15 cm³ (Z 120 x 2 : W ≈ 19 cm³)',
         'Vérifié - liernes à mi-portée'],
        ['Poteaux BA 15 x 15', '4 HA10 = 3,14 cm² → ρ = 1,4 % ; λ = 0,7 x 3,00 x √12 / 0,15 = 48,5 < 50 ; st = 15 cm ≤ 15 Øl', 'Vérifié'],
        ['Dallage', 'Double nappe HA10 e = 20 : 3,93 cm²/m par nappe et par sens > Amin = 0,13 % x 20 x 100 = 2,6 cm²/m (EC2)', 'Vérifié'],
        ['Linteau 6,00 m (ELU)', 'Mu ≈ 33,6 kN·m ; As = 2,33 cm² < 6 HA10 = 4,71 cm²', 'Vérifié'],
        ['Enrobage', 'Fondations 4 cm (XC2, sol latéritique humide) ; élévation et dallage 3 cm (climat tropical)', 'À respecter'],
    ], widths=[5.0, 15.0, 3.5])
    p = d.add_paragraph('Vérifications sommaires : les charges de la charpente, la pression du vent (NV 65 / EC1 adaptés '
                        'à la Côte d\'Ivoire) et la portance du sol doivent être confirmées par une note de calcul complète '
                        'et une étude de sol (essais LBTP) avant exécution.')
    p.runs[0].italic = True
    d.add_heading('6. Recommandations - contexte ivoirien', 1)
    puces(d, [
        'Sol latéritique : purger les poches argileuses ou organiques ; rechargement par couches de 20 cm arrosées et compactées '
        'à 95 % de l\'OPM, avec un contrôle de densité par couche.',
        'Tiges d\'ancrage : poser avec un gabarit en contreplaqué, vérifier l\'implantation et les niveaux avant bétonnage des fûts.',
        'Acier : décapage et 2 couches d\'antirouille + finition (humidité et air salin à Abidjan).',
        'Chaleur tropicale : cure humide du béton et du dallage pendant au moins 7 jours ; sciage des joints sous 24 h.',
        'Agglos : fabrication au moins 28 jours avant pose, arrosage 7 jours ; arase étanche hydrofugée sur le chaînage bas.',
        'Traitement anti-termites des fouilles et sous dallage ; gouttières et descentes EP à prévoir (non chiffrées).',
    ])
    d.save(f'{OUT}/01_Esquisse_Batiment_Industriel_30x10.docx')

# =================================================================== 2. DQE
def dqe():
    d = new_doc('DEVIS QUANTITATIF ET ESTIMATIF')
    titre(d, 'DEVIS QUANTITATIF ET ESTIMATIF (DQE)', 'Bâtiment industriel 30,00 x 10,00 m - Abidjan, Côte d\'Ivoire')
    d.add_heading('1. Bases du devis', 1)
    puces(d, [
        'Quantités issues de l\'avant-métré ci-dessous (dimensions de l\'esquisse du même jour).',
        f'Prix unitaires de vente = déboursé sec unitaire x coefficient K = {fnum(C.K, 4)} '
        f'(frais de chantier {int(C.FC*100)} %, frais généraux {int(C.FG*100)} %, aléas {int(C.ALEA*100)} %, '
        f'bénéfice {int(C.BEN*100)} %) - voir document « Déboursé sec ».',
        'Prix en FCFA, conditions du marché d\'Abidjan en octobre 2026, à actualiser avec les cotations des fournisseurs.',
        'TVA 18 % (Côte d\'Ivoire).',
        'Le gros œuvre et la structure (lots 0 à 5) correspondent à la demande ; les lots 6 à 8 (enduits, charpente de toiture + pannes Z 120 x 2 + tôles bacs 5 ondes, portail) sont des options.',
        'Le fichier Excel « Batiment_30x10_Metre_DS_DQE.xlsx » contient le même calcul avec des formules modifiables.',
    ])
    d.add_heading('2. Avant-métré', 1)
    rows = [[a['code'], a['des'], a['metre'], a['u'], fnum(a['q'], 2)] for a in C.ART]
    table(d, ['N°', 'Désignation', 'Détail du calcul', 'U', 'Quantité'], rows,
          widths=[1.0, 6.2, 7.0, 1.0, 2.2], align=['c', 'l', 'l', 'c', 'r'], fs=8)
    d.add_heading('3. Nomenclature des aciers', 1)
    rows = []
    for lot, elt, nu, lg in C.acier_detail:
        rows.append([f'Lot {lot}', elt, nu, fnum(lg, 1), fnum(lg * C.KG[nu], 1)])
    tot = {nu: sum(lg for _, _, n, lg in C.acier_detail if n == nu) for nu in C.KG}
    for nu in ('HA12', 'HA10', 'HA6'):
        rows.append(['', f'Total {nu} (hors chutes)', nu, fnum(tot[nu], 1), fnum(tot[nu] * C.KG[nu], 1)])
        rows.append(['', f'Total {nu} avec 5 % de chutes', nu, '', fnum(tot[nu] * C.KG[nu] * C.CHUTES, 1)])
        rows.append(['', f'Nombre de barres de 12 m à commander', nu, '',
                     f'{math.ceil(tot[nu] * C.KG[nu] * C.CHUTES / (12 * C.KG[nu]))} barres'])
    n = len(C.acier_detail)
    table(d, ['Lot', 'Élément', 'Nuance', 'Longueur (m)', 'Poids (kg)'], rows, widths=[1.3, 9.5, 1.5, 2.3, 2.6],
          align=['c', 'l', 'c', 'r', 'r'], bold_rows=range(n, n + 9), fs=8)
    d.add_page_break()
    d.add_heading('4. Devis quantitatif et estimatif', 1)
    rows, bold, fill = [], [], {}
    tot_lot = {}
    for lot, name, opt in C.LOTS:
        rows.append([f'LOT {lot}', name, '', '', '', '']); bold.append(len(rows) - 1); fill[len(rows) - 1] = 'D6E4F0'
        for a in [a for a in C.ART if a['lot'] == lot]:
            rows.append([a['code'], a['des'], a['u'], fnum(a['q'], 2), fnum(a['pv_u']), fnum(a['pv_tot'])])
        tot_lot[lot] = sum(a['pv_tot'] for a in C.ART if a['lot'] == lot)
        rows.append(['', f'Sous-total lot {lot}', '', '', '', fnum(tot_lot[lot])]); bold.append(len(rows) - 1)
        fill[len(rows) - 1] = 'F2F2F2'
    table(d, ['N°', 'Désignation', 'U', 'Qté', 'P.U. HT (FCFA)', 'Montant HT (FCFA)'], rows,
          widths=[1.1, 8.3, 0.9, 1.8, 2.3, 3.0], align=['c', 'l', 'c', 'r', 'r', 'r'], bold_rows=bold, fill_rows=fill, fs=8)
    d.add_heading('5. Récapitulatif', 1)
    go = sum(tot_lot[l] for l in tot_lot if not OPTION[l])
    op = sum(tot_lot[l] for l in tot_lot if OPTION[l])
    rows = [[f'Lot {l}', LOTNAME[l], fnum(tot_lot[l])] for l in tot_lot if not OPTION[l]]
    rows += [['', 'TOTAL GROS ŒUVRE ET STRUCTURE HT (A)', fnum(go)], ['', 'TVA 18 %', fnum(go * C.TVA)],
             ['', 'TOTAL GROS ŒUVRE TTC', fnum(go * (1 + C.TVA))]]
    rows += [[f'Lot {l}', LOTNAME[l], fnum(tot_lot[l])] for l in tot_lot if OPTION[l]]
    tg = go + op
    rows += [['', 'TOTAL OPTIONS HT (B)', fnum(op)], ['', 'TOTAL GÉNÉRAL HT (A + B)', fnum(tg)],
             ['', 'TVA 18 %', fnum(tg * C.TVA)], ['', 'TOTAL GÉNÉRAL TTC (A + B)', fnum(tg * (1 + C.TVA))]]
    nb = sum(1 for l in tot_lot if not OPTION[l]); no = sum(1 for l in tot_lot if OPTION[l])
    b = [nb, nb + 1, nb + 2, nb + 3 + no, nb + 4 + no, nb + 5 + no, nb + 6 + no]
    table(d, ['Lot', 'Désignation', 'Montant (FCFA)'], rows, widths=[1.5, 11.5, 4.0], align=['c', 'l', 'r'],
          bold_rows=b, fill_rows={nb + 2: 'D6E4F0', nb + 6 + no: 'D6E4F0'}, fs=9)
    ttc_go = round(go * (1 + C.TVA)); ttc = round(tg * (1 + C.TVA))
    p = d.add_paragraph()
    p.add_run('Arrêté le présent devis (gros œuvre) à la somme TTC de : ').bold = True
    p.add_run(f'{lettres(ttc_go)} francs CFA ({fnum(ttc_go)} FCFA).')
    p = d.add_paragraph()
    p.add_run('Montant TTC avec toutes les options : ').bold = True
    p.add_run(f'{lettres(ttc)} francs CFA ({fnum(ttc)} FCFA).')
    d.add_paragraph(f'Ratio gros œuvre : {fnum(go / 300)} FCFA HT/m² d\'emprise ; ratio global : {fnum(tg / 300)} FCFA HT/m².')
    d.add_paragraph()
    p = d.add_paragraph(f'Fait à Abidjan, le {DATE}'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = d.add_paragraph(f'{AUTEUR}\nTechnicien Génie Civil BTP'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].bold = True
    d.save(f'{OUT}/02_Devis_Quantitatif_Estimatif_Batiment_30x10.docx')
    return go, op

# =================================================================== 3. DÉBOURSÉ SEC
def ds():
    d = new_doc('DÉBOURSÉ SEC')
    titre(d, 'DÉBOURSÉ SEC ET SOUS-DÉTAILS DE PRIX', 'Bâtiment industriel 30,00 x 10,00 m - Abidjan, Côte d\'Ivoire')
    d.add_heading('1. Définition', 1)
    d.add_paragraph('Le déboursé sec (DS) est le coût direct de réalisation des ouvrages : matériaux + main-d\'œuvre + '
                    'matériel, sans frais de chantier, frais généraux, aléas, bénéfice ni TVA. Le petit outillage est '
                    'compté à 3 % de la main-d\'œuvre dans chaque sous-détail.')
    d.add_heading('2. Prix de base (Abidjan, octobre 2026 - à confirmer par cotations)', 1)
    for cat in ('Matériaux', 'Main-d\'œuvre', 'Matériel'):
        rows = [[lib, u, fnum(pu)] for code, lib, u, pu, ct in C.PRIX if ct == cat]
        table(d, [cat, 'Unité', 'Prix (FCFA)'], rows, widths=[11, 2, 3], align=['l', 'c', 'r'], fs=8.5)
    d.add_heading('3. Sous-détails des prix unitaires', 1)
    for a in C.ART:
        p = d.add_paragraph(); r = p.add_run(f"Article {a['code']} - {a['des']} (unité : {a['u']})"); r.bold = True
        r.font.color.rgb = BLEU
        rows = []
        for nat, lib, u, qq, pu in a['comp_v']:
            rows.append([nat, lib, u, fnum(qq, 3), fnum(pu), fnum(qq * pu)])
        po = 0.03 * a['ds_nat']['Main-d\'œuvre']
        if po:
            rows.append(['Matériel', 'Petit outillage (3 % de la MO)', '%', '3', '', fnum(po)])
        rows.append(['', f"DÉBOURSÉ SEC UNITAIRE / {a['u']}", '', '', '', fnum(a['ds_u'])])
        table(d, ['Nature', 'Composante', 'U', 'Qté/unité', 'P.U.', 'Montant'], rows,
              widths=[2.2, 7.6, 1.0, 1.8, 1.9, 2.4], align=['l', 'l', 'c', 'r', 'r', 'r'],
              bold_rows=[len(rows) - 1], fill_rows={len(rows) - 1: 'F2F2F2'}, fs=8)
    d.add_page_break()
    d.add_heading('4. Déboursé sec total par article', 1)
    rows, bold, fill = [], [], {}
    tot = {}
    for lot, name, opt in C.LOTS:
        rows.append([f'LOT {lot}', name, '', '', '', '', '', '']); bold.append(len(rows) - 1); fill[len(rows) - 1] = 'D6E4F0'
        arts = [a for a in C.ART if a['lot'] == lot]
        for a in arts:
            n = a['ds_nat']
            rows.append([a['code'], a['des'], a['u'], fnum(a['q'], 2), fnum(n['Matériaux'] * a['q']),
                         fnum(n['Main-d\'œuvre'] * a['q']), fnum(n['Matériel'] * a['q']), fnum(a['ds_tot'])])
        t = {k: sum(a['ds_nat'][k] * a['q'] for a in arts) for k in ('Matériaux', 'Main-d\'œuvre', 'Matériel')}
        t['tot'] = sum(t.values()); tot[lot] = t
        rows.append(['', f'Sous-total DS lot {lot}', '', '', fnum(t['Matériaux']), fnum(t['Main-d\'œuvre']),
                     fnum(t['Matériel']), fnum(t['tot'])]); bold.append(len(rows) - 1); fill[len(rows) - 1] = 'F2F2F2'
    table(d, ['N°', 'Désignation', 'U', 'Qté', 'Matériaux', 'Main-d\'œuvre', 'Matériel', 'DS total'], rows,
          widths=[0.9, 5.8, 0.8, 1.4, 2.2, 2.0, 1.7, 2.3], align=['c', 'l', 'c', 'r', 'r', 'r', 'r', 'r'],
          bold_rows=bold, fill_rows=fill, fs=7.5)
    d.add_heading('5. Récapitulatif du déboursé sec', 1)
    def somme(f):
        ls = [l for l in tot if f(l)]
        return {k: sum(tot[l][k] for l in ls) for k in ('Matériaux', 'Main-d\'œuvre', 'Matériel', 'tot')}
    go = somme(lambda l: not OPTION[l]); op = somme(lambda l: OPTION[l]); tg = somme(lambda l: True)
    nb = sum(1 for l in tot if not OPTION[l]); no = sum(1 for l in tot if OPTION[l])
    rows = [[f'Lot {l}', LOTNAME[l], fnum(tot[l]['Matériaux']), fnum(tot[l]['Main-d\'œuvre']), fnum(tot[l]['Matériel']),
             fnum(tot[l]['tot'])] for l in tot if not OPTION[l]]
    rows.append(['', 'DS GROS ŒUVRE ET STRUCTURE (A)', fnum(go['Matériaux']), fnum(go['Main-d\'œuvre']), fnum(go['Matériel']), fnum(go['tot'])])
    rows += [[f'Lot {l}', LOTNAME[l], fnum(tot[l]['Matériaux']), fnum(tot[l]['Main-d\'œuvre']), fnum(tot[l]['Matériel']),
              fnum(tot[l]['tot'])] for l in tot if OPTION[l]]
    rows.append(['', 'DS OPTIONS (B)', fnum(op['Matériaux']), fnum(op['Main-d\'œuvre']), fnum(op['Matériel']), fnum(op['tot'])])
    rows.append(['', 'DS TOTAL (A + B)', fnum(tg['Matériaux']), fnum(tg['Main-d\'œuvre']), fnum(tg['Matériel']), fnum(tg['tot'])])
    rows.append(['', 'Répartition gros œuvre', fnum(go['Matériaux']/go['tot']*100, 1) + ' %',
                 fnum(go[MO_K]/go['tot']*100, 1) + ' %', fnum(go['Matériel']/go['tot']*100, 1) + ' %', '100 %'])
    table(d, ['Lot', 'Désignation', 'Matériaux', 'Main-d\'œuvre', 'Matériel', 'DS total'], rows,
          widths=[1.2, 6.8, 2.5, 2.4, 2.0, 2.5], align=['c', 'l', 'r', 'r', 'r', 'r'],
          bold_rows=[nb, nb + 1 + no, nb + 2 + no, nb + 3 + no], fill_rows={nb: 'D6E4F0', nb + 2 + no: 'D6E4F0'}, fs=8)
    d.add_heading('6. Besoins en matériaux (liste de commande)', 1)
    for titre_b, base in (('Gros œuvre et structure (lots 0 à 5)', True), ('Avec options (lots 0 à 8)', False)):
        d.add_paragraph(titre_b).runs[0].bold = True
        need = C.besoins(base)
        rows = []
        for (lib, u, pu), qq in sorted(need.items(), key=lambda x: -x[0][2] * x[1]):
            if u == 'ft': continue
            arr = math.ceil(qq) if u in ('sac', 'u') else round(qq, 1)
            extra = ''
            if lib.startswith('Acier HA10'): extra = f' (≈ {math.ceil(qq / 7.40)} barres de 12 m)'
            if lib.startswith('Acier HA12'): extra = f' (≈ {math.ceil(qq / 10.66)} barres de 12 m)'
            if lib.startswith('Acier HA6'): extra = f' (≈ {math.ceil(qq / 2.66)} barres de 12 m)'
            if lib.startswith('Ciment'): extra = f' (≈ {fnum(qq * 50 / 1000, 1)} t)'
            rows.append([lib, u, fnum(arr, 0 if u in ('sac', 'u') else 1) + extra, fnum(qq * pu)])
        table(d, ['Matériau', 'U', 'Quantité', 'Montant (FCFA)'], rows, widths=[7.5, 1.2, 5.3, 3.0],
              align=['l', 'c', 'r', 'r'], fs=8)
    d.add_heading('7. Passage du déboursé sec au prix de vente', 1)
    d.add_paragraph(f'K = (1 + frais de chantier) / (1 - (frais généraux + aléas + bénéfice)) = (1 + {C.FC:.2f}) / '
                    f'(1 - ({C.FG:.2f} + {C.ALEA:.2f} + {C.BEN:.2f})) = {fnum(C.K, 4)}')
    rows = [['Déboursé sec', fnum(go['tot']), fnum(tg['tot'])],
            [f'Prix de vente HT (DS x K, PU arrondis)', fnum(sum(a['pv_tot'] for a in C.ART if not OPTION[a['lot']])),
             fnum(sum(a['pv_tot'] for a in C.ART))],
            ['Marge brute (frais + aléas + bénéfice)',
             fnum(sum(a['pv_tot'] for a in C.ART if not OPTION[a['lot']]) - go['tot']),
             fnum(sum(a['pv_tot'] for a in C.ART) - tg['tot'])]]
    table(d, ['Rubrique', 'Gros œuvre (FCFA)', 'Avec options (FCFA)'], rows, widths=[8, 4.5, 4.5],
          align=['l', 'r', 'r'], bold_rows=[1], fs=9)
    d.add_paragraph('Main-d\'œuvre estimée (gros œuvre) : ' +
                    fnum(go['Main-d\'œuvre'] / 7500) + ' hommes-jours environ (base moyenne 7 500 FCFA/h.j).')
    p = d.add_paragraph(f'Fait à Abidjan, le {DATE}'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = d.add_paragraph(f'{AUTEUR}\nTechnicien Génie Civil BTP'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].bold = True
    d.save(f'{OUT}/03_Debourse_Sec_Batiment_30x10.docx')
    return go, op

if __name__ == '__main__':
    esquisse()
    print('DQE', dqe())
    g, o = ds()
    print('DS', round(g['tot']), round(o['tot']))
