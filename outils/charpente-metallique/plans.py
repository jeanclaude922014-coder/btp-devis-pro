# -*- coding: utf-8 -*-
"""Esquisses dynamiques : coordonnées calculées par formules (feuille masquée COORD_PLANS)
et graphiques Nuage de points (XY) à lignes droites dans la feuille PLANS.

Principes :
- chaque esquisse a sa propre échelle s (même échelle en X et en Y) et ses décalages ox, oy,
  recalculés à partir de la boîte englobante réelle → mise à l'échelle automatique ;
- les axes des graphiques sont fixes (0 → XMAX, 0 → 100) avec un rapport XMAX / 100 égal
  au rapport largeur / hauteur de la zone de tracé → aucune déformation ;
- les ruptures de traits sont obtenues par #N/A (option « afficher #N/A comme cellule vide ») ;
- les cotes et repères sont des séries dont le nom (formule) est affiché en étiquette.
"""
from copy import copy
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.label import DataLabelList, DataLabel
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RichTextProperties
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

CH_W, CH_H = 25.5, 13.2           # taille des graphiques (cm)
PL_X, PL_Y, PL_W, PL_H = 0.01, 0.01, 0.98, 0.98
XMAX = round(100 * (CH_W * PL_W) / (CH_H * PL_H), 3)
YMAX = 100

MAXP, MAXIP, MAXPN, MAXPAN = 41, 13, 30, 30
MAXL, MAXLP, MAXCV, MAXPV = 10, 16, 10, 6
MAXC_LP, MAXC_PG, MAXD, MAXSKY, MAXEP = 32, 42, 4, 12, 16

BLEU, ORANGE, GRIS, ROUGE = "0070C0", "ED7D31", "7F7F7F", "C00000"
MARRON, NOIR, LATERITE, TRANSL = "C9A27E", "000000", "C55A11", "9DC3E6"


def S(color, w=1.0, dash=None, marker=None, msize=5, line=True):
    return dict(color=color, w=w, dash=dash, marker=marker, msize=msize, line=line)


class Drawing:
    def __init__(self, co, idx, title, xmin, xmax, ymin, ymax, ml=14, mr=14, mb=10, mt=10):
        self.co, self.idx, self.title, self.series = co, idx, title, []
        p = f"d{idx}_"
        dx = co.sc(p + "dx", f"MAX(0.001,({xmax})-({xmin}))")
        dy = co.sc(p + "dy", f"MAX(0.001,({ymax})-({ymin}))")
        self.s = co.sc(p + "s", f"MIN(({XMAX}-{ml}-{mr})/{dx},({YMAX}-{mb}-{mt})/{dy})")
        self.ox = co.sc(p + "ox", f"{ml}+(({XMAX}-{ml}-{mr})-{self.s}*{dx})/2-{self.s}*({xmin})")
        self.oy = co.sc(p + "oy", f"{mb}+(({YMAX}-{mb}-{mt})-{self.s}*{dy})/2-{self.s}*({ymin})")
        self.u = co.sc(p + "u", f"1/{self.s}")


class Coord:
    def __init__(self, ws, M, CO):
        self.ws, self.M, self.CO = ws, M, CO
        self.srow, self.trow, self.col = 5, 5, 7
        ws["A1"] = "COORDONNÉES DES ESQUISSES (feuille masquée, calculée) – les #N/A sont des ruptures de trait volontaires"
        ws["A1"].font = Font(name="Arial", bold=True)
        ws["A4"], ws["B4"], ws["D4"], ws["E4"] = "Grandeur", "Valeur", "Texte", "Contenu"

    def sc(self, name, expr):
        r = self.srow
        self.srow += 1
        self.ws.cell(r, 1, name)
        self.M.disp(self.CO, f"B{r}", expr)
        return f"$B${r}"

    def tx(self, name, expr):
        r = self.trow
        self.trow += 1
        self.ws.cell(r, 4, name)
        self.M.disp(self.CO, f"E{r}", expr)
        return f"'{self.CO}'!$E${r}"

    def ser(self, d, name, pts, style, label=None):
        c = self.col
        self.col += 2
        X, Y = get_column_letter(c), get_column_letter(c + 1)
        self.ws.cell(3, c, f"P{d.idx} – {name}")
        self.ws.cell(4, c, "X")
        self.ws.cell(4, c + 1, "Y")
        for i, p in enumerate(pts):
            r = 5 + i
            if p is None:
                xs = ys = "NA()"
            else:
                x, y = p[0], p[1]
                cond = p[2] if len(p) > 2 else None
                xs, ys = f"{d.ox}+{d.s}*({x})", f"{d.oy}+{d.s}*({y})"
                if cond:
                    xs, ys = f"IF({cond},{xs},NA())", f"IF({cond},{ys},NA())"
            self.M.disp(self.CO, f"{X}{r}", xs)
            self.M.disp(self.CO, f"{Y}{r}", ys)
        d.series.append(dict(name=name, col=c, r0=5, r1=4 + len(pts), style=style, label=label))

    def lab(self, d, text, x, y, cond=None, pos="ctr", size=8, rot=False, color=NOIR, bold=False):
        """Étiquette : série d'un point dont le nom (texte fixe ou référence) est affiché."""
        self.ser(d, "étiquette", [(x, y, cond)], S(NOIR, line=False),
                 label=dict(idx=0, text=text, pos=pos, size=size, rot=rot, color=color, bold=bold))

    def cote(self, d, x1, y1, x2, y2, text, horizontal, cond=None, pos=None):
        t = f"1.2*{d.u}"
        if horizontal:
            pts = [(x1, f"{y1}+{t}", cond), (x1, f"{y1}-{t}", cond), (x1, y1, cond),
                   (f"(({x1})+({x2}))/2", y1, cond), (x2, y2, cond), (x2, f"{y2}+{t}", cond), (x2, f"{y2}-{t}", cond)]
        else:
            pts = [(f"{x1}+{t}", y1, cond), (f"{x1}-{t}", y1, cond), (x1, y1, cond),
                   (x1, f"(({y1})+({y2}))/2", cond), (x2, y2, cond), (f"{x2}+{t}", y2, cond), (f"{x2}-{t}", y2, cond)]
        self.ser(d, "cote", pts, S(NOIR, 0.6),
                 label=dict(idx=3, text=text, pos=pos or ("t" if horizontal else "l"), size=8,
                            rot=not horizontal, color=NOIR, bold=False))


# ----------------------------------------------------------------------------- aides géométriques
def rect(x1, y1, x2, y2, cond=None):
    return [(x1, y1, cond), (x2, y1, cond), (x2, y2, cond), (x1, y2, cond), (x1, y1, cond), None]


def seg(x1, y1, x2, y2, cond=None):
    return [(x1, y1, cond), (x2, y2, cond), None]


def split_row(z, xa, xb, cuts, cond, clamp=False):
    """Ligne horizontale à la cote z, de xa à xb, interrompue par les baies (cuts = [(xl, xr, xm, h)]).
    Hors baie, le point de coupure devient un point de passage : la ligne reste continue."""
    def cl(v):
        return f"MAX({xa},MIN({xb},{v}))" if clamp else v
    pts = [(xa, z, cond)]
    for xl, xr, xm, h in cuts:
        pts += [(cl(f"IF({z}<{h},{xl},{xm})"), z, cond), (cl(xm), z, AND(cond, f"{z}>={h}")),
                (cl(f"IF({z}<{h},{xr},{xm})"), z, cond)]
    pts += [(xb, z, cond), None]
    return pts


def AND(*c):
    c = [x for x in c if x]
    return c[0] if len(c) == 1 else "AND(" + ",".join(c) + ")"


# ============================================================================= CONSTRUCTION
def build(wb, M, names, st, section, page_setup, protect):
    CO, PL = names["CO"], names["PL"]
    wsc = wb.create_sheet(CO)
    co = Coord(wsc, M, CO)
    sc, tx = co.sc, co.tx
    # ---- grandeurs reprises du modèle
    B = sc("B", "{portee}")
    L = sc("L", "{longueur}")
    H = sc("H", "{hauteur}")
    p = sc("p", "MAX(0.001,{pente})")
    Hm, Hf, h0 = sc("Hm", "{Hm}"), sc("Hf", "{Hf}"), sc("h0", "{h0}")
    nt, er = sc("nt", "{nt}"), sc("er", "{er}")
    n_ip, eppr = sc("n_ip", "{n_ip}"), sc("e_ppr", "{e_ppr}")
    npn, R1 = sc("np", "{np}"), sc("R1", "{R1}")
    treil, z0 = sc("treil", "{treil}"), sc("z0", "{z0}")
    kbard, kag = sc("kbard", "{kbard}"), sc("kag", "{kag}")
    nrows, nbcv = sc("n_rows", "{n_rows}"), sc("nb_cv", "{nb_cv}")
    npv, nl, npan = sc("n_pv", "{n_pv}"), sc("n_l", "{n_l}"), sc("n_pan", "{n_pan}")
    a1, a2, f1, f2 = sc("a1", "{a1}"), sc("a2", "{a2}"), sc("f1", "{f1}"), sc("f2", "{f2}")
    bfil, hfil, D = sc("b_fil", "{b_fil}"), sc("h_fil", "{h_fil}"), sc("D", "{D}")
    eprop, hs1, hfut, hmf = sc("e_prop", "{e_prop}"), sc("h_s1", "{h_s1}"), sc("h_fut", "{h_fut}"), sc("h_mf", "{h_mf}")
    erem, esab, edal = sc("e_rem", "{e_rem}"), sc("e_sable", "{e_sable}"), sc("e_dal", "{e_dal}")
    cf, cd, mai, kdbl = sc("c_fond", "{c_fond}"), sc("c_dal", "{c_dal}"), sc("maille", "{maille}"), sc("kdbl", "{kdbl}")
    fint = sc("file_int", "IF({n_int}>0,1,0)")
    deb = sc("deb", "{deb_pign}")
    nsky, asky, bsky = sc("nb_sky", "{nb_sky}"), sc("a_sky", "{a_sky}"), sc("b_sky", "{b_sky}")
    ndesc, pct = sc("nb_desc", "{nb_desc}"), sc("pct", "{pct_transl}")
    acr = sc("acr", 'IF({acrotere}="Oui",1,0)')
    hacr = sc("h_acr", "{h_acr}*" + acr)
    htr = sc("h_trav", "MAX(0.2,{h_trav}/1000)")
    lj = sc("l_jar", "{ratio_jarret}*{portee}")
    hpot = sc("h_pot", "MAX(0.15,{h_pot}/1000)")
    ltig = sc("l_tige", "{l_tige_port}")
    bch, hch, brd = sc("b_ch", "{b_ch}"), sc("h_ch", "{h_ch}"), sc("b_rd", "{b_rd}")
    nbch = sc("nb_ch", "{nb_ch}")
    hag = sc("h_ag", "{h_ag}")
    agg_full = sc("agglos_total", 'IF({fac}="Agglos 15 creux",1,0)')
    elis = sc("e_lis", "{e_lis}")
    epanr = sc("e_panr", "{e_panr}")
    pas_c = sc("pas_assises", "0.40")      # représentation : une ligne toutes les 2 assises d'agglos
    pas_c7 = sc("pas_assises_P7", "0.20")  # détail : chaque assise d'agglos

    # ---- baies (au plus 4 représentées par façade)
    doors = {}
    for F, W, n_div, e_div in [("Long-pan", L, nt, er), ("Pignon", B, n_ip, eppr)]:
        tag = "LP" if F == "Long-pan" else "PG"
        nC = sc(f"{tag}_nC", f'IF({{emp_pc}}="{F}",{{nb_pc}},0)')
        nQ = sc(f"{tag}_nQ", f'IF({{emp_pt}}="{F}",{{nb_pt}},0)')
        nd = sc(f"{tag}_nd", f"MIN({MAXD},{nC}+{nQ})")
        lst = []
        for k in range(1, MAXD + 1):
            ex = sc(f"{tag}_ex{k}", f"IF({k}<={nd},1,0)")
            w = sc(f"{tag}_w{k}", f"IF({k}<={nC},{{l_pc}},{{l_pt}})")
            h = sc(f"{tag}_h{k}", f"IF({k}<={nC},{{h_pc}},{{h_pt}})*{ex}")
            lo, hi = (f"IF({n_div}>=3,2,1)", f"IF({n_div}>=3,{n_div}-1,{n_div})") if tag == "LP" else ("1", n_div)
            pos = sc(f"{tag}_pos{k}", f"MAX({lo},MIN({hi},ROUND({k}*{n_div}/({nd}+1),0)))")
            xc = sc(f"{tag}_xc{k}", f"IF({ex}=1,({pos}-0.5)*{e_div},{W})")
            xl = sc(f"{tag}_xl{k}", f"IF({ex}=1,{xc}-{w}/2,{W})")
            xr = sc(f"{tag}_xr{k}", f"IF({ex}=1,{xc}+{w}/2,{W})")
            t = tx(f"{tag}_txt{k}", f'IF({k}<={nC},"Porte coul. ","Portillon ")&FIXED({w},2,TRUE)&" × "&FIXED({h},2,TRUE)')
            lst.append(dict(ex=ex, w=w, h=h, xc=xc, xl=xl, xr=xr, t=t, coul=f"{k}<={nC}"))
        doors[tag] = lst
    cuts = lambda tag: [(d["xl"], d["xr"], d["xc"], d["h"]) for d in doors[tag]]

    # ---- textes de cotes
    tL = tx("txt_L", '"Longueur "&FIXED({longueur},2,TRUE)&" m"')
    tB = tx("txt_B", '"Portée "&FIXED({portee},2,TRUE)&" m"')
    tE = tx("txt_e", '"Entraxe "&FIXED({er},2,TRUE)&" m ("&{nt}&" travées)"')
    tH = tx("txt_H", '"H sablière "&FIXED({hauteur},2,TRUE)&" m"')
    tHm = tx("txt_Hm", '"H égout "&FIXED({Hm},2,TRUE)&" m"')
    tHf = tx("txt_Hf", '"H faîtage "&FIXED({Hf},2,TRUE)&" m"')
    tP = tx("txt_p", '"Pente "&FIXED({pente}*100,1,TRUE)&" %"')
    tSys = tx("txt_sys", '{systeme}&" – poteaux "&{prof_pot}&IF({treil}=1," – membrures "&{nb_acc}&" "&{prof_msup}'
                         '&" / "&{nb_acc}&" "&{prof_minf}," – traverses "&{prof_trav})')
    tPan = tx("txt_pan", '"Pannes "&{prof_panne}&" : "&{np}&" par versant, e = "&FIXED({e_panr},2,TRUE)&" m"')
    tPig = tx("txt_pig", '"Poteaux de pignon "&{prof_pp}&" e = "&FIXED({e_ppr},2,TRUE)&" m"')

    D_ = []
    # ========================================================================= P1 VUE EN PLAN
    d = Drawing(co, 1, "VUE EN PLAN – IMPLANTATION", "0", L, "0", B, ml=20, mr=14, mb=16, mt=16)
    D_.append(d)
    u = d.u
    pts = []
    for i in range(MAXP):
        pts += seg(f"{er}*{i}", f"-4*{u}", f"{er}*{i}", f"{B}+4*{u}", f"{i}<={nt}")
    pts += seg(f"-4*{u}", "0", f"{L}+4*{u}", "0") + seg(f"-4*{u}", B, f"{L}+4*{u}", B)
    co.ser(d, "Axes des files", pts, S(NOIR, 0.4, "lgDashDot"))
    pts = rect(f"-{bfil}/2", f"-{bfil}/2", f"{L}+{bfil}/2", f"{B}+{bfil}/2") + \
        rect(f"{bfil}/2", f"{bfil}/2", f"{L}-{bfil}/2", f"{B}-{bfil}/2")
    co.ser(d, "Semelles filantes", pts, S(GRIS, 0.75, "dash"))
    pts = []
    for i in range(MAXP):
        c = f"{i}<={nt}"
        x = f"{er}*{i}"
        for y in ("0", B):
            pts += rect(f"{x}-{a1}/2", f"{y}-{a1}/2", f"{x}+{a1}/2", f"{y}+{a1}/2", c)
        pts += rect(f"{x}-{a1}/2", f"{B}/2-{a1}/2", f"{x}+{a1}/2", f"{B}/2+{a1}/2", AND(c, f"{fint}=1"))
    for k in range(1, MAXIP):
        c = f"{k}<={n_ip}-1"
        for x in ("0", L):
            pts += rect(f"{x}-{a2}/2", f"{k}*{eppr}-{a2}/2", f"{x}+{a2}/2", f"{k}*{eppr}+{a2}/2", c)
    co.ser(d, "Semelles isolées", pts, S(GRIS, 1.0))
    pts = []
    for i in range(MAXP):
        c = f"{i}<={nt}"
        pts += [(f"{er}*{i}", "0", c), (f"{er}*{i}", B, c), (f"{er}*{i}", f"{B}/2", AND(c, f"{fint}=1"))]
    co.ser(d, "Poteaux de portique", pts, S(BLEU, line=False, marker="square", msize=6))
    pts = []
    for k in range(1, MAXIP):
        c = f"{k}<={n_ip}-1"
        pts += [("0", f"{k}*{eppr}", c), (L, f"{k}*{eppr}", c)]
    co.ser(d, "Poteaux de pignon", pts, S(BLEU, line=False, marker="circle", msize=6))
    pts = []
    for dd in doors["LP"]:
        pts += seg(dd["xl"], f"-2.5*{u}", dd["xr"], f"-2.5*{u}", f"{dd['ex']}=1")
    for dd in doors["PG"]:
        pts += seg(f"-2.5*{u}", dd["xl"], f"-2.5*{u}", dd["xr"], f"{dd['ex']}=1")
    co.ser(d, "Baies", pts, S(NOIR, 2.5))
    for dd in doors["LP"]:
        co.lab(d, dd["t"], dd["xc"], f"-2.5*{u}", f"{dd['ex']}=1", pos="b", size=7)
    for dd in doors["PG"]:
        co.lab(d, dd["t"], f"-2.5*{u}", dd["xc"], f"{dd['ex']}=1", pos="r", size=7)
    co.cote(d, "0", f"-11*{u}", L, f"-11*{u}", tL, True, pos="b")
    co.cote(d, f"-12*{u}", "0", f"-12*{u}", B, tB, False)
    co.cote(d, "0", f"{B}+6*{u}", er, f"{B}+6*{u}", tE, True)
    for i in range(MAXP):
        co.lab(d, str(i + 1), f"{er}*{i}", f"{B}+11*{u}", f"{i}<={nt}", pos="ctr", size=8, bold=True)
    co.lab(d, "A", f"{L}+6*{u}", "0", pos="r", bold=True)
    co.lab(d, "B", f"{L}+6*{u}", B, pos="r", bold=True)

    # ========================================================================= P2 COUPE TRANSVERSALE
    d = Drawing(co, 2, "COUPE TRANSVERSALE – PORTIQUE TYPE", "0", B, f"-{D}-0.1", Hf, ml=20, mr=20, mb=12, mt=8)
    D_.append(d)
    u = d.u
    co.ser(d, "Sol / dallage", seg(f"-6*{u}", "0", f"{B}+6*{u}", "0"), S(NOIR, 1.0))
    pts = seg("0", "0", "0", Hm) + seg(B, "0", B, Hm) + \
        [("0", Hm), (f"{B}/2", Hf), (B, Hm), None] + \
        seg(f"{B}/2", "0", f"{B}/2", f"IF({treil}=1,{H},{Hf})", f"{fint}=1") + \
        seg("0", H, B, H, f"{treil}=1")
    co.ser(d, "Ossature", pts, S(BLEU, 2.25))
    a = f"({B}/MAX(1,{npan}))"
    top = lambda x: f"({Hm}+{p}*MIN({x},{B}-({x})))"
    pts = []
    for j in range(1, MAXPAN):
        x = f"{j}*{a}"
        pts += seg(x, H, x, top(x), AND(f"{treil}=1", f"{j}<={npan}-1"))
    for j in range(1, MAXPAN + 1):
        x0, x1 = f"{j - 1}*{a}", f"{j}*{a}"
        c = AND(f"{treil}=1", f"{j}<={npan}")
        left = f"{j}<={npan}/2"
        pts += seg(x0, f"IF({left},{top(x0)},{H})", x1, f"IF({left},{H},{top(x1)})", c)
    co.ser(d, "Treillis", pts, S(BLEU, 1.0))
    c = f"{treil}=0"
    pts = seg("0", f"{H}-{htr}", lj, f"{H}+{p}*{lj}", c) + seg(B, f"{H}-{htr}", f"{B}-{lj}", f"{H}+{p}*{lj}", c) + \
        [(f"{B}/2-{lj}/2", f"{Hf}-{p}*{lj}/2", c), (f"{B}/2", f"{Hf}-{htr}", c),
         (f"{B}/2+{lj}/2", f"{Hf}-{p}*{lj}/2", c), None]
    co.ser(d, "Jarrets", pts, S(BLEU, 1.5))
    pts = []
    for j in range(MAXPN):
        x = f"{j}*({B}/2)/MAX(1,{npn}-1)"
        c = f"{j}<={npn}-1"
        pts += [(x, f"{Hm}+{p}*{x}+1.2*{u}", c), (f"{B}-{x}", f"{Hm}+{p}*{x}+1.2*{u}", c)]
    co.ser(d, "Pannes", pts, S(BLEU, line=False, marker="square", msize=4))
    co.ser(d, "Faîtage", [(f"{B}/2", f"{Hf}+3*{u}")], S(BLEU, line=False, marker="triangle", msize=7))
    pts = [(f"-3.5*{u}", f"{Hm}+1.5*{u}"), (f"-3.5*{u}", f"{Hm}-0.8*{u}"), (f"-0.3*{u}", f"{Hm}-0.8*{u}"),
           (f"-0.3*{u}", f"{Hm}+1.5*{u}"), None,
           (f"{B}+3.5*{u}", f"{Hm}+1.5*{u}"), (f"{B}+3.5*{u}", f"{Hm}-0.8*{u}"), (f"{B}+0.3*{u}", f"{Hm}-0.8*{u}"),
           (f"{B}+0.3*{u}", f"{Hm}+1.5*{u}"), None]
    co.ser(d, "Chéneaux", pts, S(BLEU, 1.25))
    pts = []
    for x0, c in [("0", None), (B, None), (f"{B}/2", f"{fint}=1")]:
        zs = f"(-{D}+{eprop})"
        pts += rect(f"{x0}-{a1}/2", zs, f"{x0}+{a1}/2", f"{zs}+{hs1}", c)
        pts += seg(f"{x0}-{f1}/2", f"{zs}+{hs1}", f"{x0}-{f1}/2", "0", c) + \
            seg(f"{x0}+{f1}/2", f"{zs}+{hs1}", f"{x0}+{f1}/2", "0", c)
    co.ser(d, "Fondations", pts, S(GRIS, 1.0))
    co.cote(d, f"-9*{u}", "0", f"-9*{u}", H, tH, False)
    co.cote(d, f"{B}+9*{u}", "0", f"{B}+9*{u}", Hf, tHf, False, pos="r")
    co.cote(d, "0", f"-{D}-4*{u}", B, f"-{D}-4*{u}", tB, True, pos="b")
    co.lab(d, tP, f"{B}/4", f"{Hm}+{p}*{B}/4+5*{u}", pos="t")
    co.lab(d, tSys, f"{B}/2", f"{H}/2", pos="ctr", size=8)
    co.lab(d, tPan, f"{B}/2", f"{Hf}+7*{u}", pos="t", size=7)

    # ========================================================================= P3 ÉLÉVATION LONG-PAN
    d = Drawing(co, 3, "ÉLÉVATION LONG-PAN", "0", L, "0", f"{Hm}+{hacr}", ml=20, mr=10, mb=12, mt=10)
    D_.append(d)
    u = d.u
    co.ser(d, "Sol", seg(f"-5*{u}", "0", f"{L}+5*{u}", "0"), S(NOIR, 1.0))
    pts = []
    for i in range(MAXP):
        pts += seg(f"{er}*{i}", "0", f"{er}*{i}", Hm, f"{i}<={nt}")
    pts += seg("0", Hm, L, Hm) + [("0", Hm, f"{acr}=1"), ("0", f"{Hm}+{hacr}", f"{acr}=1"),
                                  (L, f"{Hm}+{hacr}", f"{acr}=1"), (L, Hm, f"{acr}=1"), None]
    co.ser(d, "Poteaux et sablière", pts, S(BLEU, 2.0))
    pts = []
    for k in range(MAXL):
        z = f"({z0}+{k}*{elis})"
        pts += split_row(z, "0", L, cuts("LP"), f"{k}<{nrows}")
    co.ser(d, "Lisses", pts, S(BLEU, 0.6))
    pts = []
    for k in range(1, MAXC_LP + 1):
        z = f"({k}*{pas_c})"
        pts += split_row(z, "0", L, cuts("LP"), AND(f"{kag}=1", f"{z}<{hag}"))
    co.ser(d, "Maçonnerie agglos", pts, S(MARRON, 0.5))
    pts = seg("0", f"{hch}/2", L, f"{hch}/2", f"{kag}=1") + \
        seg("0", f"{hag}/2", L, f"{hag}/2", f"{agg_full}=1") + \
        seg("0", f"{hag}-{hch}/2", L, f"{hag}-{hch}/2", f"{kag}=1")
    co.ser(d, "Chaînages", pts, S(GRIS, 2.0))
    pts = []
    for m in range(1, MAXCV + 1):
        b = f"IF({m}=1,1,IF({m}={nbcv},{nt},5*({m}-1)))"
        xa, xb = f"({b}-1)*{er}", f"{b}*{er}"
        c = f"{m}<={nbcv}"
        pts += seg(xa, "0", xb, Hm, c) + seg(xb, "0", xa, Hm, c)
    co.ser(d, "Croix de contreventement", pts, S(ORANGE, 1.25))
    pts = []
    for dd in doors["LP"]:
        c = f"{dd['ex']}=1"
        pts += seg(dd["xl"], "0", dd["xl"], Hm, c) + seg(dd["xr"], "0", dd["xr"], Hm, c) + \
            seg(dd["xl"], dd["h"], dd["xr"], dd["h"], c)
    co.ser(d, "Potelets et linteaux", pts, S(BLEU, 1.25))
    pts = []
    for dd in doors["LP"]:
        c = f"{dd['ex']}=1"
        pts += seg(dd["xl"], "0", dd["xr"], dd["h"], c) + seg(dd["xr"], "0", dd["xl"], dd["h"], c)
        pts += seg(f"{dd['xl']}-0.1", f"{dd['h']}+0.15", f"{dd['xr']}+0.1", f"{dd['h']}+0.15",
                   AND(c, dd["coul"]))
    co.ser(d, "Baies et larmiers", pts, S(NOIR, 0.75, "sysDash"))
    for dd in doors["LP"]:
        co.lab(d, dd["t"], dd["xc"], f"{dd['h']}+0.15", f"{dd['ex']}=1", pos="t", size=7)
    co.cote(d, "0", f"-6*{u}", L, f"-6*{u}", tL, True, pos="b")
    co.cote(d, f"-9*{u}", "0", f"-9*{u}", Hm, tHm, False)
    co.cote(d, "0", f"{Hm}+{hacr}+4*{u}", er, f"{Hm}+{hacr}+4*{u}", tE, True)

    # ========================================================================= P4 ÉLÉVATION PIGNON
    d = Drawing(co, 4, "ÉLÉVATION PIGNON", "0", B, "0", Hf, ml=20, mr=20, mb=12, mt=6)
    D_.append(d)
    u = d.u
    co.ser(d, "Sol", seg(f"-5*{u}", "0", f"{B}+5*{u}", "0"), S(NOIR, 1.0))
    pts = [("0", "0"), ("0", Hm), (f"{B}/2", Hf), (B, Hm), (B, "0"), None]
    for k in range(1, MAXIP):
        y = f"{k}*{eppr}"
        pts += seg(y, "0", y, f"{Hm}+{p}*MIN({y},{B}-{y})", f"{k}<={n_ip}-1")
    co.ser(d, "Poteaux de pignon et rampants", pts, S(BLEU, 2.0))
    xa = lambda z: f"IF({z}<={Hm},0,({z}-{Hm})/{p})"
    xb = lambda z: f"({B}-{xa(z)})"
    pts = []
    for k in range(MAXLP):
        z = f"({z0}+{k}*{elis})"
        pts += split_row(z, xa(z), xb(z), cuts("PG"), AND(f"{kbard}=1", f"{z}<{Hf}-0.3"), clamp=True)
    co.ser(d, "Lisses", pts, S(BLEU, 0.6))
    pts = []
    for k in range(1, MAXC_PG + 1):
        z = f"({k}*{pas_c})"
        ztop = f"IF({agg_full}=1,{Hf}-0.2,{hag})"
        pts += split_row(z, xa(z), xb(z), cuts("PG"), AND(f"{kag}=1", f"{z}<{ztop}"), clamp=True)
    co.ser(d, "Maçonnerie agglos", pts, S(MARRON, 0.5))
    pts = seg("0", f"{hch}/2", B, f"{hch}/2", f"{kag}=1") + \
        seg("0", f"{hag}/2", B, f"{hag}/2", f"{agg_full}=1") + \
        seg("0", f"{hag}-{hch}/2", B, f"{hag}-{hch}/2", f"{kag}=1") + \
        [("0", f"{Hm}-{hch}/2", f"{agg_full}=1"), (f"{B}/2", f"{Hf}-{hch}/2", f"{agg_full}=1"),
         (B, f"{Hm}-{hch}/2", f"{agg_full}=1"), None]
    co.ser(d, "Chaînages", pts, S(GRIS, 2.0))
    pts = []
    for dd in doors["PG"]:
        c = f"{dd['ex']}=1"
        pts += seg(dd["xl"], "0", dd["xl"], f"MIN({Hm},{dd['h']}+{elis})", c) + \
            seg(dd["xr"], "0", dd["xr"], f"MIN({Hm},{dd['h']}+{elis})", c) + seg(dd["xl"], dd["h"], dd["xr"], dd["h"], c)
    co.ser(d, "Potelets et linteaux", pts, S(BLEU, 1.25))
    pts = []
    for dd in doors["PG"]:
        c = f"{dd['ex']}=1"
        pts += seg(dd["xl"], "0", dd["xr"], dd["h"], c) + seg(dd["xr"], "0", dd["xl"], dd["h"], c)
        pts += seg(f"{dd['xl']}-0.1", f"{dd['h']}+0.15", f"{dd['xr']}+0.1", f"{dd['h']}+0.15", AND(c, dd["coul"]))
    co.ser(d, "Baies et larmiers", pts, S(NOIR, 0.75, "sysDash"))
    for dd in doors["PG"]:
        co.lab(d, dd["t"], dd["xc"], f"{dd['h']}+0.15", f"{dd['ex']}=1", pos="t", size=7)
    co.cote(d, "0", f"-6*{u}", B, f"-6*{u}", tB, True, pos="b")
    co.cote(d, f"-9*{u}", "0", f"-9*{u}", Hm, tHm, False)
    co.cote(d, f"{B}+9*{u}", "0", f"{B}+9*{u}", Hf, tHf, False, pos="r")
    co.lab(d, tPig, f"{B}/2", f"{Hf}+3*{u}", pos="t", size=7)

    # ========================================================================= P5 PLAN DE TOITURE
    d = Drawing(co, 5, "VUE EN PLAN TOITURE – CALEPINAGE", f"-{deb}", f"{L}+{deb}", "0", B, ml=20, mr=16, mb=14, mt=10)
    D_.append(d)
    u = d.u
    co.ser(d, "Rives et égouts", rect(f"-{deb}", "0", f"{L}+{deb}", B), S(NOIR, 1.25))
    pts = []
    for j in range(MAXPN):
        y = f"{j}*({B}/2)/MAX(1,{npn}-1)"
        c = f"{j}<={npn}-1"
        pts += seg(f"-{deb}", y, f"{L}+{deb}", y, c) + seg(f"-{deb}", f"{B}-{y}", f"{L}+{deb}", f"{B}-{y}", c)
    co.ser(d, "Pannes", pts, S(BLEU, 0.5))
    pts = []
    for i in range(MAXP):
        pts += seg(f"{er}*{i}", "0", f"{er}*{i}", B, f"{i}<={nt}")
    co.ser(d, "Fermes / traverses", pts, S(BLEU, 1.75))
    pts = []
    for i in range(1, MAXP):
        for l in (1, 2):
            x = f"({i}-1)*{er}+{l}*{er}/({nl}+1)"
            pts += seg(x, "0", x, B, AND(f"{i}<={nt}", f"{l}<={nl}"))
    co.ser(d, "Liernes", pts, S(BLEU, 0.5, "sysDot"))
    co.ser(d, "Faîtière", seg(f"-{deb}", f"{B}/2", f"{L}+{deb}", f"{B}/2"), S(NOIR, 2.25, "dash"))
    pts = []
    for m in range(1, MAXCV + 1):
        b = f"IF({m}=1,1,IF({m}={nbcv},{nt},5*({m}-1)))"
        xa_, xb_ = f"({b}-1)*{er}", f"{b}*{er}"
        for v in (0, 1):
            for r in range(1, MAXPV + 1):
                ya = f"({v}*{B}/2+({r}-1)*({B}/2)/{npv})"
                yb = f"({v}*{B}/2+{r}*({B}/2)/{npv})"
                c = AND(f"{m}<={nbcv}", f"{r}<={npv}")
                pts += seg(xa_, ya, xb_, yb, c) + seg(xa_, yb, xb_, ya, c)
    co.ser(d, "Poutre au vent", pts, S(ORANGE, 1.0))
    pts = []
    for i in range(1, MAXP):
        xc = f"(({i}-0.5)*{er})"
        w = f"({pct}*{er})"
        xl_, xr_ = f"{xc}-{w}/2", f"{xc}+{w}/2"
        for v in (0, 1):
            ya, yb = f"({v}*{B}/2+0.6*{u})", f"({v}*{B}/2+{B}/2-0.6*{u})"
            c = AND(f"{i}<={nt}", f"{pct}>0")
            pts += rect(xl_, ya, xr_, yb, c)
            hh = f"(({yb})-({ya}))"
            pts += [(xl_, ya, c), (xr_, f"{ya}+{hh}/4", c), (xl_, f"{ya}+{hh}/2", c), (xr_, f"{ya}+3*{hh}/4", c),
                    (xl_, yb, c), None]
    co.ser(d, "Translucides (hachurés)", pts, S(TRANSL, 0.75))
    pts = []
    for k in range(1, MAXSKY + 1):
        b = f"MAX(1,MIN({nt},ROUND({k}*{nt}/({nsky}+1),0)))"
        xc = f"(({b}-1)*{er}+0.25*{er})"
        yc = f"IF(MOD({k},2)=1,{B}/4,3*{B}/4)"
        c = f"{k}<={nsky}"
        pts += rect(f"{xc}-{asky}/2", f"{yc}-{bsky}/2", f"{xc}+{asky}/2", f"{yc}+{bsky}/2", c)
        pts += rect(f"{xc}-{asky}/2-0.2", f"{yc}-{bsky}/2-0.2", f"{xc}+{asky}/2+0.2", f"{yc}+{bsky}/2+0.2", c)
    co.ser(d, "Skydomes et chevêtres", pts, S(BLEU, 1.25))
    pts = []
    for k in range(1, MAXEP + 1):
        x = f"({k}-0.5)*{L}/MAX(1,{ndesc}/2)"
        c = f"{k}<={ndesc}/2"
        pts += [(x, "0", c), (x, B, c)]
    co.ser(d, "Descentes EP", pts, S("002060", line=False, marker="circle", msize=6))
    co.cote(d, "0", f"-7*{u}", L, f"-7*{u}", tL, True, pos="b")
    co.cote(d, f"-{deb}-9*{u}", "0", f"-{deb}-9*{u}", B, tB, False)
    co.lab(d, "Faîtage", f"{L}+{deb}+2*{u}", f"{B}/2", pos="r", size=7)
    co.lab(d, tP, f"{L}+{deb}+2*{u}", f"{B}/4", pos="r", size=7)
    co.lab(d, "● descentes EP", f"{L}/2", f"-2.5*{u}", pos="b", size=7)

    # ========================================================================= P6 DÉTAIL FONDATION
    zf = sc("p6_zf", f"-{D}")                      # fond de fouille
    zp = sc("p6_zp", f"{zf}+{eprop}")              # dessus béton de propreté
    zsf = sc("p6_zsf", f"{zp}+{hfil}")             # dessus semelle filante
    zag = sc("p6_zag", f"{zsf}+{hmf}")             # dessus maçonnerie de fondation
    zr = sc("p6_zr", erem)                          # dessus remblai
    zs = sc("p6_zs", f"{zr}+{esab}")                # dessus lit de sable (polyane)
    zd = sc("p6_zd", f"{zs}+{edal}")                # dessus dallage
    zsi = sc("p6_zsi", f"{zp}+{hs1}")               # dessus semelle isolée
    zft = sc("p6_zft", f"{zsi}+{hfut}")             # dessus fût
    xB = sc("p6_xB", f"{bfil}/2+1.6+{a1}/2")        # axe de la semelle isolée
    xE = sc("p6_xE", f"{xB}+{a1}/2+0.8")            # bord droit du détail
    d = Drawing(co, 6, "COUPE DE FONDATION ET DALLAGE (DÉTAIL)", f"-{bfil}/2-0.3", xE, f"{zf}-0.1",
                f"MAX({zd},{zft})+0.9", ml=16, mr=62, mb=10, mt=6)
    D_.append(d)
    u = d.u
    co.ser(d, "Terrain naturel", seg(f"-{bfil}/2-0.3", "0", "-0.075", "0"), S(NOIR, 0.6, "lgDashDot"))
    pts = rect(f"-{bfil}/2-0.05", zf, f"{bfil}/2+0.05", zp) + rect(f"-{bfil}/2", zp, f"{bfil}/2", zsf) + \
        rect(f"{xB}-{a1}/2-0.05", zf, f"{xB}+{a1}/2+0.05", zp) + rect(f"{xB}-{a1}/2", zp, f"{xB}+{a1}/2", zsi) + \
        rect(f"{xB}-{f1}/2", zsi, f"{xB}+{f1}/2", zft) + \
        [("0.075", zs), (f"{xB}-{f1}/2", zs), None, ("0.075", zd), (f"{xB}-{f1}/2", zd), None,
         (f"{xB}+{f1}/2", zs), (xE, zs), None, (f"{xB}+{f1}/2", zd), (xE, zd), None, ("0.075", zs), ("0.075", zd), None]
    co.ser(d, "Béton", pts, S(GRIS, 1.25))
    pts = rect("-0.075", zsf, "0.075", zag)
    for k in range(1, 8):
        z = f"{zsf}+{k}*{pas_c7}"
        pts += seg("-0.075", z, "0.075", z, f"{z}<{zag}")
    pts += [("-0.075", zag), ("-0.075", f"{zd}+0.6"), None, ("0.075", zag), ("0.075", f"{zd}+0.6"), None]
    co.ser(d, "Maçonnerie", pts, S(MARRON, 1.25))
    pts = seg("0.075", zr, f"{xB}-{f1}/2", zr) + seg(f"{xB}+{f1}/2", zr, xE, zr)
    for k in range(12):
        x = f"0.075+{k}*({xE}-0.075)/12"
        x2 = f"0.075+({k}+0.5)*({xE}-0.075)/12"
        pts += [(x, "0"), (x2, zr), None]
    co.ser(d, "Remblai latérite", pts, S(LATERITE, 0.6))
    co.ser(d, "Polyane", seg("0.075", f"{zs}+0.005", xE, f"{zs}+0.005"), S(NOIR, 0.6, "dash"))
    co.ser(d, "Lit de sable", seg("0.075", f"{zr}+{esab}/2", xE, f"{zr}+{esab}/2"), S("BF9000", 0.5, "sysDot"))
    # aciers
    zb = f"{zp}+{cf}"
    pts = seg(f"-{bfil}/2+{cf}", zb, f"{bfil}/2-{cf}", zb) + \
        [(f"{xB}-{a1}/2+{cf}", f"{zb}+0.15"), (f"{xB}-{a1}/2+{cf}", zb), (f"{xB}+{a1}/2-{cf}", zb),
         (f"{xB}+{a1}/2-{cf}", f"{zb}+0.15"), None]
    for sgn in ("-", "+"):
        xbar = f"{xB}{sgn}({f1}/2-{cf})"
        pts += [(f"{xbar}{sgn}0.15", f"{zb}+0.02"), (xbar, f"{zb}+0.02"), (xbar, f"{zft}-{cf}"), None]
    for k in range(12):
        z = f"{zsi}+0.05+{k}*{{e_cad_fut}}"
        pts += seg(f"{xB}-{f1}/2+{cf}", z, f"{xB}+{f1}/2-{cf}", z, f"{z}<{zft}-{cf}")
    pts += seg(f"0.075+{cd}", f"{zs}+{cd}", f"{xB}-{f1}/2-{cd}", f"{zs}+{cd}") + \
        seg(f"{xB}+{f1}/2+{cd}", f"{zs}+{cd}", xE, f"{zs}+{cd}") + \
        seg(f"0.075+{cd}", f"{zd}-{cd}", f"{xB}-{f1}/2-{cd}", f"{zd}-{cd}", f"{kdbl}=1") + \
        seg(f"{xB}+{f1}/2+{cd}", f"{zd}-{cd}", xE, f"{zd}-{cd}", f"{kdbl}=1")
    for sgn in ("-", "+"):
        xt = f"{xB}{sgn}({f1}/2-0.10)"
        pts += seg(xt, f"{zft}+0.06", xt, f"{zft}-{ltig}+0.06")
    co.ser(d, "Aciers", pts, S(ROUGE, 1.0))
    pts = []
    for k in range(4):
        pts.append((f"-{bfil}/2+{cf}+{k}*({bfil}-2*{cf})/3", f"{zb}+0.02"))
    nq = f"ROUNDDOWN(({a1}-2*{cf})/{{m_sem}},0)"
    for k in range(16):
        pts.append((f"{xB}-{a1}/2+{cf}+{k}*{{m_sem}}", f"{zb}+0.02", f"{k}<={nq}"))
    for k in range(40):
        x = f"0.2+{k}*{mai}"
        pts.append((x, f"{zs}+{cd}+0.015", AND(f"{x}<{xE}", f"ABS({x}-{xB})>{f1}/2+0.03")))
        pts.append((x, f"{zd}-{cd}-0.015", AND(f"{x}<{xE}", f"ABS({x}-{xB})>{f1}/2+0.03", f"{kdbl}=1")))
    co.ser(d, "Aciers (sections)", pts, S(ROUGE, line=False, marker="circle", msize=3))
    pts = seg(f"{xB}-{{L_pl_pied_port}}/2", f"{zft}+0.025", f"{xB}+{{L_pl_pied_port}}/2", f"{zft}+0.025") + \
        seg(f"{xB}-{hpot}/2", f"{zft}+0.05", f"{xB}-{hpot}/2", f"{zft}+0.85") + \
        seg(f"{xB}+{hpot}/2", f"{zft}+0.05", f"{xB}+{hpot}/2", f"{zft}+0.85")
    co.ser(d, "Platine et poteau", pts, S(BLEU, 2.0))
    co.cote(d, f"-{bfil}/2-0.18", zf, f"-{bfil}/2-0.18", "0", tx("p6_tD", '"Prof. "&FIXED({D},2,TRUE)&" m"'), False)
    co.cote(d, f"-{bfil}/2", f"{zf}-0.06", f"{bfil}/2", f"{zf}-0.06",
            tx("p6_tf", '"Semelle filante "&{b_fil}*100&" × "&{h_fil}*100'), True, pos="b")
    co.cote(d, f"{xB}-{a1}/2", f"{zf}-0.06", f"{xB}+{a1}/2", f"{zf}-0.06",
            tx("p6_ts", '"Semelle "&{a1}*100&" × "&{a1}*100&" × "&{h_s1}*100'), True, pos="b")
    lx = f"{xE}+5*{u}"
    ytop = f"(MAX({zd},{zft})+0.75)"
    reps = [(f"{zft}+0.06", tx("p6_t6", '"Tiges d\'ancrage "&{d_tige_port}&" L = "&FIXED({l_tige_port},2,TRUE)&" m + platine"'), f"{xB}+{f1}/2-0.10"),
            (f"({zs}+{zd})/2", tx("p6_t1", '"Dallage "&{e_dal}*100&" cm – "&{ferr}&" – "&{ha_ix}&"/"&{ha_iy}'
                                           '&IF({kdbl}=1," + "&{ha_sx}&"/"&{ha_sy},"")&" maille "&{maille}*100&" cm"'), xE),
            (f"{zs}+0.005", tx("p6_t2", '"Polyane sur lit de sable "&{e_sable}*100&" cm"'), xE),
            (f"{zr}/2", tx("p6_t3", '"Remblai latérite compactée 95 % OPM – "&{e_rem}*100&" cm"'), xE),
            ("0", "TN (terrain naturel)", "-0.075"),
            (f"({zsi}+{zft})/2", tx("p6_t4", '"Fût "&{f1}*100&" × "&{f1}*100&" h "&FIXED({h_fut},2,TRUE)&" m – "'
                                            '&{nb_bfut}&" "&{ha_fut}&" + cadres "&{ha_cad}&" e = "&{e_cad_fut}*100'), f"{xB}+{f1}/2"),
            (f"{zp}+{cf}", tx("p6_t5", '"Quadrillage "&{ha_sem}&" maille "&{m_sem}*100&" cm"'), f"{xB}+{a1}/2-{cf}"),
            (f"({zp}+{zsi})/2", tx("p6_t9", '"Semelle isolée béton armé "&{dos_ba}&" kg/m³ sur propreté "&{dos_prop}&" kg/m³"'), f"{xB}+{a1}/2")]
    pts = []
    for k, (z, t, xa_) in enumerate(reps):
        yl = f"{ytop}-{k}*6.5*{u}"
        pts += [(xa_, z), (f"{xE}+1*{u}", z), (f"{xE}+4*{u}", yl), None]
        co.lab(d, t, lx, yl, pos="r", size=7)
    co.ser(d, "Repères", pts, S(NOIR, 0.4))
    co.lab(d, tx("p6_t7", '"Agglos 15 pleins h "&FIXED({h_mf},2,TRUE)&" m"'), "-0.075", f"({zsf}+{zag})/2",
           pos="l", size=7, rot=True)
    co.lab(d, tx("p6_t8", '"Filants "&{nb_fil}&" "&{ha_fil}&" + rép. "&{ha_rep}&" e = "&{e_rep}*100'),
           "0", f"{zf}-0.30", pos="b", size=7)

    # ========================================================================= P7 MUR AGGLOS
    E = sc("p7_e", er)
    Hw = sc("p7_h", f"MAX(0.5,{hag})")
    k7 = f"{kag}=1"
    d = Drawing(co, 7, "ÉLÉVATION MUR AGGLOS – PANNEAU TYPE", "0", E, "0", f"{Hw}+0.5", ml=16, mr=58, mb=12, mt=8)
    D_.append(d)
    u = d.u
    co.ser(d, "Sol", seg(f"-3*{u}", "0", f"{E}+3*{u}", "0", k7), S(NOIR, 1.0))
    co.ser(d, "Poteaux métalliques", seg("0", "0", "0", f"{Hw}+0.4", k7) + seg(E, "0", E, f"{Hw}+0.4", k7),
           S(BLEU, 3.0))
    zint = f"{Hw}/2"
    pts = rect("0", "0", E, hch, k7) + rect("0", f"{Hw}-{hch}", E, Hw, k7) + \
        rect("0", f"{zint}-{hch}/2", E, f"{zint}+{hch}/2", AND(k7, f"{agg_full}=1"))
    for j in (1, 2, 3):
        xc = f"{E}*{j}/4"
        pts += rect(f"{xc}-{brd}/2", hch, f"{xc}+{brd}/2", f"{Hw}-{hch}", k7)
    co.ser(d, "Béton (chaînages, raidisseurs)", pts, S(GRIS, 1.25))
    pts = []
    for k in range(1, 61):
        z = f"{k}*{pas_c7}"
        c = AND(k7, f"{z}>{hch}", f"{z}<{Hw}-{hch}",
                f"OR({agg_full}=0,ABS({z}-{zint})>{hch}/2)")
        for j in range(4):
            x1 = "0" if j == 0 else f"{E}*{j}/4+{brd}/2"
            x2 = E if j == 3 else f"{E}*{j + 1}/4-{brd}/2"
            pts += seg(x1, z, x2, z, c)
    co.ser(d, "Agglos 15 creux", pts, S(MARRON, 0.5))
    pts = []
    for zc, c in [(f"{hch}/2", k7), (f"{Hw}-{hch}/2", k7), (zint, AND(k7, f"{agg_full}=1"))]:
        for sg in ("-", "+"):
            z = f"{zc}{sg}({hch}/2-{cf}*0.6)"
            pts += seg("0", z, E, z, c)
    for j in (1, 2, 3):
        for sg in ("-", "+"):
            x = f"{E}*{j}/4{sg}({brd}/2-{cf}*0.5)"
            pts += seg(x, f"{cf}", x, f"{Hw}-{cf}", k7)
    co.ser(d, "Aciers", pts, S(ROUGE, 0.75))
    co.cote(d, "0", f"-6*{u}", E, f"-6*{u}", tx("p7_tE", '"Panneau entre poteaux "&FIXED({er},2,TRUE)&" m"'), True,
            cond=k7, pos="b")
    co.cote(d, f"-6*{u}", "0", f"-6*{u}", Hw, tx("p7_tH", '"Mur h "&FIXED({h_ag},2,TRUE)&" m"'), False, cond=k7)
    lx = f"{E}+4*{u}"
    co.lab(d, tx("p7_t1", '"Chaînage haut "&{b_ch}*100&" × "&{h_ch}*100&" – "&{nb_chl}&" "&{ha_chl}&" + cad. "&{ha_chc}&" e = "&{e_chc}*100'),
           lx, f"{Hw}-{hch}/2", k7, pos="r", size=7)
    co.lab(d, tx("p7_t2", '"Chaînage intermédiaire "&{b_ch}*100&" × "&{h_ch}*100&" – "&{nb_chl}&" "&{ha_chl}&" + cad. "&{ha_chc}'),
           lx, zint, AND(k7, f"{agg_full}=1"), pos="r", size=7)
    co.lab(d, tx("p7_t3", '"Chaînage bas "&{b_ch}*100&" × "&{h_ch}*100&" – "&{nb_chl}&" "&{ha_chl}&" + cad. "&{ha_chc}'),
           lx, f"{hch}/2", k7, pos="r", size=7)
    co.lab(d, tx("p7_t4", '{nb_rd}&" raidisseurs "&{b_rd}*100&" × "&{b_rd}*100&" – "&{nb_chl}&" "&{ha_chl}&" + cad. "&{ha_chc}&" e = "&{e_chc}*100'),
           f"{E}/2", f"{Hw}+3*{u}", k7, pos="t", size=7)
    co.lab(d, "SANS OBJET : façade en bardage bac acier (vue affichée en « Agglos 15 creux » ou « Mixte »)",
           f"{E}/2", f"{Hw}/2", f"{kag}=0", pos="ctr", size=11, bold=True, color=ROUGE)

    # --------------------------------------------------------------------- feuille PLANS
    wsc.sheet_state = "hidden"
    for col, w in zip("ABCDE", [16, 12, 3, 12, 60]):
        wsc.column_dimensions[col].width = w
    protect(wsc)
    build_plans_sheet(wb, M, names, D_, st, page_setup, protect)


# ============================================================================= GRAPHIQUES
def rich(size, color=NOIR, rot=False, bold=False):
    cp = CharacterProperties(sz=int(size * 100), b=bold, solidFill=color)
    bp = RichTextProperties(rot=-5400000 if rot else 0, vert="horz", anchor="ctr", wrap="none")
    return RichText(bodyPr=bp, p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def make_chart(wsc, d):
    ch = ScatterChart()
    ch.scatterStyle = "lineMarker"
    ch.style = None
    ch.legend = None
    ch.title = None
    for ax, mx in ((ch.x_axis, XMAX), (ch.y_axis, YMAX)):
        ax.scaling.min, ax.scaling.max = 0, mx
        ax.delete = True
        ax.majorGridlines = None
        ax.minorGridlines = None
    ch.width, ch.height = CH_W, CH_H
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                           x=PL_X, y=PL_Y, w=PL_W, h=PL_H))
    ch.display_blanks = "gap"
    ch.visible_cells_only = False
    for sd in d.series:
        xr = Reference(wsc, min_col=sd["col"], min_row=sd["r0"], max_row=sd["r1"])
        yr = Reference(wsc, min_col=sd["col"] + 1, min_row=sd["r0"], max_row=sd["r1"])
        s = Series(yr, xr)
        lab = sd["label"]
        if lab and lab["text"].startswith("'"):
            s.tx = SeriesLabel(strRef=StrRef(f=lab["text"]))
        elif lab:
            s.tx = SeriesLabel(v=lab["text"])
        else:
            s.tx = SeriesLabel(v=sd["name"])
        s.smooth = False
        sty = sd["style"]
        gp = GraphicalProperties()
        if sty["line"]:
            gp.line.solidFill = sty["color"]
            gp.line.width = int(sty["w"] * 12700)
            if sty["dash"]:
                gp.line.prstDash = sty["dash"]
        else:
            gp.line.noFill = True
        s.graphicalProperties = gp
        if sty["marker"]:
            s.marker = Marker(symbol=sty["marker"], size=sty["msize"])
            mgp = GraphicalProperties(solidFill=sty["color"])
            mgp.line.solidFill = sty["color"]
            s.marker.graphicalProperties = mgp
        else:
            s.marker = Marker(symbol="none")
        if lab:
            dl = DataLabel(idx=lab["idx"], showSerName=True, showVal=False, showCatName=False, showLegendKey=False,
                           showPercent=False, showBubbleSize=False, dLblPos=lab["pos"],
                           txPr=rich(lab["size"], lab["color"], lab["rot"], lab["bold"]))
            s.dLbls = DataLabelList(dLbl=[dl], showSerName=False, showVal=False, showCatName=False,
                                    showLegendKey=False, showPercent=False, showBubbleSize=False)
        ch.series.append(s)
    return ch


def build_plans_sheet(wb, M, names, drawings, st, page_setup, protect):
    PL, CO = names["PL"], names["CO"]
    ws = wb.create_sheet(PL)
    wsc = wb[CO]
    ws.column_dimensions["A"].width = 1.5
    for i in range(2, 14):
        ws.column_dimensions[get_column_letter(i)].width = 11.3
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    BLOCK = 36
    legend = ("Couleurs : charpente en bleu · contreventements en orange · béton en gris · aciers en rouge · "
              "maçonnerie en marron clair · cotes et baies en noir · translucides en bleu clair · remblai en latérite")
    gab = ('"Gabarit "&{gab}&" : portée "&FIXED({portee},2,TRUE)&" m × longueur "&FIXED({longueur},2,TRUE)'
           '&" m – H sablière "&FIXED({hauteur},2,TRUE)&" m – pente "&FIXED({pente}*100,1,TRUE)&" % – "&{systeme}&" – façade : "&{facade}')
    for k, d in enumerate(drawings):
        r0 = 1 + k * BLOCK
        c = ws.cell(r0, 2, f"PLANCHE {k + 1} / {len(drawings)} – {d.title}")
        st(c, bold=True, color="FFFFFF", fill="1F4E79", size=12, border=False)
        for col in range(3, 14):
            ws.cell(r0, col).fill = copy(c.fill)
        ch = make_chart(wsc, d)
        ws.add_chart(ch, f"B{r0 + 1}")
        rc = r0 + 27
        # cartouche
        ws.merge_cells(start_row=rc, start_column=2, end_row=rc, end_column=9)
        st(ws.cell(rc, 2, d.title), bold=True, size=11)
        ws.merge_cells(start_row=rc + 1, start_column=2, end_row=rc + 1, end_column=9)
        st(ws.cell(rc + 1, 2), size=8, wrap=True)
        ws.row_dimensions[rc + 1].height = 24
        M.disp(PL, f"B{rc + 1}", gab)
        ws.merge_cells(start_row=rc + 2, start_column=2, end_row=rc + 2, end_column=9)
        st(ws.cell(rc + 2, 2, legend), size=7, italic=True, wrap=True)
        ws.row_dimensions[rc + 2].height = 22
        ws.merge_cells(start_row=rc + 3, start_column=2, end_row=rc + 4, end_column=9)
        st(ws.cell(rc + 3, 2), size=8, wrap=True)
        if k == 5:
            M.disp(PL, f"B{rc + 3}", '"Légende aciers : semelles isolées "&{ha_sem}&" maille "&{m_sem}*100&" cm ; fûts "'
                   '&{nb_bfut}&" "&{ha_fut}&" + cadres "&{ha_cad}&" e = "&{e_cad_fut}*100&" cm ; semelle filante "&{nb_fil}&" "'
                   '&{ha_fil}&" + "&{ha_rep}&" e = "&{e_rep}*100&" cm ; dallage "&{ferr}&" "&{ha_ix}&" (X) / "&{ha_iy}&" (Y)"'
                   '&IF({kdbl}=1," + nappe sup. "&{ha_sx}&" / "&{ha_sy}&" + chaises "&{ha_chaise},"")&", maille "&{maille}*100&" cm."')
        elif k == 6:
            M.disp(PL, f"B{rc + 3}", '"Sections : chaînages "&{b_ch}*100&" × "&{h_ch}*100&" cm et raidisseurs "&{b_rd}*100'
                   '&" × "&{b_rd}*100&" cm – "&{nb_chl}&" "&{ha_chl}&" + cadres "&{ha_chc}&" e = "&{e_chc}*100'
                   '&" cm ; linteaux idem avec appuis de "&{appui_lin}*100&" cm."')
        elif k == 0:
            M.disp(PL, f"B{rc + 3}", '"Poteaux de portique ■ "&{prof_pot}&" sur semelles "&{a1}*100&" × "&{a1}*100'
                   '&" ; poteaux de pignon ● "&{prof_pp}&" sur semelles "&{a2}*100&" × "&{a2}*100&" ; semelles filantes en pointillés."')
        elif k == 4:
            M.disp(PL, f"B{rc + 3}", '"Translucides : "&FIXED({pct_transl}*100,0,TRUE)&" % – skydomes : "&{nb_sky}'
                   '&" – descentes EP : "&{nb_desc}&" – poutre au vent : "&{nb_cv}&" travées – liernes : "&{n_l}&" ligne(s) par travée."')
        else:
            M.disp(PL, f"B{rc + 3}", '"Croix de Saint-André dans les travées d\'extrémité et une travée sur 5 ("&{nb_cv}'
                   '&" travées). "&IF({kbard}=1,"Lisses "&{prof_lisse}&" e = "&FIXED({e_lis},2,TRUE)&" m","Façade maçonnée : lisses sans objet")'
                   '&IF({kag}=1," ; maçonnerie agglos 15 creux sur "&FIXED({z0},2,TRUE)&" m avec chaînages gris.",".")')
        st(ws.cell(rc, 10, f"Planche n° {k + 1} / {len(drawings)}"), bold=True, h="center")
        ws.merge_cells(start_row=rc, start_column=10, end_row=rc, end_column=11)
        st(ws.cell(rc, 12, "Date :"), h="right", size=9)
        c = ws.cell(rc, 13, "=TODAY()")
        st(c, fmt="dd/mm/yyyy", size=9, h="center")
        ws.merge_cells(start_row=rc + 1, start_column=10, end_row=rc + 1, end_column=13)
        st(ws.cell(rc + 1, 10, "ESQUISSE NON CONTRACTUELLE"), bold=True, color="C00000", h="center")
        ws.merge_cells(start_row=rc + 2, start_column=10, end_row=rc + 2, end_column=13)
        st(ws.cell(rc + 2, 10, "BON POUR ACCORD"), bold=True, h="center", fill="FBE5D6")
        st(ws.cell(rc + 3, 10, "Date :"), size=9)
        st(ws.cell(rc + 4, 10, "Signature :"), size=9)
        for rr in range(rc, rc + 6):
            for cc in range(2, 14):
                ws.cell(rr, cc).border = box
        ws.merge_cells(start_row=rc + 3, start_column=11, end_row=rc + 3, end_column=13)
        ws.merge_cells(start_row=rc + 4, start_column=11, end_row=rc + 5, end_column=13)
        ws.merge_cells(start_row=rc + 5, start_column=2, end_row=rc + 5, end_column=9)
        st(ws.cell(rc + 5, 2, "Échelle : automatique (même échelle en X et en Y) – cotes en mètres"), size=7, italic=True)
        if k < len(drawings) - 1:
            ws.row_breaks.append(Break(id=r0 + BLOCK - 1))
    page_setup(ws, rows_title=None)
    ws.print_area = f"A1:M{len(drawings) * BLOCK}"
    ws.page_setup.fitToHeight = len(drawings)
    ws.sheet_view.showGridLines = False
