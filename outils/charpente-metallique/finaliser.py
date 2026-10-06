# -*- coding: utf-8 -*-
"""Finalisation du classeur généré par openpyxl :
1) graphiques : titre automatique supprimé et option « afficher #N/A comme cellule vide »
   (ruptures de traits des esquisses) ;
2) valeurs en cache : le classeur est recalculé par LibreOffice sur une copie, puis les
   résultats sont réinjectés dans les cellules formules du fichier d'origine (qui garde
   ses graphiques, validations, protections). Excel recalcule de toute façon à l'ouverture.
"""
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

NA_EXT = ('<extLst><ext uri="{56B9EC1D-385E-4148-901F-78D8002777C0}" '
          'xmlns:c16r3="http://schemas.microsoft.com/office/drawing/2017/03/chart">'
          '<c16r3:dataDisplayOptions16><c16r3:dispNaAsBlank val="1"/></c16r3:dataDisplayOptions16>'
          '</ext></extLst>')

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculerEtEnregistrer()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def _rewrite_zip(path, transform):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            new = transform(item.filename, data)
            zout.writestr(item, new if new is not None else data)
    os.replace(tmp, path)


def patch_charts(path):
    def tr(name, data):
        if not re.match(r"xl/charts/chart\d+\.xml$", name):
            return None
        s = data.decode("utf-8")
        if "<autoTitleDeleted" not in s:
            s = s.replace("<chart>", '<chart><autoTitleDeleted val="1"/>', 1)
        if "dispNaAsBlank" not in s:
            s = s.replace("</chart>", NA_EXT + "</chart>", 1)
        return s.encode("utf-8")
    _rewrite_zip(path, tr)


def recalc_copy(src, timeout=300):
    """Recalcule une copie avec LibreOffice ; renvoie son chemin, ou None si indisponible."""
    if not shutil.which("soffice"):
        return None
    work = Path(tempfile.mkdtemp(prefix="recalc_"))
    dst = work / "copie.xlsx"
    shutil.copy(src, dst)
    prof = work / "profil"
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    url = prof.as_uri()
    subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                   capture_output=True, timeout=120, env=env)
    mdir = prof / "user" / "basic" / "Standard"
    if not mdir.exists():
        return None
    (mdir / "Module1.xba").write_text(MACRO)
    before = dst.stat().st_mtime_ns
    subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
                    "vnd.sun.star.script:Standard.Module1.RecalculerEtEnregistrer?language=Basic&location=application",
                    str(dst)], capture_output=True, timeout=timeout, env=env)
    if dst.stat().st_mtime_ns == before:
        return None
    return str(dst)


def inject_values(path, recalculated):
    from openpyxl import load_workbook
    wb = load_workbook(recalculated, data_only=True)
    with zipfile.ZipFile(path) as z:
        wbxml = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rid_target.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    sheets = {}
    for name, rid in re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbxml):
        name = (name.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
                .replace("&quot;", '"').replace("&apos;", "'"))
        name = re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), name)
        t = rid_target[rid].lstrip("/")
        sheets["xl/" + t if not t.startswith("xl/") else t] = name
    cell_re = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)

    def tr(fname, data):
        if fname not in sheets:
            return None
        ws = wb[sheets[fname]]

        def rep(m):
            ref, attrs, f = m.group(1), m.group(2), m.group(3)
            attrs = re.sub(r'\s*t="[^"]*"', "", attrs)
            v = ws[ref].value
            if v is None:
                return f'<c r="{ref}"{attrs} t="str"><f>{f}</f><v></v></c>'
            if isinstance(v, bool):
                return f'<c r="{ref}"{attrs} t="b"><f>{f}</f><v>{int(v)}</v></c>'
            if isinstance(v, (int, float)):
                return f'<c r="{ref}"{attrs}><f>{f}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
            if hasattr(v, "toordinal"):  # date
                from openpyxl.utils.datetime import to_excel
                return f'<c r="{ref}"{attrs}><f>{f}</f><v>{to_excel(v)}</v></c>'
            s = str(v)
            if s.startswith("#") and s.rstrip("!?/0").upper() in ("#N/A", "#VALUE", "#REF", "#DIV", "#NUM", "#NAME", "#NULL"):
                return f'<c r="{ref}"{attrs} t="e"><f>{f}</f><v>{escape(s)}</v></c>'
            return f'<c r="{ref}"{attrs} t="str"><f>{f}</f><v>{escape(s)}</v></c>'
        return cell_re.sub(rep, data.decode("utf-8")).encode("utf-8")
    _rewrite_zip(path, tr)


def finaliser(path, recalc=True):
    patch_charts(path)
    if not recalc:
        return None
    rc = recalc_copy(path)
    if rc is None:
        print("LibreOffice indisponible : valeurs non injectées (Excel recalculera à l'ouverture).")
        return None
    inject_values(path, rc)
    return rc
