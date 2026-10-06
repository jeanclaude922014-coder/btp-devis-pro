#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère le classeur Excel de quantitatif, d'estimation et d'esquisses
de bâtiments industriels en charpente métallique (gros œuvre compris).

Usage : python3 generer_classeur.py [chemin_sortie.xlsx]

Normes de référence : EC3 (charpente), EC1 (charges, vent – pas de neige),
BAEL 91 révisé 99 et EC2 (béton armé, fondations). Contexte : Côte d'Ivoire.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

from moteur import Model, absref
import plans
import finaliser

# ----------------------------------------------------------------------------- noms
ACC, PAR, BP, BA = "ACCUEIL", "PARAMÈTRES", "BASE_PROFILÉS", "BASE_ACIERS"
PRE, QC, GO, DQE = "PRÉDIMENSIONNEMENT", "QUANTITATIF_CHARPENTE", "GROS_ŒUVRE", "DQE_DEVIS"
REC, PL, CO, CA = "RÉCAP_GABARITS", "PLANS", "COORD_PLANS", "CALC_GABARITS"

GABARITS = [("10 × 30", 10, 30), ("15 × 20", 15, 20), ("15 × 30", 15, 30), ("20 × 20", 20, 20),
            ("20 × 40", 20, 40), ("30 × 40", 30, 40), ("40 × 40", 40, 40)]
PERSO = "PERSONNALISÉ"
SCENARIOS = ["Façade du paramétrage", "Bardage bac acier", "Agglos 15 creux"]
FACADES = ["Bardage bac acier", "Agglos 15 creux", "Mixte"]
DIAMS = ["HA6", "HA8", "HA10", "HA12", "HA14", "HA16"]

# ----------------------------------------------------------------------------- styles
FONT = "Arial"
C_BLEU, C_BLEU_CLAIR, C_ORANGE, C_ORANGE_CLAIR = "1F4E79", "DDEBF7", "C55A11", "FBE5D6"
C_JAUNE, C_ROUGE, C_GRIS = "FFFF00", "FF0000", "F2F2F2"
NF_INT, NF_2, NF_PCT = "#,##0", "#,##0.00", "0.0%"
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL = lambda c: PatternFill("solid", start_color=c, end_color=c)
RED_FILL, WHITE_BOLD = FILL(C_ROUGE), Font(name=FONT, bold=True, color="FFFFFF")


def st(c, bold=False, fill=None, color="000000", size=10, fmt=None, h=None, wrap=False, border=True,
       italic=False):
    c.font = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
    if fill:
        c.fill = FILL(fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if border:
        c.border = BORDER


def yellow(c, fmt=None):
    st(c, fill=C_JAUNE, fmt=fmt, bold=True, color="000000")
    c.protection = Protection(locked=False)


def title(ws, text, sub=None, width_cols=10):
    ws["A1"] = text
    st(ws["A1"], bold=True, size=14, color="FFFFFF", fill=C_BLEU, border=False)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    ws.row_dimensions[1].height = 24
    if sub:
        ws["A2"] = sub
        st(ws["A2"], italic=True, size=9, color=C_ORANGE, border=False)


def header(ws, row, labels, col0=2, fill=C_BLEU):
    for i, t in enumerate(labels):
        c = ws.cell(row, col0 + i, t)
        st(c, bold=True, color="FFFFFF", fill=fill, h="center", wrap=True)
    ws.row_dimensions[row].height = 30


def section(ws, row, text, col0=2, ncols=9):
    c = ws.cell(row, col0, text)
    st(c, bold=True, color="FFFFFF", fill=C_ORANGE)
    for i in range(1, ncols):
        st(ws.cell(row, col0 + i), fill=C_ORANGE)


def page_setup(ws, rows_title="1:4", fit_height=0):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Page &P / &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.left.text = "&A"
    ws.oddFooter.left.size = 8
    if rows_title:
        ws.print_title_rows = rows_title


def protect(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def dv_list(ws, cells, items=None, ref=None):
    f = ref if ref else '"' + ",".join(items) + '"'
    dv = DataValidation(type="list", formula1=f, allow_blank=False, showDropDown=False)
    dv.error, dv.errorTitle = "Choisir une valeur de la liste.", "Saisie"
    ws.add_data_validation(dv)
    for c in cells:
        dv.add(c)


def dv_num(ws, cell, lo, hi, msg):
    dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi))
    dv.errorStyle, dv.error, dv.errorTitle = "warning", msg, "Valeur inhabituelle"
    dv.prompt, dv.promptTitle = msg, "Plage conseillée"
    ws.add_data_validation(dv)
    dv.add(cell)


def red_if(ws, rng, formula):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=RED_FILL, font=WHITE_BOLD))


M = Model()
GAB_DEFAUT = ["20 × 40"]
PARAM_TEST = {}


# ----------------------------------------------------------------------------- aides formules
def XL(e, col):
    """Recherche dans BASE_PROFILÉS par RECHERCHEX, repli INDEX/EQUIV (anciennes versions)."""
    keys = f"'{BP}'!$B$5:$B$200"
    vals = f"'{BP}'!${col}$5:${col}$200"
    return (f"IFERROR(_xlfn.XLOOKUP({e},{keys},{vals}),"
            f"IFERROR(INDEX({vals},MATCH({e},{keys},0)),0))")


def XA(e):
    """Poids (kg/ml) d'un acier HA par RECHERCHEX dans BASE_ACIERS."""
    keys, vals = f"'{BA}'!$B$5:$B$20", f"'{BA}'!$D$5:$D$20"
    return (f"IFERROR(_xlfn.XLOOKUP({e},{keys},{vals}),"
            f"IFERROR(INDEX({vals},MATCH({e},{keys},0)),0))")


def XD(e):
    """Diamètre (m) d'un acier HA."""
    keys, vals = f"'{BA}'!$B$5:$B$20", f"'{BA}'!$C$5:$C$20"
    return (f"IFERROR(_xlfn.XLOOKUP({e},{keys},{vals}),"
            f"IFERROR(INDEX({vals},MATCH({e},{keys},0)),0))/1000")


def PU_HA(e):
    s = "0"
    for d in reversed(DIAMS):
        s = f'IF({e}="{d}",{{pu_{d}}},{s})'
    return s


# ============================================================================= BASES
PROFILES = [
    # (désignation, famille, kg/ml, h mm)
    *[(f"IPE {h}", "IPE", w, h) for h, w in [(160, 15.8), (180, 18.8), (200, 22.4), (220, 26.2), (240, 30.7),
                                           (270, 36.1), (300, 42.2), (330, 49.1), (360, 57.1), (400, 66.3),
                                           (450, 77.6), (500, 90.7), (550, 106.0), (600, 122.0)]],
    *[(f"HEA {n}", "HEA", w, h) for n, w, h in [(140, 24.7, 133), (160, 30.4, 152), (180, 35.5, 171),
                                              (200, 42.3, 190), (220, 50.5, 210), (240, 60.3, 230),
                                              (260, 68.2, 250), (280, 76.4, 270), (300, 88.3, 290),
                                              (320, 97.6, 310), (340, 105.0, 330), (360, 112.0, 350),
                                              (400, 125.0, 390)]],
    *[(f"UPN {h}", "UPN", w, h) for h, w in [(80, 8.64), (100, 10.6), (120, 13.4), (140, 16.0), (160, 18.8),
                                           (180, 22.0), (200, 25.3), (220, 29.4), (240, 33.2), (260, 37.9),
                                           (300, 46.2)]],
    *[(f"L {a}", "Cornière égale", w, int(a.split("x")[0])) for a, w in
      [("40x4", 2.42), ("45x4.5", 3.06), ("50x5", 3.77), ("60x6", 5.42), ("70x7", 7.38), ("80x8", 9.63),
       ("90x9", 12.2), ("100x10", 15.0)]],
    *[(f"Z {h}", "Panne Z (formée à froid)", w, h) for h, w in [(120, 3.45), (140, 3.80), (160, 4.35),
                                                              (180, 4.75), (200, 5.20), (220, 6.90),
                                                              (250, 7.60)]],
    *[(f"C {h}", "Profil C (formé à froid)", w, h) for h, w in [(120, 3.92), (140, 4.24), (160, 4.57),
                                                              (180, 5.20), (200, 5.50), (220, 6.95),
                                                              (250, 7.60)]],
    *[(f"Tube {t}", "Tube carré", w, int(t.split("x")[0])) for t, w in
      [("40x40x3", 3.30), ("50x50x3", 4.25), ("60x60x3", 5.19), ("80x80x4", 9.22), ("100x100x4", 11.7),
       ("100x100x5", 14.4), ("120x120x5", 17.5)]],
    *[(f"Tube Ø{t}", "Tube rond", w, int(float(t.split("x")[0]))) for t, w in
      [("60.3x3.2", 4.51), ("76.1x3.2", 5.75), ("88.9x4", 8.38), ("114.3x4", 10.9)]],
    *[(f"Rond {d}", "Rond plein (tirants, liernes, tiges)", w, d) for d, w in
      [(10, 0.617), (12, 0.888), (14, 1.21), (16, 1.58), (18, 2.00), (20, 2.47), (22, 2.98), (24, 3.55),
       (27, 4.49), (30, 5.55)]],
]
ACIERS = [("HA6", 6, 0.222), ("HA8", 8, 0.395), ("HA10", 10, 0.617), ("HA12", 12, 0.888),
          ("HA14", 14, 1.208), ("HA16", 16, 1.578), ("HA20", 20, 2.466), ("HA25", 25, 3.853)]


def build_bases(wb):
    ws = wb.create_sheet(BP)
    title(ws, "BASE DES PROFILÉS – poids en kg/ml", "Valeurs indicatives (catalogues usuels). "
          "Recherche par RECHERCHEX dans toutes les feuilles. Ajouter de nouveaux profilés en fin de tableau.", 6)
    header(ws, 4, ["Désignation", "Famille", "kg/ml", "h (mm)", "Remarque"])
    for i, (n, f, w, h) in enumerate(PROFILES):
        r = 5 + i
        for j, v in enumerate([n, f, w, h]):
            c = ws.cell(r, 2 + j, v)
            st(c, fmt=NF_2 if j == 2 else None)
            c.protection = Protection(locked=False)
            c.fill = FILL(C_JAUNE)
        st(ws.cell(r, 6), size=8)
    for r in range(5 + len(PROFILES), 201):
        for j in range(4):
            c = ws.cell(r, 2 + j)
            c.protection = Protection(locked=False)
    ws.cell(5, 6, "Membrures de treillis : poids multiplié par le nombre de cornières accolées")
    for col, w in zip("ABCDEF", [2, 18, 30, 10, 10, 60]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    page_setup(ws)

    ws = wb.create_sheet(BA)
    title(ws, "BASE DES ACIERS HA – poids en kg/ml", "Aciers à haute adhérence FeE500 – barres commerciales de 12 m", 6)
    header(ws, 4, ["Désignation", "Ø (mm)", "kg/ml", "Section (cm²)", "Remarque"])
    for i, (n, d, w) in enumerate(ACIERS):
        r = 5 + i
        ws.cell(r, 2, n)
        ws.cell(r, 3, d)
        ws.cell(r, 4, w)
        ws.cell(r, 5, f"=PI()*C{r}^2/400")
        for j in range(4):
            st(ws.cell(r, 2 + j), fmt=["@", "0", "0.000", "0.00"][j])
        st(ws.cell(r, 6), size=8)
    ws.cell(5, 6, "Poids = 7 850 kg/m³ × π Ø² / 4")
    for col, w in zip("ABCDEF", [2, 14, 10, 10, 14, 40]):
        ws.column_dimensions[col].width = w
    page_setup(ws)


# ============================================================================= PARAMÈTRES
class ParamSheet:
    def __init__(self, ws):
        self.ws, self.row = ws, 5

    def sec(self, text):
        self.row += 1
        section(self.ws, self.row, text, 2, 4)
        self.row += 1

    def p(self, name, label, value, unit="", comment="", fmt=None, items=None, rng=None):
        ws, r = self.ws, self.row
        self.row += 1
        ws.cell(r, 2, label)
        st(ws.cell(r, 2))
        c = ws.cell(r, 3, value)
        if fmt is None:
            fmt = {"%": NF_PCT, "FCFA": NF_INT}.get(unit.split("/")[0].strip() if unit else "", None)
            if fmt is None:
                fmt = "0.00" if isinstance(value, float) else ("0" if isinstance(value, int) else "@")
        yellow(c, fmt)
        c.alignment = Alignment(horizontal="center")
        st(ws.cell(r, 4, unit), size=9, h="center")
        st(ws.cell(r, 5, comment), size=8, italic=True, wrap=False)
        if items:
            dv_list(ws, [f"C{r}"], items)
        if rng:
            dv_num(ws, f"C{r}", rng[0], rng[1], f"Plage conseillée : {rng[0]} à {rng[1]} {unit}")
        M.param(name, PAR, f"C{r}")
        return f"C{r}"

    def v(self, name, label, expr, unit="", comment="", fmt=NF_2, eng=None):
        ws, r = self.ws, self.row
        self.row += 1
        st(ws.cell(r, 2, label), italic=True)
        st(ws.cell(r, 3), fill=C_BLEU_CLAIR, fmt=fmt, bold=True, h="center")
        st(ws.cell(r, 4, unit), size=9, h="center")
        st(ws.cell(r, 5, comment), size=8, italic=True)
        M.var(name, expr, PAR, f"C{r}", fmt, eng)


def build_parametres(wb):
    ws = wb.create_sheet(PAR)
    title(ws, "PARAMÈTRES – toutes les données saisissables",
          "Cellules JAUNES = saisie. Cellules bleu clair = valeurs calculées. Les gabarits prédéfinis "
          "utilisent la hauteur et la pente ci-dessous ; PERSONNALISÉ utilise sa propre ligne du tableau (col. G à K).", 11)
    header(ws, 4, ["Paramètre", "Valeur", "Unité", "Plage conseillée / commentaire"])
    P = ParamSheet(ws)
    P.row = 4

    P.sec("1. GABARIT ET GÉOMÉTRIE")
    rh = P.p("hauteur_def", "Hauteur sous sablière (gabarits prédéfinis)", 6.0, "m", "4 à 12 m", rng=(3, 15))
    rp = P.p("pente_def", "Pente de toiture (gabarits prédéfinis)", 0.16, "%", "10 à 20 % conseillé ; < 5 % = alerte")
    red_if(ws, rp, f"{rp}<0.05")
    # tableau des gabarits
    gr0 = 5
    header(ws, gr0 - 1, ["Gabarit", "Portée (m)", "Longueur (m)", "Hauteur (m)", "Pente"], col0=7)
    for i, (n, b, l) in enumerate(GABARITS + [(PERSO, 40, 48)]):
        r = gr0 + i
        st(ws.cell(r, 7, n), bold=True)
        if n == PERSO:
            for col, v, f in [(8, 40.0, "0.00"), (9, 48.0, "0.00"), (10, 6.0, "0.00"), (11, 0.16, NF_PCT)]:
                yellow(ws.cell(r, col, v), f)
        else:
            for col, v, f in [(8, b, "0.00"), (9, l, "0.00"), (10, f"=$C${rh[1:]}", "0.00"),
                              (11, f"=$C${rp[1:]}", NF_PCT)]:
                st(ws.cell(r, col, v), fmt=f, h="center")
    st(ws.cell(gr0 + 8, 7, "Ligne PERSONNALISÉ : saisir portée, longueur, hauteur et pente. "
                         "Valeurs proposées = modèle de référence 40 × 48 m, H 6 m, pente 16 %."),
       size=8, italic=True, border=False)
    red_if(ws, f"K{gr0 + 7}", f"K{gr0 + 7}<0.05")
    G_RANGE = f"'{PAR}'!$G${gr0}:$G${gr0 + 7}"

    def gab_eng(col):
        return lambda g, s: f"'{PAR}'!${col}${gr0 + g}"

    P.v("gab_actif", "Gabarit actif (choisi en ACCUEIL)", "{gab}", "", "", fmt="@",
        eng=lambda g, s: f'"{GABARITS[g][0]}"')
    for name, lab, col, fmt in [("portee", "Portée active", "H", "0.00"), ("longueur", "Longueur active", "I", "0.00"),
                                ("hauteur", "Hauteur sous sablière active", "J", "0.00"),
                                ("pente", "Pente active", "K", NF_PCT)]:
        P.v(name, lab, f"INDEX('{PAR}'!${col}${gr0}:${col}${gr0 + 7},MATCH({{gab}},{G_RANGE},0))",
            "%" if fmt == NF_PCT else "m", "Valeur du gabarit actif", fmt=fmt, eng=gab_eng(col))

    P.sec("2. CHARPENTE – TRAMES ET SYSTÈME")
    P.p("e_port", "Entraxe des portiques", 6.0, "m", "6 m par défaut", rng=(4, 8))
    P.p("e_pan", "Entraxe des pannes", 1.35, "m", "1,20 à 1,50 m", rng=(1.2, 1.5))
    P.p("e_lis", "Entraxe des lisses de bardage", 1.80, "m", "1,50 à 2,00 m", rng=(1.5, 2.0))
    P.p("e_pp", "Espacement des poteaux de pignon", 6.0, "m", "5 à 6 m", rng=(5, 6))
    P.p("deb_pign", "Débord des pannes aux pignons", 0.30, "m", "")
    P.p("sys_2030", "Système pour portée de 20 à 30 m", "IPE renforcé", "", "IPE renforcé ou Ferme treillis",
        items=["IPE renforcé", "Ferme treillis"])
    P.p("file_int", "File de poteaux intermédiaires (file centrale)", "Non", "", "Conseillé au-delà de 30 m",
        items=["Oui", "Non"])
    P.p("ratio_treil", "Hauteur de ferme aux appuis = portée / …", 30, "", "25 à 35")
    P.p("pas_treil", "Pas des nœuds du treillis", 2.0, "m", "1,5 à 2,5 m")
    P.p("ratio_jarret", "Longueur de jarret (en % de la portée)", 0.10, "%", "8 à 12 %")
    P.p("nb_acc", "Nombre de cornières accolées par membrure", 2, "u", "2 cornières dos à dos")

    P.sec("3. OUVERTURES – SKYDOMES")
    P.p("nb_pc", "Portes coulissantes : nombre", 1, "u", "")
    P.p("l_pc", "Portes coulissantes : largeur", 4.0, "m", "")
    P.p("h_pc", "Portes coulissantes : hauteur", 4.5, "m", "")
    P.p("emp_pc", "Portes coulissantes : emplacement", "Pignon", "", "Long-pan ou Pignon", items=["Long-pan", "Pignon"])
    P.p("nb_pt", "Portillons : nombre", 1, "u", "Modèle de référence : portillon de 5,50 m en long-pan")
    P.p("l_pt", "Portillons : largeur", 5.5, "m", "")
    P.p("h_pt", "Portillons : hauteur", 4.0, "m", "")
    P.p("emp_pt", "Portillons : emplacement", "Long-pan", "", "Long-pan ou Pignon", items=["Long-pan", "Pignon"])
    P.p("deb_larm", "Débord du larmier de porte coulissante", 0.20, "m", "ml = largeur + 0,20 m")
    P.p("nb_sky", "Nombre de skydomes", 4, "u", "")
    P.p("a_sky", "Skydome : longueur", 1.20, "m", "")
    P.p("b_sky", "Skydome : largeur", 1.20, "m", "")

    P.sec("4. COUVERTURE – BARDAGE – FAÇADE")
    P.p("acrotere", "Présence d'acrotère", "Non", "", "Oui / Non", items=["Oui", "Non"])
    P.p("h_acr", "Hauteur d'acrotère", 1.00, "m", "")
    P.p("esp_baion", "Espacement des baïonnettes d'acrotère", 1.25, "m", "1,0 à 1,5 m", rng=(1.0, 1.5))
    P.p("pct_transl", "Translucides (% de la surface de couverture)", 0.10, "%", "10 % par défaut")
    P.p("rec_tole", "Recouvrements des tôles (majoration)", 0.10, "%", "8 à 12 %")
    c = P.p("ep_tole", "Épaisseur des tôles (bac acier)", 0.45, "mm", "0,40 à 0,50 mm mini (climat tropical humide)")
    red_if(ws, c, f"{c}<0.4")
    P.p("S_desc", "Surface de toiture par descente EP", 180.0, "m²", "150 à 200 m²", rng=(150, 200))
    P.p("vis_m2", "Vis autoforeuses par m² de tôle", 6, "u/m²", "environ 6 /m²")
    P.p("riv_ml", "Rivets par ml de chéneau", 5, "u/ml", "")
    P.p("facade", "Type de façade", "Bardage bac acier", "", "Bardage / Agglos 15 creux / Mixte", items=FACADES)
    P.p("h_soub", "Hauteur du soubassement en agglos (variante Mixte)", 1.50, "m", "")

    P.sec("5. ACIER – ASSEMBLAGES – PROTECTION ANTICORROSION")
    P.p("pct_chutes", "Chutes sur profilés", 0.05, "%", "5 % par défaut")
    P.p("pct_boul", "Boulonnerie : minimum en % du tonnage", 0.02, "%", "On retient le maximum entre détail et %")
    P.p("rho", "Masse volumique de l'acier", 7850, "kg/m³", "")
    P.p("L_barre", "Longueur commerciale des barres et profilés", 12.0, "m", "")
    P.p("protection", "Protection anticorrosion", "Galvanisation à chaud", "",
        "Climat tropical humide : protection obligatoire", items=["Galvanisation à chaud", "Peinture anticorrosion"])
    P.p("ratio_m2t", "Surface à peindre par tonne d'acier", 25.0, "m²/t", "20 à 30 m²/t")
    P.p("kg_echant", "Poids d'une échantignole (taquet de panne)", 1.80, "kg/u", "")
    P.p("kg_ecl_pan", "Poids d'une éclisse de panne", 2.50, "kg/u", "")
    P.p("nb_b_noeud", "Boulons HR M20 par assemblage poteau-traverse", 8, "u", "")
    P.p("nb_b_fait", "Boulons HR M20 par faîtage", 8, "u", "")
    P.p("nb_b_ech", "Boulons M12 par échantignole", 2, "u", "")
    P.p("nb_b_cv", "Boulons HR M16 par extrémité de diagonale", 2, "u", "")
    P.p("nb_b_lis", "Boulons M12 par attache de lisse", 2, "u", "")
    P.p("nb_b_ecl", "Boulons M12 par éclisse de panne", 4, "u", "")
    P.p("kg_HR20", "Poids boulon HR M20 (écrou + rondelles)", 0.45, "kg/u", "")
    P.p("kg_HR16", "Poids boulon HR M16 (écrou + rondelles)", 0.25, "kg/u", "")
    P.p("kg_M12", "Poids boulon ordinaire M12 (écrou + rondelle)", 0.10, "kg/u", "")
    ronds = [f"Rond {d}" for d in (16, 18, 20, 22, 24, 27, 30)]
    P.p("d_tige_port", "Tiges d'ancrage poteaux de portique : diamètre", "Rond 24", "", "", items=ronds)
    P.p("l_tige_port", "Tiges d'ancrage poteaux de portique : longueur", 0.60, "m", "")
    P.p("nb_tige_port", "Tiges d'ancrage par poteau de portique", 4, "u", "")
    P.p("d_tige_pig", "Tiges d'ancrage poteaux de pignon : diamètre", "Rond 20", "", "", items=ronds)
    P.p("l_tige_pig", "Tiges d'ancrage poteaux de pignon : longueur", 0.50, "m", "")
    P.p("nb_tige_pig", "Tiges d'ancrage par poteau de pignon", 2, "u", "")

    P.sec("6. GROS ŒUVRE – TERRASSEMENTS ET FONDATIONS")
    P.p("ep_dec", "Épaisseur de décapage", 0.20, "m", "")
    P.p("deb_dec", "Débord de décapage autour de l'emprise", 1.00, "m", "")
    P.p("surl", "Surlargeur de fouille (de chaque côté)", 0.10, "m", "")
    P.p("D", "Profondeur de fondation", 1.20, "m", "1,20 m (sol latéritique)")
    P.p("foison", "Coefficient de foisonnement", 1.25, "", "1,25")
    P.p("e_mur", "Épaisseur des murs (agglos de 15)", 0.15, "m", "")
    P.p("e_rem_c", "Épaisseur du remblai compacté (latérite 95 % OPM)", "20 cm", "", "20 cm / 40 cm",
        items=["20 cm", "40 cm"])
    P.p("e_sable", "Lit de sable de réglage sous dallage", 0.05, "m", "5 cm")
    P.p("e_prop", "Béton de propreté", 0.05, "m", "5 cm")
    P.p("a1", "Semelle isolée poteau de portique : côté", 1.20, "m", "120 × 120 cm")
    P.p("h_s1", "Semelle isolée poteau de portique : hauteur", 0.40, "m", "")
    P.p("f1", "Fût poteau de portique : côté", 0.40, "m", "40 × 40 cm")
    P.p("a2", "Semelle isolée poteau de pignon : côté", 0.80, "m", "80 × 80 cm")
    P.p("h_s2", "Semelle isolée poteau de pignon : hauteur", 0.30, "m", "")
    P.p("f2", "Fût poteau de pignon : côté", 0.30, "m", "30 × 30 cm")
    P.p("h_fut", "Hauteur des fûts (amorces) jusqu'au dallage", 1.20, "m", "1,20 m")
    P.p("m_sem", "Maille du quadrillage des semelles isolées", 0.15, "m", "")
    P.p("ha_sem", "Quadrillage inférieur des semelles : acier", "HA12", "", "", items=DIAMS)
    P.p("ha_fut", "Aciers longitudinaux des fûts", "HA12", "", "", items=DIAMS)
    P.p("nb_bfut", "Nombre de barres par fût", 4, "u", "4 HA12")
    P.p("ha_cad", "Cadres des fûts", "HA6", "", "", items=DIAMS)
    P.p("e_cad_fut", "Espacement des cadres des fûts", 0.15, "m", "15 cm")
    P.p("b_fil", "Semelle filante : largeur", 0.60, "m", "60 cm")
    P.p("h_fil", "Semelle filante : hauteur", 0.20, "m", "20 cm")
    P.p("ha_fil", "Semelle filante : aciers filants", "HA10", "", "", items=DIAMS)
    P.p("nb_fil", "Semelle filante : nombre de filants", 4, "u", "4 HA10")
    P.p("ha_rep", "Semelle filante : aciers de répartition", "HA8", "", "", items=DIAMS)
    P.p("e_rep", "Semelle filante : espacement de la répartition", 0.20, "m", "20 cm")
    P.p("c_fond", "Enrobage des fondations", 0.05, "m", "")
    P.p("retour", "Longueur des retours d'ancrage", 0.15, "m", "")
    P.p("crochet", "Longueur des crochets de cadres", 0.10, "m", "")
    P.p("h_mf", "Maçonnerie de fondation en agglos 15 pleins : hauteur", 1.00, "m", "1,00 m")

    P.sec("7. DALLAGE")
    c_ed = P.p("e_dal_c", "Épaisseur du dallage", "20 cm", "", "15 cm / 20 cm", items=["15 cm", "20 cm"])
    c_fe = P.p("ferr", "Ferraillage du dallage", "Nappe simple", "", "", items=["Nappe simple", "Nappe double"])
    red_if(ws, f"{c_ed}:{c_ed}", f'AND({c_ed}="15 cm",{c_fe}="Nappe double")')
    red_if(ws, f"{c_fe}:{c_fe}", f'AND({c_ed}="15 cm",{c_fe}="Nappe double")')
    P.p("ha_ix", "Nappe inférieure – sens X (long)", "HA10", "", "HA8 / HA10 / HA12", items=["HA8", "HA10", "HA12"])
    P.p("ha_iy", "Nappe inférieure – sens Y (travers)", "HA10", "", "", items=["HA8", "HA10", "HA12"])
    P.p("ha_sx", "Nappe supérieure – sens X (si nappe double)", "HA8", "", "", items=["HA8", "HA10", "HA12"])
    P.p("ha_sy", "Nappe supérieure – sens Y (si nappe double)", "HA8", "", "", items=["HA8", "HA10", "HA12"])
    P.p("maille_c", "Maille des aciers du dallage", "20 cm", "", "15 / 20 / 25 cm", items=["15 cm", "20 cm", "25 cm"])
    P.p("c_dal", "Enrobage du dallage", 0.03, "m", "")
    P.p("ha_chaise", "Chaises (nappe double) : acier", "HA8", "", "", items=DIAMS)
    P.p("ch_m2", "Chaises par m² (nappe double)", 4, "u/m²", "4 /m²")
    P.p("l_chaise", "Longueur développée d'une chaise", 0.60, "m", "")
    P.p("esp_joint", "Espacement des joints de retrait sciés", 5.0, "m", "5 à 6 m", rng=(5, 6))
    P.p("maj_poly", "Majoration du polyane (recouvrements, remontées)", 0.10, "%", "10 %")

    P.sec("8. VARIANTE ÉLÉVATION EN AGGLOS 15 CREUX")
    P.p("b_ch", "Chaînages : largeur", 0.15, "m", "15 × 20 cm")
    P.p("h_ch", "Chaînages : hauteur", 0.20, "m", "")
    P.p("ha_chl", "Chaînages, raidisseurs, linteaux : aciers longitudinaux", "HA10", "", "", items=DIAMS)
    P.p("nb_chl", "Nombre de barres longitudinales", 4, "u", "4 HA10")
    P.p("ha_chc", "Chaînages, raidisseurs, linteaux : cadres", "HA6", "", "", items=DIAMS)
    P.p("e_chc", "Espacement des cadres", 0.20, "m", "20 cm")
    P.p("b_rd", "Poteaux raidisseurs : côté", 0.15, "m", "15 × 15 cm")
    P.p("nb_rd", "Poteaux raidisseurs par panneau", 3, "u", "3 par panneau entre poteaux métalliques")
    P.p("appui_lin", "Appui des linteaux de chaque côté", 0.20, "m", "20 cm")
    P.p("enduit", "Enduit 2 faces", "Non", "", "Oui / Non", items=["Oui", "Non"])
    P.p("ep_enduit", "Épaisseur d'enduit", 0.02, "m", "")

    P.sec("9. DOSAGES ET RATIOS (CPJ 42.5)")
    P.p("dos_prop", "Dosage béton de propreté", 150, "kg/m³", "150 kg/m³")
    P.p("dos_ba", "Dosage béton armé", 350, "kg/m³", "350 kg/m³")
    P.p("dos_mort", "Dosage mortier", 300, "kg/m³", "300 kg/m³")
    P.p("sac", "Poids d'un sac de ciment CPJ 42.5", 50, "kg", "")
    P.p("sab_b", "Sable par m³ de béton", 0.40, "m³/m³", "")
    P.p("g1_b", "Gravier 5/15 par m³ de béton", 0.30, "m³/m³", "")
    P.p("g2_b", "Gravier 15/25 par m³ de béton", 0.50, "m³/m³", "")
    P.p("sab_m", "Sable par m³ de mortier", 1.10, "m³/m³", "")
    P.p("mort_m2", "Mortier de pose par m² d'agglos 15", 0.015, "m³/m²", "", fmt="0.000")
    P.p("agg_m2", "Agglos par m²", 12.5, "u/m²", "12,5 /m²")
    P.p("casse", "Casse des agglos", 0.05, "%", "5 %")
    P.p("rec_ha", "Recouvrement des aciers (× Ø)", 40, "Ø", "40 Ø")
    P.p("chutes_ha", "Chutes sur aciers HA", 0.05, "%", "5 %")
    P.p("pct_mo", "Main d'œuvre et matériel gros œuvre (% des matériaux)", 0.40, "%", "")

    P.sec("10. PRIX UNITAIRES (FCFA HT, rendu Abidjan)")
    for n, lab, v, u in [
        ("pu_acier", "Acier de charpente fourni et fabriqué", 1150, "FCFA/kg"),
        ("pu_montage", "Montage de la charpente", 250, "FCFA/kg"),
        ("pu_galva", "Galvanisation à chaud", 450, "FCFA/kg"),
        ("pu_peint", "Peinture anticorrosion (primaire + finition)", 3500, "FCFA/m²"),
        ("pu_tige", "Tige d'ancrage avec écrous et rondelles", 6000, "FCFA/u"),
        ("pu_gab", "Gabarit de pose des tiges", 7500, "FCFA/u"),
        ("pu_tole", "Couverture bac acier (fourniture et pose)", 6500, "FCFA/m²"),
        ("pu_transl", "Panneaux translucides", 12000, "FCFA/m²"),
        ("pu_fait", "Faîtière", 4500, "FCFA/ml"),
        ("pu_rive", "Rives", 3500, "FCFA/ml"),
        ("pu_larm", "Larmiers et bavettes de couverture", 3000, "FCFA/ml"),
        ("pu_chen", "Chéneaux", 15000, "FCFA/ml"),
        ("pu_naiss", "Naissances EP", 7500, "FCFA/u"),
        ("pu_desc", "Descentes EP PVC Ø110", 6000, "FCFA/ml"),
        ("pu_sky", "Skydome (fourni-posé)", 150000, "FCFA/u"),
        ("pu_vis", "Vis autoforeuses", 75, "FCFA/u"),
        ("pu_riv", "Rivets", 25, "FCFA/u"),
        ("pu_bard", "Bardage bac acier (fourniture et pose)", 6000, "FCFA/m²"),
        ("pu_larm_b", "Larmiers de bardage", 3000, "FCFA/ml"),
        ("pu_larm_p", "Larmiers de porte", 3500, "FCFA/ml"),
        ("pu_rail", "Rail haut de porte coulissante", 15000, "FCFA/ml"),
        ("pu_pc", "Portes coulissantes", 45000, "FCFA/m²"),
        ("pu_pt", "Portillons", 50000, "FCFA/m²"),
        ("pu_dec", "Décapage", 500, "FCFA/m²"),
        ("pu_fouil", "Fouilles", 4500, "FCFA/m³"),
        ("pu_rem", "Remblai latérite fourni et compacté", 9000, "FCFA/m³"),
        ("pu_evac", "Évacuation des déblais", 3500, "FCFA/m³"),
        ("pu_ciment", "Ciment CPJ 42.5 (sac de 50 kg)", 5500, "FCFA/sac"),
        ("pu_sable", "Sable", 12000, "FCFA/m³"),
        ("pu_g1", "Gravier 5/15", 22000, "FCFA/m³"),
        ("pu_g2", "Gravier 15/25", 20000, "FCFA/m³"),
        *[(f"pu_{d}", f"Acier {d}", 800 if d == "HA6" else (750 if d == "HA8" else 725), "FCFA/kg") for d in DIAMS],
        ("pu_ap", "Agglos 15 pleins", 450, "FCFA/u"),
        ("pu_ac", "Agglos 15 creux", 375, "FCFA/u"),
        ("pu_poly", "Film polyane", 400, "FCFA/m²"),
        ("pu_coff", "Bois de coffrage", 3500, "FCFA/m²"),
        ("pu_joint", "Joints de retrait sciés", 1500, "FCFA/ml"),
        ("pu_cure", "Cure du béton", 250, "FCFA/m²"),
    ]:
        P.p(n, lab, v, u, "", fmt=NF_INT)

    # ---- tableau des platines (colonnes G à L)
    r0 = 17
    st(ws.cell(r0 - 1, 7, "PLATINES ET PIÈCES PLATES (dimensions modifiables)"), bold=True, color=C_ORANGE,
       border=False)
    header(ws, r0, ["Pièce", "L (m)", "l (m)", "e (m)", "kg/u"], col0=7)
    for i, (n, lab, L, l, e) in enumerate([
        ("pl_pied_port", "Platine de pied – poteau de portique", 0.45, 0.35, 0.020),
        ("pl_pied_pig", "Platine de pied – poteau de pignon", 0.30, 0.25, 0.015),
        ("pl_tete", "Platine de tête de poteau / about de jarret", 0.55, 0.25, 0.020),
        ("pl_tete_pig", "Platine de tête – poteau de pignon", 0.25, 0.20, 0.012),
        ("pl_ecl_fait", "Éclisse de faîtage (une plaque)", 0.50, 0.25, 0.015),
        ("pl_gousset", "Gousset de nœud de treillis", 0.30, 0.30, 0.010),
    ]):
        r = r0 + 1 + i
        st(ws.cell(r, 7, lab))
        for j, v in enumerate([L, l, e]):
            yellow(ws.cell(r, 8 + j, v), "0.000")
        st(ws.cell(r, 11, f"=H{r}*I{r}*J{r}*{M.params['rho']}"), fmt=NF_2, fill=C_BLEU_CLAIR)
        M.param("kg_" + n, PAR, f"K{r}")
        M.param("L_" + n, PAR, f"H{r}")
    for col, w in zip("ABCDEFGHIJK", [2, 52, 18, 9, 46, 3, 42, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    page_setup(ws)
    return G_RANGE


# ============================================================================= PRÉDIMENSIONNEMENT
RULES = [  # borne sup. portée, système, poteau IPE, traverse IPE, poteau treillis, memb sup, memb inf, montants, diagonales
    (10, "Portique IPE + jarrets", "IPE 220", "IPE 200", "HEA 180", "L 50x5", "L 45x4.5", "L 40x4", "L 40x4"),
    (15, "Portique IPE + jarrets", "IPE 270", "IPE 240", "HEA 200", "L 60x6", "L 50x5", "L 40x4", "L 45x4.5"),
    (20, "Portique IPE + jarrets", "IPE 330", "IPE 300", "HEA 220", "L 70x7", "L 60x6", "L 45x4.5", "L 50x5"),
    (25, "IPE renforcé ou treillis", "IPE 400", "IPE 360", "HEA 240", "L 70x7", "L 60x6", "L 50x5", "L 50x5"),
    (30, "IPE renforcé ou treillis", "IPE 450", "IPE 400", "HEA 260", "L 80x8", "L 70x7", "L 50x5", "L 60x6"),
    (35, "Ferme treillis", "IPE 500", "IPE 450", "HEA 280", "L 90x9", "L 80x8", "L 50x5", "L 60x6"),
    (40, "Ferme treillis", "IPE 550", "IPE 500", "HEA 300", "L 100x10", "L 90x9", "L 60x6", "L 70x7"),
    (50, "Ferme treillis + file intermédiaire", "IPE 600", "IPE 550", "HEA 340", "L 100x10", "L 100x10", "L 60x6",
     "L 80x8"),
    (1000, "Ferme treillis + file intermédiaire", "IPE 600", "IPE 600", "HEA 400", "L 100x10", "L 100x10", "L 70x7",
     "L 90x9"),
]
RULES_PIG = [(6, "IPE 180"), (8, "IPE 220"), (10, "IPE 270"), (12, "IPE 300"), (1000, "IPE 360")]  # selon Hf
RULES_PAN = [(5, "Z 160", "C 140"), (6, "Z 200", "C 160"), (7, "Z 220", "C 180"), (1000, "Z 250", "C 200")]


def build_predim(wb):
    ws = wb.create_sheet(PRE)
    title(ws, "PRÉDIMENSIONNEMENT – choix automatique des profilés (modifiable en jaune)",
          "Portée ≤ 20 m : portique IPE avec jarrets · 20 à 30 m : IPE renforcé ou ferme treillis · "
          "> 30 m : ferme treillis (cornières) + alerte file intermédiaire. Valeurs de prédimensionnement à "
          "confirmer par une note de calcul EC3 / EC1 (vent).", 10)
    # --- règles
    section(ws, 4, "A. RÈGLES DE CHOIX SELON LA PORTÉE", 2, 9)
    header(ws, 5, ["Portée ≤ (m)", "Système indicatif", "Poteau (portique IPE)", "Traverse IPE",
                   "Poteau (treillis)", "Membrure sup.", "Membrure inf.", "Montants", "Diagonales"])
    for i, rule in enumerate(RULES):
        for j, v in enumerate(rule):
            c = ws.cell(6 + i, 2 + j, v)
            yellow(c) if j >= 2 else st(c, h="center" if j == 0 else None)
    rr0, rr1 = 6, 5 + len(RULES)
    rng = lambda col: f"'{PRE}'!${col}${rr0}:${col}${rr1}"
    r = rr1 + 2
    header(ws, r, ["H faîtage ≤ (m)", "Poteau de pignon"])
    header(ws, r, ["Entraxe ≤ (m)", "Panne", "Lisse"], col0=6)
    for i, (h, p) in enumerate(RULES_PIG):
        st(ws.cell(r + 1 + i, 2, h), h="center")
        yellow(ws.cell(r + 1 + i, 3, p))
    for i, (e, pa, li) in enumerate(RULES_PAN):
        st(ws.cell(r + 1 + i, 6, e), h="center")
        yellow(ws.cell(r + 1 + i, 7, pa))
        yellow(ws.cell(r + 1 + i, 8, li))
    pg0, pg1 = r + 1, r + len(RULES_PIG)
    pn0, pn1 = r + 1, r + len(RULES_PAN)

    # --- géométrie
    r = max(pg1, pn1) + 2
    section(ws, r, "B. GÉOMÉTRIE ET GRANDEURS CALCULÉES", 2, 9)
    header(ws, r + 1, ["Grandeur", "Valeur", "Unité", "Règle de calcul"])
    G = [r + 2]

    def g(name, label, expr, unit="", rule="", fmt=NF_2, eng=None):
        rw = G[0]
        G[0] += 1
        st(ws.cell(rw, 2, label))
        st(ws.cell(rw, 3), fmt=fmt, fill=C_BLEU_CLAIR, h="center", bold=True)
        st(ws.cell(rw, 4, unit), size=9, h="center")
        st(ws.cell(rw, 5, rule), size=8, italic=True)
        ws.merge_cells(start_row=rw, start_column=5, end_row=rw, end_column=10)
        M.var(name, expr, PRE, f"C{rw}", fmt, eng)

    g("idx_regle", "Ligne de règle (selon portée)", f'COUNTIF({rng("B")},"<"&{{portee}})+1', "", "Rang de la portée dans le tableau A", "0")
    g("systeme", "Système structurel retenu",
      'IF({portee}<=20,"Portique IPE avec jarrets",IF({portee}<=30,IF({sys_2030}="Ferme treillis",'
      '"Ferme treillis","Portique IPE renforcé"),"Ferme treillis"))', "", "Selon portée et choix 20–30 m", "@")
    g("treil", "Ferme treillis (1 = oui)", 'IF({systeme}="Ferme treillis",1,0)', "", "", "0")
    g("nt", "Nombre de travées", "MAX(1,ROUNDUP({longueur}/{e_port},0))", "u", "ARRONDI.SUP(L / entraxe)", "0")
    g("N", "Nombre de portiques N", "{nt}+1", "u", "N = ARRONDI.SUP(Longueur / entraxe) + 1", "0")
    g("er", "Entraxe réel des portiques", "{longueur}/{nt}", "m", "L / nombre de travées")
    g("R1", "Longueur rampante d'un versant", "({portee}/2)/COS(ATAN({pente}))", "m", "(portée / 2) / cos(arctan(pente))")
    g("h0", "Hauteur de ferme aux appuis", "{treil}*{portee}/{ratio_treil}", "m", "portée / ratio (0 si portique IPE)")
    g("Hm", "Hauteur des murs à l'égout", "{hauteur}+{h0}", "m", "H sablière + hauteur de ferme aux appuis")
    g("Hf", "Hauteur au faîtage", "{Hm}+{pente}*{portee}/2", "m", "Hm + pente × portée / 2")
    g("n_ip", "Intervalles entre poteaux de pignon (par pignon)", "MAX(1,ROUNDUP({portee}/{e_pp},0))", "u",
      "ARRONDI.SUP(portée / espacement)", "0")
    g("e_ppr", "Espacement réel des poteaux de pignon", "{portee}/{n_ip}", "m", "")
    g("n_pp", "Poteaux de pignon (2 pignons)", "2*({n_ip}-1)", "u", "2 × (intervalles − 1)", "0")
    g("n_int", "Poteaux intermédiaires (file centrale)", 'IF({file_int}="Oui",{N},0)', "u", "1 par portique si file centrale", "0")
    g("n_pan", "Panneaux du treillis", "{treil}*2*ROUNDUP({portee}/2/{pas_treil},0)", "u", "nombre pair ≈ portée / pas", "0")
    g("np", "Pannes par versant", "ROUNDUP({R1}/{e_pan},0)+1", "u", "ARRONDI.SUP(rampant / entraxe) + 1", "0")
    g("e_panr", "Entraxe réel des pannes", "{R1}/({np}-1)", "m", "")
    g("Lp", "Longueur des pannes (avec débords)", "{longueur}+2*{deb_pign}", "m", "L + 2 × débord")
    g("n_l", "Lignes de liernes par travée", "IF({er}>6,2,1)", "u", "mi-portée ; au tiers si entraxe > 6 m", "0")
    g("nb_cv", "Travées contreventées", "IF({nt}<=2,{nt},2+INT(({nt}-1)/5))", "u",
      "2 travées d'extrémité + 1 travée sur 5", "0")
    g("n_pv", "Panneaux de poutre au vent par versant", "MAX(1,ROUND({R1}/{er},0))", "u", "panneaux ≈ carrés", "0")
    g("S_bat", "Surface couverte (emprise)", "{portee}*{longueur}", "m²", "portée × longueur")
    g("P", "Périmètre", "2*({portee}+{longueur})", "m", "2 × (portée + longueur)")
    g("fac", "Type de façade (scénario)", "{facade}", "", "PARAMÈTRES (comparatif : bardage puis agglos)", "@",
      eng=lambda gi, s: ["{facade}", f'"{FACADES[0]}"', f'"{FACADES[1]}"'][s])
    g("z0", "Hauteur de maçonnerie d'élévation (base du bardage)",
      f'IF({{fac}}="{FACADES[1]}",{{Hm}},IF({{fac}}="Mixte",MIN({{h_soub}},{{Hm}}),0))', "m",
      "0 en bardage · Hm en agglos · soubassement en mixte")
    g("kbard", "Bardage présent (1 = oui)", f'IF({{fac}}="{FACADES[1]}",0,1)', "", "0 si façade en agglos", "0")
    g("kag", "Maçonnerie d'élévation présente (1 = oui)", f'IF({{fac}}="{FACADES[0]}",0,1)', "", "", "0")
    g("ktreil_pv", "Panneaux de treillis (calcul sécurisé)", "MAX(1,{n_pan})", "u", "évite la division par 0", "0")
    g("n_lp_pc", "Portes coulissantes en long-pan", 'IF({emp_pc}="Long-pan",{nb_pc},0)', "u", "", "0")
    g("n_lp_pt", "Portillons en long-pan", 'IF({emp_pt}="Long-pan",{nb_pt},0)', "u", "", "0")
    g("S_ouv", "Surface totale des ouvertures", "{nb_pc}*{l_pc}*{h_pc}+{nb_pt}*{l_pt}*{h_pt}", "m²", "")
    g("L_ouv", "Largeur cumulée des ouvertures", "{nb_pc}*{l_pc}+{nb_pt}*{l_pt}", "m", "")
    g("S_ouv_b", "Ouvertures dans la zone bardée",
      "{kbard}*({nb_pc}*{l_pc}*MAX(0,{h_pc}-{z0})+{nb_pt}*{l_pt}*MAX(0,{h_pt}-{z0}))", "m²", "partie au-dessus de z0")
    g("S_ouv_ag", "Ouvertures dans la maçonnerie",
      "{kag}*({nb_pc}*{l_pc}*MIN({h_pc},{z0})+{nb_pt}*{l_pt}*MIN({h_pt},{z0}))", "m²", "partie sous z0")
    g("n_rows", "Rangées de lisses en long-pan", "{kbard}*ROUNDUP(({Hm}-{z0})/{e_lis},0)", "u",
      "ARRONDI.SUP(hauteur bardée / entraxe)", "0")
    g("rc_pc", "Rangées de lisses coupées par une porte coulissante",
      "MIN({n_rows},MAX(0,ROUNDUP(({h_pc}-{z0})/{e_lis},0)))", "u", "", "0")
    g("rc_pt", "Rangées de lisses coupées par un portillon",
      "MIN({n_rows},MAX(0,ROUNDUP(({h_pt}-{z0})/{e_lis},0)))", "u", "", "0")
    g("S_pig_b", "Surface brute des 2 pignons dans la zone bardée",
      "{kbard}*2*({portee}*({Hm}-{z0})+{pente}*{portee}^2/4)", "m²", "trapèze : B × h + pente × B² / 4")
    g("S_couv", "Surface de couverture (rampants)", "2*{R1}*{Lp}", "m²", "2 × rampant × longueur couverte")
    g("nb_desc", "Descentes EP", "MAX(2,2*ROUNDUP({S_bat}/(2*{S_desc}),0))", "u",
      "1 pour 150 à 200 m², réparties sur les 2 long-pans", "0")
    g("pu_prot_kg", "Coût de protection ramené au kg",
      '{pu_galva}*IF({protection}="Galvanisation à chaud",1,0)+{pu_peint}*{ratio_m2t}/1000*IF({protection}="Galvanisation à chaud",0,1)',
      "FCFA/kg", "galvanisation au kg ou peinture × m²/t", NF_INT)

    # --- profilés retenus
    r = G[0] + 1
    section(ws, r, "C. PROFILÉS RETENUS", 2, 9)
    header(ws, r + 1, ["Élément", "Profilé automatique", "Forçage (saisie)", "Profilé retenu", "kg/ml", "h (mm)",
                       "Règle", "", ""])
    rows = [
        ("pot", "Poteaux de portique",
         f'IF({{treil}}=1,INDEX({rng("F")},{{idx_regle}}),INDEX({rng("D")},{{idx_regle}}))', "Tableau A selon portée et système"),
        ("trav", "Traverses (portique IPE)", f'INDEX({rng("E")},{{idx_regle}})', "Tableau A"),
        ("msup", "Membrures supérieures (treillis)", f'INDEX({rng("G")},{{idx_regle}})', "Tableau A – cornières accolées"),
        ("minf", "Membrures inférieures (treillis)", f'INDEX({rng("H")},{{idx_regle}})', "Tableau A – cornières accolées"),
        ("mont", "Montants (treillis)", f'INDEX({rng("I")},{{idx_regle}})', "Tableau A"),
        ("diag", "Diagonales (treillis)", f'INDEX({rng("J")},{{idx_regle}})', "Tableau A"),
        ("pp", "Poteaux de pignon",
         f"INDEX('{PRE}'!$C${pg0}:$C${pg1},COUNTIF('{PRE}'!$B${pg0}:$B${pg1},\"<\"&{{Hf}})+1)", "Selon hauteur au faîtage"),
        ("pint", "Poteaux intermédiaires (file centrale)", '"HEA 200"', "Valeur proposée"),
        ("panne", "Pannes",
         f"INDEX('{PRE}'!$G${pn0}:$G${pn1},COUNTIF('{PRE}'!$F${pn0}:$F${pn1},\"<\"&{{er}})+1)", "Selon entraxe des portiques"),
        ("lisse", "Lisses de bardage",
         f"INDEX('{PRE}'!$H${pn0}:$H${pn1},COUNTIF('{PRE}'!$F${pn0}:$F${pn1},\"<\"&{{er}})+1)", "Selon entraxe des portiques"),
        ("cvv", "Contreventements verticaux (croix)", '"L 70x7"', "Cornière proposée"),
        ("cvt", "Contreventements de couverture (poutre au vent)", '"L 60x6"', "Cornière proposée"),
        ("lierne", "Liernes et bretelles", '"Rond 12"', "Rond Ø12"),
        ("chev", "Chevêtres de skydome", '"L 60x6"', "Cadre en cornière"),
        ("baion", "Baïonnettes d'acrotère", '"L 50x5"', "Cornière"),
        ("bai", "Linteaux et potelets de baies", '"UPN 140"', "UPN"),
    ]
    for i, (k, lab, auto, rule) in enumerate(rows):
        rw = r + 2 + i
        st(ws.cell(rw, 2, lab))
        st(ws.cell(rw, 3), h="center", fill=C_GRIS)
        yellow(ws.cell(rw, 4))
        st(ws.cell(rw, 5), h="center", bold=True, fill=C_BLEU_CLAIR)
        st(ws.cell(rw, 6), fmt=NF_2, h="center")
        st(ws.cell(rw, 7), fmt="0", h="center")
        st(ws.cell(rw, 8, rule), size=8, italic=True)
        M.param("force_" + k, PRE, f"D{rw}")
        M.var("auto_" + k, auto, PRE, f"C{rw}", "@")
        M.var("prof_" + k, f'IF({{force_{k}}}="",{{auto_{k}}},{{force_{k}}})', PRE, f"E{rw}", "@",
              eng=(lambda kk: (lambda gi, s: "{auto_" + kk + "}"))(k))
        mult = "*{nb_acc}" if k in ("msup", "minf") else ""
        M.var("kg_" + k, XL("{prof_" + k + "}", "D") + mult, PRE, f"F{rw}", NF_2)
        M.var("h_" + k, XL("{prof_" + k + "}", "E"), PRE, f"G{rw}", "0")
        red_if(ws, f"F{rw}", f"F{rw}=0")
    r = r + 2 + len(rows) + 1
    # --- alertes
    section(ws, r, "D. ALERTES DE PRÉDIMENSIONNEMENT", 2, 9)
    alerts = [
        ('IF(AND({portee}>30,{file_int}="Non"),"ALERTE : portée > 30 m sans poteau intermédiaire – prévoir une '
         'file centrale (indispensable au-delà de 40 m sans appui)","OK : portée / appuis")'),
        ('IF({pente}<0.05,"ALERTE : pente < 5 % – risque de stagnation, revoir la pente","OK : pente ≥ 5 %")'),
        ('IF(AND({e_dal_c}="15 cm",{ferr}="Nappe double"),"ALERTE : dallage 15 cm en nappe double (peu courant)",'
         '"OK : dallage / ferraillage")'),
        ('IF(OR({kg_m2}<18,{kg_m2}>45),"ALERTE : ratio acier "&FIXED({kg_m2},1,TRUE)&" kg/m² hors fourchette 18–45",'
         '"OK : ratio acier "&FIXED({kg_m2},1,TRUE)&" kg/m²")'),
        ('IF({ep_tole}<0.4,"ALERTE : tôle < 0,40 mm","OK : épaisseur de tôle")'),
        ('IF(AND({er}>6.5,{prof_panne}="Z 250"),"ALERTE : entraxe de portiques élevé, vérifier les pannes",'
         '"OK : entraxe des portiques")'),
    ]
    for i, a in enumerate(alerts):
        rw = r + 1 + i
        st(ws.cell(rw, 2), bold=True)
        ws.merge_cells(start_row=rw, start_column=2, end_row=rw, end_column=10)
        M.disp(PRE, f"B{rw}", a)
        red_if(ws, f"B{rw}", f'LEFT(B{rw},6)="ALERTE"')
    for col, w in zip("ABCDEFGHIJ", [2, 44, 20, 16, 16, 12, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    page_setup(ws, rows_title="1:2")
    return alerts


# ============================================================================= LIGNES DE QUANTITATIF
class Lines:
    """Feuille de lignes : Désignation | Profilé/Section | Unité | Quantité | L unit. | L tot. | kg/ml | Poids | Règle."""

    def __init__(self, ws, sheet, row):
        self.ws, self.sheet, self.row = ws, sheet, row
        self.tags = {}

    def sec(self, text):
        self.row += 1
        section(self.ws, self.row, text, 2, 9)
        self.row += 1

    def line(self, key, label, unit, prof=None, q=None, lu=None, lt=None, kgml=None, kg=None, rule="", tag=None,
             qfmt=None):
        ws, r, S = self.ws, self.row, self.sheet
        self.row += 1
        st(ws.cell(r, 2, label), wrap=True)
        st(ws.cell(r, 4, unit), h="center", size=9)
        st(ws.cell(r, 10, rule), size=8, italic=True, wrap=True)
        for col in range(3, 10):
            if col != 4:
                st(ws.cell(r, col), h="center")
        if prof:
            M.var(key + "_p", prof, S, f"C{r}", "@")
        if lt is None and lu is not None and q is not None:
            lt = f"{{{key}_q}}*{{{key}_lu}}"
        if lu is None and lt is not None and q is not None:
            lu = f"IF({{{key}_q}}>0,{{{key}_lt}}/{{{key}_q}},0)"
        if kg is None and kgml is not None:
            kg = f"{{{key}_lt}}*{{{key}_kgml}}" if lt is not None else f"{{{key}_q}}*{{{key}_kgml}}"
        if qfmt is None:
            qfmt = NF_INT if unit in ("u", "barres", "sacs") else NF_2
        for suf, ex, col, fmt in [("q", q, "E", qfmt), ("lu", lu, "F", NF_2), ("lt", lt, "G", NF_2),
                                  ("kgml", kgml, "H", "0.000"), ("kg", kg, "I", NF_2)]:
            if ex is not None:
                M.var(f"{key}_{suf}", ex, S, f"{col}{r}", fmt)
        if q is not None and rule:
            ws.cell(r, 5).comment = Comment(rule, "Règle de calcul", width=320, height=90)
        if tag:
            self.tags.setdefault(tag, []).append(key)
        return key

    def total(self, name, label, expr, unit="kg", col="I", fmt=NF_2, rule=""):
        ws, r = self.ws, self.row
        self.row += 1
        st(ws.cell(r, 2, label), bold=True, fill=C_ORANGE_CLAIR)
        for c in range(3, 11):
            st(ws.cell(r, c), fill=C_ORANGE_CLAIR)
        ws.cell(r, 4, unit).font = Font(name=FONT, bold=True)
        st(ws.cell(r, 10, rule), size=8, italic=True, fill=C_ORANGE_CLAIR)
        M.var(name, expr, self.sheet, f"{col}{r}", fmt)
        ws.cell(r, ord(col) - 64).font = Font(name=FONT, bold=True)
        ws.cell(r, ord(col) - 64).number_format = fmt
        return name


def sum_of(keys, suf="kg"):
    return "+".join("{" + k + "_" + suf + "}" for k in keys) if keys else "0"


def lines_sheet(wb, name, ttl, sub):
    ws = wb.create_sheet(name)
    title(ws, ttl, sub, 10)
    header(ws, 4, ["Désignation", "Profilé / section", "Unité", "Quantité", "Long. unitaire (m)",
                   "Long. totale (m)", "kg/ml ou kg/u", "Poids total (kg)", "Règle de calcul (aussi en commentaire)"])
    for col, w in zip("ABCDEFGHIJ", [2, 50, 16, 8, 12, 12, 13, 11, 14, 62]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    page_setup(ws)
    return ws


def build_qc(wb):
    ws = lines_sheet(wb, QC, "QUANTITATIF CHARPENTE MÉTALLIQUE (EC3 / EC1)",
                     "Une ligne par élément. Toutes les quantités se recalculent depuis PARAMÈTRES et PRÉDIMENSIONNEMENT.")
    Lq = Lines(ws, QC, 4)
    L = Lq.line
    Lq.sec("A1. OSSATURE")
    L("pot", "Poteaux de portiques (2 par portique)", "ml", "{prof_pot}", "2*{N}", "{Hm}", kgml="{kg_pot}",
      rule="2 × N poteaux × hauteur à l'égout (H + hauteur de ferme)", tag="acier")
    L("pp", "Poteaux de pignon (hauteur variable avec la pente)", "ml", "{prof_pp}", "{n_pp}", None,
      "2*(({n_ip}-1)*{Hm}+{pente}*{e_ppr}*INT({n_ip}^2/4))", "{kg_pp}",
      rule="2 pignons × Σ (Hm + pente × min(x ; B − x)) ; Σ min = (B / n) × ENT(n² / 4)", tag="acier")
    L("pint", "Poteaux intermédiaires (file centrale)", "ml", "{prof_pint}", "{n_int}",
      "IF({treil}=1,{hauteur},{Hf})", kgml="{kg_pint}", rule="1 par portique si file centrale = Oui", tag="acier")
    L("trav", "Traverses de portique (2 versants)", "ml", "{prof_trav}", "{N}*(1-{treil})", "2*{R1}",
      kgml="{kg_trav}", rule="Longueur rampante = (portée / 2) / cos(arctan(pente)) × 2", tag="acier")
    L("jar", "Jarrets (renforts poteau-traverse)", "ml", "{prof_trav}", "2*{N}*(1-{treil})",
      "{ratio_jarret}*{portee}", kgml="{kg_trav}", rule="2 par portique, longueur = % × portée, découpés dans le profil de traverse", tag="acier")
    L("msup", "Ferme treillis – membrures supérieures", "ml", "{prof_msup}", "{N}*{treil}", "2*{R1}",
      kgml="{kg_msup}", rule="1 par ferme, longueur = 2 × rampant ; poids × cornières accolées", tag="acier")
    L("minf", "Ferme treillis – membrures inférieures", "ml", "{prof_minf}", "{N}*{treil}", "{portee}",
      kgml="{kg_minf}", rule="1 par ferme, longueur = portée", tag="acier")
    L("mont", "Ferme treillis – montants", "ml", "{prof_mont}", "{N}*{treil}*({n_pan}+1)", None,
      "{N}*{treil}*(({n_pan}+1)*{h0}+{pente}*({portee}/{ktreil_pv})*INT({n_pan}^2/4))", "{kg_mont}",
      rule="(n + 1) montants par ferme, hauteur h0 + pente × min(x ; B − x)", tag="acier")
    L("diag", "Ferme treillis – diagonales", "ml", "{prof_diag}", "{N}*{treil}*{n_pan}",
      "{treil}*SQRT(({portee}/{ktreil_pv})^2+({h0}+{pente}*{portee}/4)^2)", kgml="{kg_diag}",
      rule="n diagonales par ferme, longueur √(pas² + h moyenne²)", tag="acier")
    L("pl_pp", "Platines de pied – poteaux de portique", "u", '"PL"', "2*{N}+{n_int}", kgml="{kg_pl_pied_port}",
      rule="1 par poteau de portique ou intermédiaire", tag="acier")
    L("pl_pg", "Platines de pied – poteaux de pignon", "u", '"PL"', "{n_pp}", kgml="{kg_pl_pied_pig}",
      rule="1 par poteau de pignon", tag="acier")
    L("pl_tt", "Platines de tête de poteau et d'about de jarret", "u", '"PL"', "IF({treil}=1,2,4)*{N}",
      kgml="{kg_pl_tete}", rule="Portique : 2 platines par nœud × 2 nœuds ; treillis : 1 platine de tête par poteau", tag="acier")
    L("pl_tg", "Platines de tête – poteaux de pignon", "u", '"PL"', "{n_pp}", kgml="{kg_pl_tete_pig}",
      rule="1 par poteau de pignon", tag="acier")
    L("gous", "Goussets de nœuds de treillis", "u", '"PL"', "{N}*{treil}*(2*{n_pan}+2)", kgml="{kg_pl_gousset}",
      rule="2 (n + 1) nœuds par ferme", tag="acier")
    L("ecf", "Éclisses de faîtage (paires de plaques)", "u", '"PL"', "{N}", kgml="2*{kg_pl_ecl_fait}",
      rule="1 paire par portique", tag="acier")
    L("ecp", "Éclisses de pannes (à chaque raboutage)", "u", '"PL"', "2*{np}*MAX(0,ROUNDUP({Lp}/{L_barre},0)-1)",
      kgml="{kg_ecl_pan}", rule="Pannes × (ARRONDI.SUP(longueur / 12 m) − 1)", tag="acier")
    Lq.sec("A2. STABILITÉ")
    L("cvv", "Contreventements verticaux – croix de Saint-André (long-pans)", "ml", "{prof_cvv}",
      "{nb_cv}*2*2", "SQRT({er}^2+{Hm}^2)", kgml="{kg_cvv}",
      rule="Travées d'extrémité + 1 sur 5 ; 2 long-pans × 2 diagonales", tag="acier")
    L("cvt", "Contreventements de couverture – poutre au vent", "ml", "{prof_cvt}",
      "{nb_cv}*2*{n_pv}*2", "SQRT({er}^2+({R1}/{n_pv})^2)", kgml="{kg_cvt}",
      rule="Mêmes travées ; 2 versants × panneaux × 2 diagonales", tag="acier")
    L("lier", "Liernes de pannes", "ml", "{prof_lierne}", "{nt}*{n_l}*2*({np}-1)", "{e_panr}",
      kgml="{kg_lierne}", rule="À mi-portée (au tiers si entraxe > 6 m), entre pannes successives", tag="acier")
    L("bret", "Bretelles de faîtage", "ml", "{prof_lierne}", "{nt}*{n_l}*2",
      "SQRT(({er}/({n_l}+1))^2+{e_panr}^2)", kgml="{kg_lierne}", rule="2 par ligne de liernes au faîtage", tag="acier")
    L("lll", "Liernes de lisses – long-pans", "ml", "{prof_lierne}", "{nt}*{n_l}*2*{n_rows}",
      "IF({n_rows}>0,({Hm}-{z0})/{n_rows},0)", kgml="{kg_lierne}",
      rule="Par ligne de liernes : hauteur bardée ; 0 en façade agglos", tag="acier")
    L("llp", "Liernes de lisses – pignons", "ml", "{prof_lierne}", "{kbard}*2*{n_ip}", None,
      "{kbard}*2*{n_ip}*(({Hm}-{z0})+{pente}*{portee}/4)", "{kg_lierne}",
      rule="1 par intervalle de pignon, hauteur moyenne", tag="acier")
    Lq.sec("A3. COUVERTURE")
    L("pan", "Pannes", "ml", "{prof_panne}", "2*{np}", "{Lp}", kgml="{kg_panne}",
      rule="Par versant : ARRONDI.SUP(rampant / entraxe) + 1 ; longueur = L + débords", tag="acier")
    L("ech", "Échantignoles (taquets de pannes)", "u", '"Taquet"', "2*{np}*{N}", kgml="{kg_echant}",
      rule="1 par panne et par portique", tag="acier")
    L("chev", "Chevêtres de skydomes", "ml", "{prof_chev}", "{nb_sky}", "2*({a_sky}+{b_sky})",
      kgml="{kg_chev}", rule="1 cadre par skydome", tag="acier")
    L("baio", "Baïonnettes d'acrotère", "ml", "{prof_baion}",
      'IF({acrotere}="Oui",ROUNDUP({P}/{esp_baion},0),0)', "{h_acr}", kgml="{kg_baion}",
      rule="1 tous les 1,0 à 1,5 m de périmètre si acrotère", tag="acier")
    L("tole", "Tôles bac acier (ép. selon paramètre)", "m²", '"Bac acier "&FIXED({ep_tole},2,TRUE)&" mm"',
      "{S_couv}*(1+{rec_tole})-{pct_transl}*{S_couv}-{nb_sky}*{a_sky}*{b_sky}",
      rule="rampant × longueur × 2 × (1 + recouvrement) − translucides − skydomes")
    L("transl", "Panneaux sandwich translucides", "m²", '"Translucide"', "{pct_transl}*{S_couv}",
      rule="% × surface de couverture")
    L("fait", "Faîtières", "ml", '"Faîtière"', "{Lp}", rule="Longueur couverte")
    L("rive", "Rives", "ml", '"Rive"', "4*{R1}", rule="2 pignons × 2 versants × rampant")
    L("larm", "Larmiers / bavettes de couverture", "ml", '"Bavette"', "2*{Lp}", rule="2 égouts")
    L("chen", "Chéneaux d'eau (2 côtés)", "ml", '"Chéneau"', "2*{longueur}", rule="2 × longueur")
    L("naiss", "Naissances EP", "u", '"Naissance"', "{nb_desc}", rule="1 descente pour 150 à 200 m² de toiture")
    L("desc", "Descentes EP", "ml", '"PVC Ø110"', "{nb_desc}*{Hm}", rule="Nombre × hauteur")
    L("sky", "Skydomes", "u", '"Skydome"', "{nb_sky}", rule="Paramètre")
    L("riv", "Rivets de chéneaux", "u", '"Rivet"', "{chen_q}*{riv_ml}", rule="ml de chéneaux × rivets / ml")
    Lq.sec("A4. BARDAGE ET BAIES")
    L("lis_lp", "Lisses de bardage – long-pans", "ml", "{prof_lisse}", "2*{n_rows}*{nt}", None,
      "MAX(0,2*{n_rows}*{longueur}-{n_lp_pc}*{l_pc}*{rc_pc}-{n_lp_pt}*{l_pt}*{rc_pt})", "{kg_lisse}",
      rule="Rangées × 2 × L − ouvertures coupées ; 0 si agglos, au-dessus du soubassement si mixte", tag="acier")
    L("lis_pg", "Lisses de bardage – pignons", "ml", "{prof_lisse}", "ROUNDUP({lis_pg_lt}/{e_ppr},0)", None,
      "MAX(0,{S_pig_b}/{e_lis}-({nb_pc}-{n_lp_pc})*{l_pc}*{rc_pc}-({nb_pt}-{n_lp_pt})*{l_pt}*{rc_pt})",
      "{kg_lisse}", rule="Surface du trapèze bardé / entraxe − ouvertures coupées", tag="acier")
    L("lin", "Linteaux de baies", "ml", "{prof_bai}", "{nb_pc}+{nb_pt}", None, "{L_ouv}", "{kg_bai}",
      rule="1 au-dessus de chaque porte ou portillon", tag="acier")
    L("potl", "Potelets de baies", "ml", "{prof_bai}", "2*({nb_pc}+{nb_pt})", "{Hm}", kgml="{kg_bai}",
      rule="2 par baie, toute hauteur", tag="acier")
    L("bard", "Bardage bac acier", "m²", '"Bac acier"',
      '({kbard}*2*{longueur}*({Hm}-{z0})+{S_pig_b}-{S_ouv_b})*(1+{rec_tole})+'
      'IF({acrotere}="Oui",{kbard}*{P}*{h_acr},0)',
      rule="Long-pans + pignons en trapèze − ouvertures, × (1 + recouvrement) ; 0 si agglos")
    L("larp", "Larmiers de portes coulissantes", "ml", '"Larmier"', "{nb_pc}*({l_pc}+{deb_larm})",
      rule="Largeur de porte + 0,20 m")
    L("rail", "Rails hauts de portes coulissantes", "ml", '"Rail"', "{nb_pc}*2*{l_pc}", rule="2 × largeur par porte")
    L("larb", "Larmiers de bardage (pied + angles)", "ml", '"Larmier"',
      "{kbard}*(MAX(0,{P}-{L_ouv})+4*({Hm}-{z0}))", rule="Périmètre − portes + 4 angles ; 0 si agglos")
    Lq.sec("A5. ASSEMBLAGES ET ANCRAGES")
    L("b_nd", "Boulons HR M20 – assemblages poteau-traverse / jarrets", "u", '"HR M20"', "2*{N}*{nb_b_noeud}",
      kgml="{kg_HR20}", rule="2 nœuds par portique × boulons par nœud", tag="boul")
    L("b_ft", "Boulons HR M20 – faîtage", "u", '"HR M20"', "{N}*{nb_b_fait}", kgml="{kg_HR20}",
      rule="1 faîtage par portique", tag="boul")
    L("b_pn", "Boulons M12 – pannes sur échantignoles", "u", '"M12"', "{ech_q}*{nb_b_ech}", kgml="{kg_M12}",
      rule="Échantignoles × boulons", tag="boul")
    L("b_ec", "Boulons M12 – éclisses de pannes", "u", '"M12"', "{ecp_q}*{nb_b_ecl}", kgml="{kg_M12}",
      rule="Éclisses × boulons", tag="boul")
    L("b_cv", "Boulons HR M16 – contreventements", "u", '"HR M16"', "({cvv_q}+{cvt_q})*2*{nb_b_cv}",
      kgml="{kg_HR16}", rule="Diagonales × 2 extrémités × boulons", tag="boul")
    L("b_ls", "Boulons M12 – lisses", "u", '"M12"', "({lis_lp_q}+{lis_pg_q})*2*{nb_b_lis}", kgml="{kg_M12}",
      rule="Tronçons de lisses × 2 attaches × boulons", tag="boul")
    L("tg_po", "Tiges d'ancrage – poteaux de portique", "ml", "{d_tige_port}", "(2*{N}+{n_int})*{nb_tige_port}",
      "{l_tige_port}", kgml=XL("{d_tige_port}", "D"), rule="4 par poteau de portique", tag="boul")
    L("tg_pg", "Tiges d'ancrage – poteaux de pignon", "ml", "{d_tige_pig}", "{n_pp}*{nb_tige_pig}",
      "{l_tige_pig}", kgml=XL("{d_tige_pig}", "D"), rule="2 par poteau de pignon", tag="boul")
    L("ecr", "Écrous de tiges d'ancrage", "u", '"Écrou"', "2*({tg_po_q}+{tg_pg_q})", rule="2 par tige")
    L("rond", "Rondelles de tiges d'ancrage", "u", '"Rondelle"', "2*({tg_po_q}+{tg_pg_q})", rule="2 par tige")
    L("gab", "Gabarits de pose des tiges", "u", '"Gabarit"', "2*{N}+{n_int}+{n_pp}", rule="1 par poteau")
    L("vis", "Vis autoforeuses", "u", '"Vis"', "ROUNDUP({vis_m2}*({tole_q}+{transl_q}+{bard_q}),0)",
      rule="≈ 6 par m² de tôle, translucide et bardage")
    Lq.sec("A6. TOTAUX CHARPENTE")
    T = Lq.total
    T("kg_oss", "Acier de structure (profilés, platines, pièces)", sum_of(Lq.tags["acier"]), "kg",
      rule="Somme des lignes acier A1 à A4")
    T("kg_boul_d", "Boulonnerie et tiges d'ancrage (détail)", sum_of(Lq.tags["boul"]), "kg", rule="Somme A5")
    T("kg_boul", "Boulonnerie retenue", "MAX({kg_boul_d},{pct_boul}*{kg_oss})", "kg",
      rule="Maximum (détail ; % du tonnage)")
    T("kg_chutes", "Chutes", "{pct_chutes}*{kg_oss}", "kg", rule="% chutes × acier de structure")
    T("kg_tot", "TONNAGE TOTAL CHARPENTE", "{kg_oss}+{kg_boul}+{kg_chutes}", "kg", rule="Structure + boulonnerie + chutes")
    T("t_tot", "Tonnage total", "{kg_tot}/1000", "t", fmt=NF_2)
    T("kg_m2", "Ratio acier", "{kg_tot}/{S_bat}", "kg/m²", rule="Contrôle : 18 à 45 kg/m²")
    red_if(ws, M.vars["kg_m2"]["coord"], f'OR({M.vars["kg_m2"]["coord"]}<18,{M.vars["kg_m2"]["coord"]}>45)')
    T("galva_kg", "Protection : galvanisation (acier à traiter)", "{kg_oss}", "kg", rule="Acier de structure")
    T("peint_m2", "Protection : peinture (surface)", "{kg_oss}/1000*{ratio_m2t}", "m²", rule="t × m²/t")
    return Lq


def build_go(wb):
    ws = lines_sheet(wb, GO, "GROS ŒUVRE ET FONDATIONS (BAEL 91 mod. 99 / EC2)",
                     "Sol latéritique : remblai en latérite compactée à 95 % OPM. Ciment CPJ 42.5. "
                     "Aciers : longueurs avec recouvrement de 40 Ø et barres de 12 m.")
    Lg = Lines(ws, GO, 4)
    L = Lg.line
    Lg.sec("B0. DONNÉES DE CALCUL")
    for k, lab, ex, u, rule in [
        ("e_dal", "Épaisseur du dallage", "VALUE(LEFT({e_dal_c},2))/100", "m", "Liste 15 / 20 cm"),
        ("e_rem", "Épaisseur du remblai compacté", "VALUE(LEFT({e_rem_c},2))/100", "m", "Liste 20 / 40 cm"),
        ("maille", "Maille du dallage", "VALUE(LEFT({maille_c},2))/100", "m", "Liste 15 / 20 / 25 cm"),
        ("kdbl", "Nappe double (1 = oui)", 'IF({ferr}="Nappe double",1,0)', "", ""),
        ("n_s1", "Semelles isolées 1 (poteaux de portique et file centrale)", "2*{N}+{n_int}", "u", ""),
        ("n_s2", "Semelles isolées 2 (poteaux de pignon)", "{n_pp}", "u", ""),
        ("L_fil", "Longueur de semelle filante", "MAX(0,{P}-2*{N}*{a1}-{n_pp}*{a2})", "m",
         "Périmètre − largeur des semelles isolées traversées"),
        ("Bx", "Dallage : longueur intérieure (sens X)", "{longueur}-{e_mur}", "m", "L − épaisseur de mur"),
        ("By", "Dallage : largeur intérieure (sens Y)", "{portee}-{e_mur}", "m", "B − épaisseur de mur"),
        ("S_int", "Surface intérieure", "{Bx}*{By}", "m²", ""),
        ("h_ag", "Hauteur de maçonnerie d'élévation", "{z0}", "m", "selon type de façade"),
        ("nb_ch", "Nombre de chaînages horizontaux", f'IF({{fac}}="{FACADES[1]}",3,IF({{fac}}="Mixte",2,0))', "u",
         "Agglos : bas, intermédiaire, haut · Mixte : bas et haut"),
        ("n_panx", "Panneaux de maçonnerie (travées du périmètre)", "{kag}*(2*{nt}+2*{n_ip})", "u", ""),
        ("c_beton", "Coût matériaux d'1 m³ de béton armé",
         "{dos_ba}/{sac}*{pu_ciment}+{sab_b}*{pu_sable}+{g1_b}*{pu_g1}+{g2_b}*{pu_g2}", "FCFA/m³", ""),
        ("c_mort", "Coût matériaux d'1 m³ de mortier", "{dos_mort}/{sac}*{pu_ciment}+{sab_m}*{pu_sable}", "FCFA/m³", ""),
    ]:
        r = Lg.row
        Lg.row += 1
        st(ws.cell(r, 2, lab), italic=True)
        st(ws.cell(r, 4, u), size=9, h="center")
        st(ws.cell(r, 5), fill=C_BLEU_CLAIR, bold=True, h="center")
        st(ws.cell(r, 10, rule), size=8, italic=True)
        fmt = NF_INT if u in ("u", "", "FCFA/m³") else NF_2
        M.var(k, ex, GO, f"E{r}", fmt)

    def lap(len_expr, d):
        """Longueur d'une barre continue avec recouvrements de 40 Ø tous les 12 m."""
        return f"({len_expr}+MAX(0,ROUNDUP(({len_expr})/{{L_barre}},0)-1)*{{rec_ha}}*{XD(d)})"

    def steel(key, label, d, q, lu, rule, tag):
        L(key, label, "barres", d, q, lu, kgml=XA(d), rule=rule, tag=tag)
        Lg.tags.setdefault("ha", []).append(key)

    Lg.sec("B1. TERRASSEMENTS")
    L("dec", "Décapage de l'emprise", "m²", None, "({portee}+2*{deb_dec})*({longueur}+2*{deb_dec})",
      rule="(B + 2 débords) × (L + 2 débords)")
    L("fp1", "Fouilles en puits – semelles isolées 1", "m³", None, "{n_s1}*({a1}+2*{surl})^2*{D}",
      rule="n × (a + 2 surlargeurs)² × profondeur 1,20 m", tag="fouil")
    L("fp2", "Fouilles en puits – semelles isolées 2", "m³", None, "{n_s2}*({a2}+2*{surl})^2*{D}",
      rule="n × (a + 2 surlargeurs)² × profondeur", tag="fouil")
    L("fr", "Fouilles en rigole – semelles filantes", "m³", None, "{L_fil}*({b_fil}+2*{surl})*{D}",
      rule="Longueur filante × (b + 2 surlargeurs) × profondeur", tag="fouil")
    L("rem", "Remblai latérite compactée sous dallage (approvisionnement)", "m³", None,
      "{S_int}*{e_rem}*{foison}", rule="Surface intérieure × épaisseur × foisonnement 1,25 (95 % OPM)")
    L("sab", "Lit de sable de réglage (nid de sable)", "m³", None, "{S_int}*{e_sable}", rule="Surface × 5 cm")
    L("evac", "Évacuation des déblais", "m³", None, "({dec_q}*{ep_dec}+{fp1_q}+{fp2_q}+{fr_q})*{foison}",
      rule="(décapage + fouilles) × foisonnement")

    for i, (n, a, hs, f, lab) in enumerate([("s1", "a1", "h_s1", "f1", "portique (120 × 120)"),
                                            ("s2", "a2", "h_s2", "f2", "pignon (80 × 80)")]):
        Lg.sec(f"B{2 + i}. SEMELLES ISOLÉES – POTEAUX DE {lab.upper()}")
        nn = "{n_" + n + "}"
        L(f"{n}_bp", "Béton de propreté (5 cm)", "m³", None, f"{nn}*{{{a}}}^2*{{e_prop}}", rule="n × a² × 5 cm", tag="bp")
        L(f"{n}_ba", "Béton armé – semelles", "m³", None, f"{nn}*{{{a}}}^2*{{{hs}}}", rule="n × a² × h", tag="ba")
        L(f"{n}_bf", "Béton armé – fûts (amorces)", "m³", None, f"{nn}*{{{f}}}^2*{{h_fut}}",
          rule="n × côté² × hauteur du fût (1,20 m)", tag="ba")
        L(f"{n}_cf", "Coffrage semelles et fûts", "m²", None, f"{nn}*(4*{{{a}}}*{{{hs}}}+4*{{{f}}}*{{h_fut}})",
          rule="n × (4 a h + 4 côté × h fût)", tag="coff")
        steel(f"{n}_aq", "Aciers – quadrillage inférieur", "{ha_sem}",
              f"{nn}*2*(ROUNDDOWN(({{{a}}}-2*{{c_fond}})/{{m_sem}},0)+1)", f"{{{a}}}-2*{{c_fond}}+2*{{retour}}",
              "2 sens × (ENT((a − 2c) / maille) + 1) barres ; L = a − 2c + 2 retours", "fond")
        steel(f"{n}_af", "Aciers – fûts (barres longitudinales)", "{ha_fut}", f"{nn}*{{nb_bfut}}",
              f"{{h_fut}}+{{{hs}}}-{{c_fond}}+{{retour}}", "4 barres par fût, ancrées dans la semelle", "fond")
        steel(f"{n}_ac", "Aciers – cadres des fûts", "{ha_cad}",
              f"{nn}*(ROUNDDOWN({{h_fut}}/{{e_cad_fut}},0)+1)", f"4*({{{f}}}-2*{{c_fond}})+2*{{crochet}}",
              "Cadres tous les 15 cm", "fond")

    Lg.sec("B4. SEMELLES FILANTES (SOUS MURS)")
    L("sf_bp", "Béton de propreté (5 cm)", "m³", None, "{L_fil}*{b_fil}*{e_prop}", rule="L × b × 5 cm", tag="bp")
    L("sf_ba", "Béton armé – semelle filante 60 × 20", "m³", None, "{L_fil}*{b_fil}*{h_fil}", rule="L × b × h", tag="ba")
    L("sf_cf", "Coffrage semelle filante", "m²", None, "2*{L_fil}*{h_fil}", rule="2 faces × L × h", tag="coff")
    steel("sf_al", "Aciers filants (4 HA10)", "{ha_fil}", "{nb_fil}", lap("{L_fil}", "{ha_fil}"),
          "4 filants × longueur avec recouvrements de 40 Ø", "fond")
    steel("sf_ar", "Aciers de répartition (HA8 e = 20 cm)", "{ha_rep}", "ROUNDDOWN({L_fil}/{e_rep},0)+1",
          "{b_fil}-2*{c_fond}", "L / espacement + 1 ; L = b − 2c", "fond")

    Lg.sec("B5. MAÇONNERIE DE FONDATION (AGGLOS 15 PLEINS)")
    L("mf_s", "Maçonnerie agglos 15 pleins (h = 1,00 m)", "m²", None, "{P}*{h_mf}", rule="Périmètre × 1,00 m")
    L("mf_u", "Agglos 15 pleins", "u", None, "ROUNDUP({mf_s_q}*{agg_m2}*(1+{casse}),0)",
      rule="12,5 /m² + 5 % de casse")
    L("mf_m", "Mortier de pose", "m³", None, "{mf_s_q}*{mort_m2}", rule="m² × mortier / m²", tag="mort")

    Lg.sec("B6. DALLAGE INTÉRIEUR")
    L("poly", "Film polyane", "m²", None, "{S_int}*(1+{maj_poly})", rule="Surface intérieure × 1,10")
    L("dal", "Béton du dallage", "m³", None, "{S_int}*{e_dal}", rule="Surface × épaisseur (15 ou 20 cm)", tag="ba")
    steel("di_x", "Nappe inférieure – sens X", "{ha_ix}", "ROUNDDOWN({By}/{maille},0)+1",
          lap("{Bx}-2*{c_dal}", "{ha_ix}"), "(côté / maille) + 1 barres ; L avec recouvrements 40 Ø", "dal")
    steel("di_y", "Nappe inférieure – sens Y", "{ha_iy}", "ROUNDDOWN({Bx}/{maille},0)+1",
          lap("{By}-2*{c_dal}", "{ha_iy}"), "(côté / maille) + 1 barres", "dal")
    steel("ds_x", "Nappe supérieure – sens X (si nappe double)", "{ha_sx}", "{kdbl}*(ROUNDDOWN({By}/{maille},0)+1)",
          lap("{Bx}-2*{c_dal}", "{ha_sx}"), "0 en nappe simple", "dal")
    steel("ds_y", "Nappe supérieure – sens Y (si nappe double)", "{ha_sy}", "{kdbl}*(ROUNDDOWN({Bx}/{maille},0)+1)",
          lap("{By}-2*{c_dal}", "{ha_sy}"), "0 en nappe simple", "dal")
    steel("chai", "Chaises (nappe double, 4 /m²)", "{ha_chaise}", "{kdbl}*ROUNDUP({ch_m2}*{S_int},0)", "{l_chaise}",
          "4 chaises / m² en nappe double", "dal")
    L("joint", "Joints de retrait sciés", "ml", None,
      "(ROUNDUP({Bx}/{esp_joint},0)-1)*{By}+(ROUNDUP({By}/{esp_joint},0)-1)*{Bx}", rule="Tous les 5 à 6 m dans les 2 sens")
    L("cure", "Cure du béton", "m²", None, "{S_int}", rule="Surface du dallage")

    Lg.sec("B7. VARIANTE ÉLÉVATION EN AGGLOS 15 CREUX (selon type de façade)")
    L("ag_s", "Surface des murs (− ouvertures, pignons en trapèze)", "m²", None,
      f'MAX(0,{{kag}}*({{P}}*{{h_ag}}+IF({{fac}}="{FACADES[1]}",2*{{pente}}*{{portee}}^2/4,0))-{{S_ouv_ag}})',
      rule="Périmètre × hauteur + pignons (agglos) − ouvertures")
    L("ag_u", "Agglos 15 creux", "u", None, "ROUNDUP({ag_s_q}*{agg_m2}*(1+{casse}),0)", rule="12,5 /m² + 5 % de casse")
    L("ag_m", "Mortier de pose", "m³", None, "{ag_s_q}*{mort_m2}", rule="m² × mortier / m²", tag="mort")
    L("ch_l", "Chaînages 15 × 20 (bas, intermédiaire, haut + rampants)", "ml", None,
      f'{{nb_ch}}*{{P}}+IF({{fac}}="{FACADES[1]}",4*{{R1}},0)', rule="Nombre de chaînages × périmètre (+ rampants des pignons)")
    L("ch_b", "Béton des chaînages", "m³", None, "{ch_l_q}*{b_ch}*{h_ch}", rule="ml × 0,15 × 0,20", tag="ba")
    L("ch_c", "Coffrage des chaînages", "m²", None, "{ch_l_q}*2*{h_ch}", rule="2 faces", tag="coff")
    steel("ch_a", "Chaînages – aciers longitudinaux (4 HA10)", "{ha_chl}", "{nb_chl}",
          "{ch_l_q}*(1+{rec_ha}*" + XD("{ha_chl}") + "/{L_barre})", "4 barres × ml × (1 + 40 Ø / 12 m)", "elev")
    steel("ch_k", "Chaînages – cadres HA6 e = 20 cm", "{ha_chc}", "ROUNDDOWN({ch_l_q}/{e_chc},0)+{nb_ch}",
          "2*({b_ch}+{h_ch})-8*{c_fond}+2*{crochet}", "ml / 0,20 + 1 par chaînage", "elev")
    L("rd_n", "Poteaux raidisseurs 15 × 15 (3 par panneau)", "u", None, "{nb_rd}*{n_panx}",
      rule="3 × nombre de travées sur le périmètre")
    L("rd_b", "Béton des raidisseurs", "m³", None, "{rd_n_q}*{b_rd}^2*{h_ag}", rule="n × 0,15² × hauteur du mur", tag="ba")
    L("rd_c", "Coffrage des raidisseurs", "m²", None, "{rd_n_q}*2*{b_rd}*{h_ag}", rule="2 faces coffrées", tag="coff")
    steel("rd_a", "Raidisseurs – aciers (4 HA10)", "{ha_chl}", "{rd_n_q}*{nb_chl}", "{h_ag}+2*{retour}",
          "4 barres × (hauteur + 2 retours)", "elev")
    steel("rd_k", "Raidisseurs – cadres HA6 e = 20 cm", "{ha_chc}", "{rd_n_q}*(ROUNDDOWN({h_ag}/{e_chc},0)+1)",
          "4*{b_rd}-8*{c_fond}+2*{crochet}", "hauteur / 0,20 + 1 par raidisseur", "elev")
    L("li_l", "Linteaux BA au-dessus des baies (appuis 20 cm)", "ml", None,
      "{kag}*(IF({h_pc}<{h_ag},{nb_pc}*({l_pc}+2*{appui_lin}),0)+IF({h_pt}<{h_ag},{nb_pt}*({l_pt}+2*{appui_lin}),0))",
      rule="Baies dans la maçonnerie : largeur + 2 × 0,20 m")
    L("li_b", "Béton des linteaux", "m³", None, "{li_l_q}*{b_ch}*{h_ch}", rule="ml × 0,15 × 0,20", tag="ba")
    L("li_c", "Coffrage des linteaux", "m²", None, "{li_l_q}*(2*{h_ch}+{b_ch})", rule="2 joues + fond", tag="coff")
    steel("li_a", "Linteaux – aciers (4 HA10)", "{ha_chl}", "{nb_chl}", "{li_l_q}", "4 barres × ml", "elev")
    steel("li_k", "Linteaux – cadres HA6", "{ha_chc}", "ROUNDDOWN({li_l_q}/{e_chc},0)", "2*({b_ch}+{h_ch})-8*{c_fond}+2*{crochet}",
          "ml / 0,20", "elev")
    L("end", "Enduit 2 faces (option)", "m²", None, 'IF({enduit}="Oui",2*{ag_s_q},0)', rule="2 × surface des murs")
    L("end_m", "Mortier d'enduit", "m³", None, "{end_q}*{ep_enduit}", rule="m² × épaisseur", tag="mort")

    Lg.sec("B8. RÉCAPITULATIF DES MATÉRIAUX DE GROS ŒUVRE")
    T = Lg.total
    t = Lg.tags
    T("V_bp", "Béton de propreté (150 kg/m³)", sum_of(t["bp"], "q"), "m³", col="E")
    T("V_ba", "Béton armé (350 kg/m³)", sum_of(t["ba"], "q"), "m³", col="E")
    T("V_mort", "Mortier (300 kg/m³)", sum_of(t["mort"], "q"), "m³", col="E")
    T("V_beton", "Béton total", "{V_bp}+{V_ba}", "m³", col="E")
    T("ciment", "Ciment CPJ 42.5", "ROUNDUP(({V_bp}*{dos_prop}+{V_ba}*{dos_ba}+{V_mort}*{dos_mort})/{sac},0)",
      "sacs", col="E", fmt=NF_INT, rule="Σ volumes × dosages / 50 kg")
    T("sable", "Sable", "{V_beton}*{sab_b}+{V_mort}*{sab_m}+{sab_q}", "m³", col="E",
      rule="béton × 0,40 + mortier × 1,10 + lit de sable")
    T("g1", "Gravier 5/15", "{V_beton}*{g1_b}", "m³", col="E")
    T("g2", "Gravier 15/25", "{V_beton}*{g2_b}", "m³", col="E")
    for d in DIAMS:
        expr = "+".join(f'IF({{{k}_p}}="{d}",{{{k}_kg}},0)' for k in t["ha"])
        T(f"kg_{d}", f"Aciers {d} (chutes comprises)", f"({expr})*(1+{{chutes_ha}})", "kg", col="I",
          rule="Somme des lignes de ce diamètre × (1 + chutes)")
        M.var(f"bar_{d}", f'ROUNDUP({{kg_{d}}}/({XA(chr(34) + d + chr(34))}*{{L_barre}}),0)', GO,
              f"E{Lg.row - 1}", NF_INT)
        ws.cell(Lg.row - 1, 6, "← barres de 12 m").font = Font(name=FONT, size=8, italic=True)
    T("kg_HA", "Aciers HA – total", "+".join("{kg_" + d + "}" for d in DIAMS), "kg", col="I")
    T("n_ap", "Agglos 15 pleins", "{mf_u_q}", "u", col="E", fmt=NF_INT)
    T("n_ac", "Agglos 15 creux", "{ag_u_q}", "u", col="E", fmt=NF_INT)
    T("polyane", "Polyane", "{poly_q}", "m²", col="E")
    T("coffrage", "Bois de coffrage", sum_of(t["coff"], "q"), "m²", col="E")
    # coût de la façade maçonnée (pour le comparatif)
    elev = [k for k in t["ha"] if k.startswith(("ch_", "rd_", "li_"))]
    acier_elev = "+".join(f"{{{k}_kg}}*{PU_HA('{' + k + '_p}')}" for k in elev)
    T("cout_fac_ag", "Coût matériaux + main d'œuvre de la façade maçonnée",
      f"({{ag_u_q}}*{{pu_ac}}+({{ch_b_q}}+{{rd_b_q}}+{{li_b_q}})*{{c_beton}}+({{ag_m_q}}+{{end_m_q}})*{{c_mort}}"
      f"+({acier_elev})*(1+{{chutes_ha}})+({{ch_c_q}}+{{rd_c_q}}+{{li_c_q}})*{{pu_coff}})*(1+{{pct_mo}})",
      "FCFA", col="I", fmt=NF_INT, rule="Pour le comparatif des façades (hors TVA)")
    return Lg


# ============================================================================= DQE
def build_dqe(wb):
    ws = wb.create_sheet(DQE)
    title(ws, "DÉTAIL QUANTITATIF ET ESTIMATIF (DQE) – DEVIS", "Montants en FCFA HT. Prix unitaires saisis "
          "dans PARAMÈTRES (section 10). TVA 18 %.", 7)
    header(ws, 4, ["N°", "Désignation", "Unité", "Quantité", "PU (FCFA)", "Montant (FCFA)"])
    R = [5]
    lots = []

    def lot(text):
        R[0] += 1
        section(ws, R[0], text, 2, 6)
        R[0] += 1
        lots.append((text, []))

    def item(no, label, unit, q, pu):
        r = R[0]
        R[0] += 1
        key = "dqe_" + no.replace(".", "_")
        st(ws.cell(r, 2, no), h="center")
        st(ws.cell(r, 3, label))
        cu = ws.cell(r, 4)
        st(cu, h="center")
        if unit.startswith("="):
            M.disp(DQE, f"D{r}", unit[1:])
        else:
            cu.value = unit
        st(ws.cell(r, 5), fmt=NF_2, h="right")
        st(ws.cell(r, 6), fmt=NF_INT, h="right")
        st(ws.cell(r, 7), fmt=NF_INT, h="right")
        M.var(key + "_q", q, DQE, f"E{r}", NF_2)
        M.var(key + "_pu", pu, DQE, f"F{r}", NF_INT)
        M.var(key + "_m", f"{{{key}_q}}*{{{key}_pu}}", DQE, f"G{r}", NF_INT)
        lots[-1][1].append(key)

    def subtotal(name):
        r = R[0]
        R[0] += 1
        st(ws.cell(r, 3, "Sous-total " + lots[-1][0].split("–")[0].strip()), bold=True, fill=C_ORANGE_CLAIR)
        for c in (2, 4, 5, 6):
            st(ws.cell(r, c), fill=C_ORANGE_CLAIR)
        st(ws.cell(r, 7), bold=True, fill=C_ORANGE_CLAIR, fmt=NF_INT)
        M.var(name, sum_of(lots[-1][1], "m"), DQE, f"G{r}", NF_INT)

    galva = '{protection}="Galvanisation à chaud"'
    lot("LOT 1 – CHARPENTE MÉTALLIQUE")
    item("1.1", "Fourniture et fabrication de la charpente (profilés, platines, boulonnerie, chutes)", "kg", "{kg_tot}", "{pu_acier}")
    item("1.2", "Transport et montage de la charpente", "kg", "{kg_tot}", "{pu_montage}")
    item("1.3", "Protection anticorrosion (galvanisation ou peinture, climat tropical humide)",
         f'=IF({galva},"kg","m²")', f"IF({galva},{{galva_kg}},{{peint_m2}})", f"IF({galva},{{pu_galva}},{{pu_peint}})")
    item("1.4", "Tiges d'ancrage avec écrous et rondelles", "u", "{tg_po_q}+{tg_pg_q}", "{pu_tige}")
    item("1.5", "Gabarits de pose des tiges d'ancrage", "u", "{gab_q}", "{pu_gab}")
    subtotal("st1")
    lot("LOT 2 – COUVERTURE")
    for no, lab, u, q, pu in [("2.1", "Couverture en tôles bac acier", "m²", "{tole_q}", "{pu_tole}"),
                              ("2.2", "Panneaux translucides", "m²", "{transl_q}", "{pu_transl}"),
                              ("2.3", "Faîtières", "ml", "{fait_q}", "{pu_fait}"),
                              ("2.4", "Rives", "ml", "{rive_q}", "{pu_rive}"),
                              ("2.5", "Larmiers et bavettes de couverture", "ml", "{larm_q}", "{pu_larm}"),
                              ("2.6", "Chéneaux d'eau", "ml", "{chen_q}", "{pu_chen}"),
                              ("2.7", "Naissances EP", "u", "{naiss_q}", "{pu_naiss}"),
                              ("2.8", "Descentes EP PVC Ø110", "ml", "{desc_q}", "{pu_desc}"),
                              ("2.9", "Skydomes avec chevêtres", "u", "{sky_q}", "{pu_sky}"),
                              ("2.10", "Vis autoforeuses", "u", "{vis_q}", "{pu_vis}"),
                              ("2.11", "Rivets de chéneaux", "u", "{riv_q}", "{pu_riv}")]:
        item(no, lab, u, q, pu)
    subtotal("st2")
    lot("LOT 3 – BARDAGE ET MENUISERIES MÉTALLIQUES")
    for no, lab, u, q, pu in [("3.1", "Bardage en bac acier", "m²", "{bard_q}", "{pu_bard}"),
                              ("3.2", "Larmiers de bardage (pied et angles)", "ml", "{larb_q}", "{pu_larm_b}"),
                              ("3.3", "Larmiers de portes coulissantes", "ml", "{larp_q}", "{pu_larm_p}"),
                              ("3.4", "Rails hauts de portes coulissantes", "ml", "{rail_q}", "{pu_rail}"),
                              ("3.5", "Portes coulissantes", "m²", "{nb_pc}*{l_pc}*{h_pc}", "{pu_pc}"),
                              ("3.6", "Portillons", "m²", "{nb_pt}*{l_pt}*{h_pt}", "{pu_pt}")]:
        item(no, lab, u, q, pu)
    subtotal("st3")
    lot("LOT 4 – TERRASSEMENTS")
    for no, lab, u, q, pu in [("4.1", "Décapage de l'emprise", "m²", "{dec_q}", "{pu_dec}"),
                              ("4.2", "Fouilles en puits et en rigole (profondeur 1,20 m)", "m³",
                               "{fp1_q}+{fp2_q}+{fr_q}", "{pu_fouil}"),
                              ("4.3", "Remblai en latérite compactée à 95 % OPM", "m³", "{rem_q}", "{pu_rem}"),
                              ("4.4", "Évacuation des déblais", "m³", "{evac_q}", "{pu_evac}")]:
        item(no, lab, u, q, pu)
    subtotal("st4")
    lot("LOT 5 – GROS ŒUVRE (matériaux et main d'œuvre)")
    items5 = [("5.1", "Ciment CPJ 42.5 (sacs de 50 kg)", "sacs", "{ciment}", "{pu_ciment}"),
              ("5.2", "Sable", "m³", "{sable}", "{pu_sable}"),
              ("5.3", "Gravier 5/15", "m³", "{g1}", "{pu_g1}"),
              ("5.4", "Gravier 15/25", "m³", "{g2}", "{pu_g2}")]
    items5 += [(f"5.{5 + i}", f"Aciers {d}", "kg", f"{{kg_{d}}}", f"{{pu_{d}}}") for i, d in enumerate(DIAMS)]
    n = 5 + len(DIAMS)
    items5 += [(f"5.{n}", "Agglos 15 pleins", "u", "{n_ap}", "{pu_ap}"),
               (f"5.{n + 1}", "Agglos 15 creux", "u", "{n_ac}", "{pu_ac}"),
               (f"5.{n + 2}", "Film polyane", "m²", "{polyane}", "{pu_poly}"),
               (f"5.{n + 3}", "Bois de coffrage", "m²", "{coffrage}", "{pu_coff}"),
               (f"5.{n + 4}", "Joints de retrait sciés", "ml", "{joint_q}", "{pu_joint}"),
               (f"5.{n + 5}", "Cure du béton", "m²", "{cure_q}", "{pu_cure}")]
    for it in items5:
        item(*it)
    mat = sum_of(lots[-1][1], "m")
    item(f"5.{n + 6}", "Main d'œuvre et matériel gros œuvre (% des matériaux du lot 5)", "Fft", "1",
         f"{{pct_mo}}*({mat})")
    subtotal("st5")
    R[0] += 1
    for name, lab, ex in [("total_ht", "TOTAL HT", "{st1}+{st2}+{st3}+{st4}+{st5}"),
                          ("tva", "TVA 18 %", "{total_ht}*0.18"),
                          ("total_ttc", "TOTAL TTC", "{total_ht}+{tva}")]:
        r = R[0]
        R[0] += 1
        st(ws.cell(r, 3, lab), bold=True, color="FFFFFF", fill=C_BLEU, size=11)
        for c in (2, 4, 5, 6):
            st(ws.cell(r, c), fill=C_BLEU)
        st(ws.cell(r, 7), bold=True, color="FFFFFF", fill=C_BLEU, fmt=NF_INT, size=11)
        M.var(name, ex, DQE, f"G{r}", NF_INT)
    r = R[0]
    st(ws.cell(r, 3, "Coût HT au m² couvert"), italic=True)
    st(ws.cell(r, 7), fmt=NF_INT, italic=True)
    M.var("cout_m2", "{total_ht}/{S_bat}", DQE, f"G{r}", NF_INT)
    r += 1
    st(ws.cell(r, 3, "Coût de la façade bardée (bardage, lisses et liernes, larmiers)"), italic=True)
    st(ws.cell(r, 7), fmt=NF_INT, italic=True)
    M.var("cout_fac_b", "{bard_q}*{pu_bard}+{larb_q}*{pu_larm_b}+({lis_lp_kg}+{lis_pg_kg}+{lll_kg}+{llp_kg})"
          "*(1+{pct_chutes})*({pu_acier}+{pu_montage}+{pu_prot_kg})", DQE, f"G{r}", NF_INT)
    r += 1
    st(ws.cell(r, 3, "Coût de la façade maçonnée (agglos, chaînages, raidisseurs, linteaux, enduit)"), italic=True)
    st(ws.cell(r, 7), fmt=NF_INT, italic=True)
    M.var("cout_fac", "{cout_fac_b}+{cout_fac_ag}", DQE, f"H{r}", NF_INT)
    M.disp(DQE, f"G{r}", "{cout_fac_ag}", NF_INT)
    ws.cell(r, 8).font = Font(name=FONT, size=7, color="BFBFBF")
    for col, w in zip("ABCDEFGH", [2, 7, 70, 8, 14, 12, 18, 4]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    page_setup(ws)


# ============================================================================= MOTEUR GABARITS
ENG_COL0 = 4  # colonne D


def eng_col(g, s):
    return get_column_letter(ENG_COL0 + g * 3 + s)


def assign_erows():
    for i, n in enumerate(M.vars):
        M.vars[n]["erow"] = 5 + i


def build_engine(wb):
    ws = wb.create_sheet(CA)
    ws["A1"] = "MOTEUR DE CALCUL DES 7 GABARITS (feuille masquée) – mêmes règles que les feuilles visibles"
    ws["A1"].font = Font(name=FONT, bold=True)
    ws["A2"], ws["A3"] = "Gabarit", "Scénario de façade"
    for g, (n, _, _) in enumerate(GABARITS):
        for s in range(3):
            c = eng_col(g, s)
            ws[f"{c}2"], ws[f"{c}3"] = n, SCENARIOS[s]
    names = list(M.vars)
    for i, n in enumerate(names):
        v, r = M.vars[n], 5 + i
        ws.cell(r, 1, n)
        ws.cell(r, 2, f"{v['sheet']}!{v['coord']}")
        for g in range(len(GABARITS)):
            for s in range(3):
                c = eng_col(g, s)
                ex = v["eng"](g, s) if v["eng"] else v["expr"]
                ws[f"{c}{r}"] = "=" + M.render(ex, c)
    ws.sheet_state = "hidden"
    protect(ws)


def eng_ref(name, g, s=0):
    return f"'{CA}'!${eng_col(g, s)}${M.vars[name]['erow']}"


# ============================================================================= RÉCAP
def build_recap(wb):
    ws = wb.create_sheet(REC)
    title(ws, "RÉCAPITULATIF COMPARATIF DES 7 GABARITS",
          "Calculé avec les paramètres en vigueur (hauteur, pente, prix…) et les profilés automatiques. "
          "Façades : variante bardage bac acier et variante agglos 15 creux.", 12)
    cols = [("Tonnage acier (t)", "t_tot", NF_2), ("kg/m²", "kg_m2", NF_2), ("Tôles (m²)", "tole_q", NF_2),
            ("Béton (m³)", "V_beton", NF_2), ("Aciers HA (kg)", "kg_HA", NF_2), ("Coût HT (FCFA)", "total_ht", NF_INT),
            ("Coût HT / m²", "cout_m2", NF_INT)]
    header(ws, 4, ["Gabarit", "Système", "Surface (m²)"] + [c[0] for c in cols] +
           ["Coût façade bardage", "Coût façade agglos"])
    for g, (n, b, l) in enumerate(GABARITS):
        r = 5 + g
        st(ws.cell(r, 2, n), bold=True, h="center")
        ws.cell(r, 3, "=" + eng_ref("systeme", g))
        st(ws.cell(r, 3), size=9)
        ws.cell(r, 4, "=" + eng_ref("S_bat", g))
        st(ws.cell(r, 4), fmt=NF_2)
        for j, (_, name, fmt) in enumerate(cols):
            c = ws.cell(r, 5 + j, "=" + eng_ref(name, g))
            st(c, fmt=fmt)
        st(ws.cell(r, 12, "=" + eng_ref("cout_fac", g, 1)), fmt=NF_INT)
        st(ws.cell(r, 13, "=" + eng_ref("cout_fac", g, 2)), fmt=NF_INT)
    red_if(ws, "F5:F11", "OR(F5<18,F5>45)")
    r = 13
    st(ws.cell(r, 2, "Contrôle de cohérence (gabarit actif) : écart entre le moteur et le DQE"), bold=True,
       border=False)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    gi = f"MATCH({{gab}},'{PAR}'!$G$5:$G$11,0)"
    row_ht = M.vars["total_ht"]["erow"]
    first, last = eng_col(0, 0), eng_col(len(GABARITS) - 1, 2)
    M.disp(REC, f"H{r}",
           f'IFERROR(IF(ABS(INDEX(\'{CA}\'!${first}${row_ht}:${last}${row_ht},1,({gi}-1)*3+1)-{{total_ht}})<1,'
           f'"OK (écart nul)","ÉCART : vérifier les forçages de profilés"),"PERSONNALISÉ : sans objet")')
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=11)
    st(ws.cell(r, 8), bold=True)
    red_if(ws, f"H{r}", f'LEFT(H{r},5)="ÉCART"')
    st(ws.cell(r + 1, 2, "Le comparatif utilise les profilés automatiques : un forçage en PRÉDIMENSIONNEMENT ne s'applique "
                         "qu'au gabarit actif. Coûts façade = bardage + lisses + liernes + larmiers, ou agglos + "
                         "chaînages + raidisseurs + linteaux (+ enduit), main d'œuvre comprise."),
       size=8, italic=True, border=False)
    for col, w in zip("ABCDEFGHIJKLM", [2, 11, 26, 11, 11, 9, 11, 11, 13, 15, 12, 16, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    page_setup(ws)


# ============================================================================= ACCUEIL
def build_accueil(ws, G_RANGE, alerts):
    title(ws, "QUANTITATIF – ESTIMATION – ESQUISSES : BÂTIMENT INDUSTRIEL EN CHARPENTE MÉTALLIQUE", None, 8)
    ws["A1"].font = Font(name=FONT, bold=True, size=12, color="FFFFFF")
    st(ws["A2"], border=False)
    ws["B2"] = ("Charpente : EC3 / EC1 (vent, sans neige) · Béton armé et fondations : BAEL 91 mod. 99 / EC2 · "
                "Contexte : Côte d'Ivoire – sol latéritique, climat tropical humide · Monnaie : FCFA · TVA 18 %")
    st(ws["B2"], size=9, italic=True, color=C_ORANGE, border=False)
    ws["B4"] = "MODE D'EMPLOI : 1) choisir un gabarit ci-dessous ; 2) ajuster les cellules JAUNES de PARAMÈTRES ; " \
               "3) lire les résultats ici, dans le DQE et dans PLANS."
    st(ws["B4"], size=9, bold=True, border=False)
    st(ws.cell(6, 2, "GABARIT (portée × longueur)"), bold=True, fill=C_ORANGE_CLAIR)
    c = ws.cell(6, 4, GAB_DEFAUT[0])
    yellow(c)
    c.font = Font(name=FONT, bold=True, size=13)
    c.alignment = Alignment(horizontal="center")
    dv_list(ws, ["D6"], ref="=" + G_RANGE)
    M.param("gab", ACC, "D6")
    st(ws.cell(6, 5, "← liste déroulante (PERSONNALISÉ : dimensions dans PARAMÈTRES)"), size=8, italic=True,
       border=False)
    st(ws.cell(7, 2, "Légende : cellule jaune = saisie ; cellule bleu clair = résultat ; rouge = contrôle en alerte"),
       size=8, italic=True, border=False)
    rows = [("Portée", "{portee}", "m", NF_2), ("Longueur", "{longueur}", "m", NF_2),
            ("Hauteur sous sablière", "{hauteur}", "m", NF_2), ("Hauteur au faîtage", "{Hf}", "m", NF_2),
            ("Pente", "{pente}", "%", NF_PCT), ("Système structurel", "{systeme}", "", "@"),
            ("Nombre de portiques", "{N}", "u", NF_INT), ("Type de façade", "{facade}", "", "@"),
            ("SURFACE COUVERTE", "{S_bat}", "m²", NF_2), ("TONNAGE ACIER CHARPENTE", "{t_tot}", "t", NF_2),
            ("Ratio acier", "{kg_m2}", "kg/m²", NF_2), ("Tôles de couverture", "{tole_q}", "m²", NF_2),
            ("BÉTON (propreté + armé)", "{V_beton}", "m³", NF_2), ("Aciers HA", "{kg_HA}", "kg", NF_2),
            ("COÛT TOTAL HT", "{total_ht}", "FCFA", NF_INT), ("TVA 18 %", "{tva}", "FCFA", NF_INT),
            ("COÛT TOTAL TTC", "{total_ttc}", "FCFA", NF_INT), ("Coût HT au m²", "{cout_m2}", "FCFA/m²", NF_INT)]
    header(ws, 9, ["Résumé", "", "Valeur", "Unité"])
    for i, (lab, ex, u, fmt) in enumerate(rows):
        r = 10 + i
        big = lab.isupper()
        st(ws.cell(r, 2, lab), bold=big)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        st(ws.cell(r, 4), fmt=fmt, bold=True, fill=C_BLEU_CLAIR, h="right", size=11 if big else 10)
        st(ws.cell(r, 5, u), size=9)
        M.disp(ACC, f"D{r}", ex, fmt)
        if lab == "Ratio acier":
            red_if(ws, f"D{r}", f"OR(D{r}<18,D{r}>45)")
        if lab == "Pente":
            red_if(ws, f"D{r}", f"D{r}<0.05")
    r = 10 + len(rows) + 1
    header(ws, r, ["Contrôles automatiques", "", "", ""])
    for i, a in enumerate(alerts[:5]):
        rw = r + 1 + i
        ws.merge_cells(start_row=rw, start_column=2, end_row=rw, end_column=7)
        st(ws.cell(rw, 2), bold=True)
        M.disp(ACC, f"B{rw}", a)
        red_if(ws, f"B{rw}", f'LEFT(B{rw},6)="ALERTE"')
    rw = r + 7
    for i, t in enumerate(["Feuilles : PARAMÈTRES (saisies) · BASE_PROFILÉS / BASE_ACIERS (poids) · PRÉDIMENSIONNEMENT · "
                           "QUANTITATIF_CHARPENTE · GROS_ŒUVRE · DQE_DEVIS · RÉCAP_GABARITS · PLANS (7 esquisses dynamiques).",
                           "Les esquisses sont des graphiques Nuage de points alimentés par la feuille masquée COORD_PLANS : "
                           "aucune macro, compatible Excel mobile (Microsoft 365).",
                           "Feuilles protégées sans mot de passe (Révision > Ôter la protection) ; seules les cellules jaunes "
                           "sont modifiables.",
                           "Estimation de prédimensionnement : à confirmer par une note de calcul (EC3 / EC1 / BAEL) avant exécution."]):
        st(ws.cell(rw + i, 2, t), size=8, italic=True, border=False)
    for col, w in zip("ABCDEFGH", [2, 30, 14, 22, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    page_setup(ws, rows_title=None, fit_height=1)


# ============================================================================= ÉCRITURE
def write_all(wb):
    for name, v in M.vars.items():
        ws = wb[v["sheet"]]
        c = ws[v["coord"]]
        c.value = "=" + M.render(v["expr"], "A")
        if v["fmt"]:
            c.number_format = v["fmt"]
    for sheet, coord, ex, fmt in M.disps:
        c = wb[sheet][coord]
        c.value = "=" + M.render(ex, "A")
        if fmt:
            c.number_format = fmt


def main(out, gabarit="20 × 40", recalc=True, plans_seuls=False, params=None):
    GAB_DEFAUT[0] = gabarit
    PARAM_TEST.update(params or {})
    wb = Workbook()
    acc = wb.active
    acc.title = ACC
    G_RANGE = build_parametres(wb)
    build_bases(wb)
    alerts = build_predim(wb)
    build_qc(wb)
    build_go(wb)
    build_dqe(wb)
    assign_erows()
    build_recap(wb)
    build_accueil(acc, G_RANGE, alerts)
    plans.build(wb, M, dict(PL=PL, CO=CO, PAR=PAR), st, section, page_setup, protect)
    build_engine(wb)
    write_all(wb)
    order = [ACC, PAR, BP, BA, PRE, QC, GO, DQE, REC, PL, CO, CA]
    wb._sheets = [wb[n] for n in order]
    for n in order:
        if n not in (CA,):
            protect(wb[n])
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.title = "Quantitatif charpente métallique"
    for name, val in PARAM_TEST.items():   # surcharge de paramètres (tests)
        sh, co = M.params[name].replace("$", "").split("!")
        wb[sh.strip("'")][co].value = val
    if plans_seuls:
        for n in order:
            if n != PL:
                wb[n].sheet_state = "hidden"
        wb.active = order.index(PL)
    wb.save(out)
    print("Classeur écrit :", out, "–", len(M.vars), "variables moteur")
    return finaliser.finaliser(out, recalc)


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("sortie", nargs="?", default="Quantitatif_Charpente_Metallique.xlsx")
    ap.add_argument("--gabarit", default="20 × 40", help="gabarit sélectionné à l'ouverture")
    ap.add_argument("--sans-recalcul", action="store_true", help="ne pas injecter les valeurs calculées")
    ap.add_argument("--plans-seuls", action="store_true", help="(test) masque les autres feuilles")
    ap.add_argument("--param", action="append", default=[], help="(test) nom=valeur")
    a = ap.parse_args()
    pv = {}
    for kv in a.param:
        k, v = kv.split("=", 1)
        try:
            v = float(v)
        except ValueError:
            pass
        pv[k] = v
    main(a.sortie, a.gabarit, not a.sans_recalcul, a.plans_seuls, pv)
